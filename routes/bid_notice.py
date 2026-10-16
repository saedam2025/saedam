"""업무지원 > 입찰공고관리.

나라장터(조달청) 입찰공고정보서비스에서 방과후·늘봄 관련 용역 입찰공고를 담당자가 [공고가져오기]를 누를 때마다
(최근 4주치) 찾아 쌓아 둔다. 자동 수집은 없다. 검색어로 찾은 공고 중 지역(기본 경기도)·기관(기본 초등학교) 조건에
맞는 것만 저장하고, 이미 쌓아 둔 공고는 다시 기록하지 않으며, 정정·취소 공고가 나오면 같은 공고에 이력으로 남긴다.
새 소식은 설정 화면에서 등록한 메일 주소(최대 5곳)로 직접 보낼 수 있다.

- 조회·해석: ``services/g2b_bid.py``
- 설정(인증키·검색어·지역·받는 메일)은 ``admin_settings``의 ``bid_notice_settings``에 저장하고 인증키는 암호화한다.
- 공고마다 진행 상태(입찰예정 등)를 고른다. 낙찰 결과는 [공고가져오기] 창의 [낙찰 결과 가져오기]를 누를 때만
  개찰일시가 지난 공고를 나라장터 낙찰정보에서 찾아 저장한다(자동으로는 가져오지 않는다).
- 첨부파일은 [첨부 받기]를 누르는 순간 나라장터에서 받아 압축파일 하나로 바로 내려 준다(서버에 보관하지 않는다).
"""

from __future__ import annotations

import base64
import hashlib
import json
import os
import re
import shutil
import smtplib
import tempfile
import threading
import zipfile
from datetime import datetime, timedelta, timezone
from email.header import Header
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.utils import formataddr
from html import escape
from io import BytesIO

from cryptography.fernet import Fernet, InvalidToken
from flask import Blueprint, jsonify, render_template, request, send_file, session

from services import g2b_bid as g2b

from .contacts import _mail_credentials
from .database import get_db
from .security import admin_required, is_admin_session, load_credential_secret
from .storage import DATA_ROOT

bid_notice_bp = Blueprint('bid_notice', __name__)

KST = timezone(timedelta(hours=9))
SETTINGS_KEY = 'bid_notice_settings'
ENV_SERVICE_KEY = 'G2B_SERVICE_KEY'
INTRANET_URL = 'https://works.saedam.org/bid-notice/'

# 공고명은 학교마다 표현이 달라("방과후학교 프로그램 운영 용역" 등) 긴 문구 대신 핵심 낱말로 찾는다.
DEFAULT_KEYWORDS = ('방과후', '늘봄', '맞춤형', '온동네 돌봄')
DEFAULT_EXCLUDE_WORDS = ('급식', '통학차량', '폐기물')
# 기관 조건(초등학교)과 상관없이 모으는 검색어: 학교가 아니라 교육지원청이 발주하는 사업
DEFAULT_ORG_FREE_KEYWORDS = ('온동네 돌봄',)
DEFAULT_REGION = '경기도'
DEFAULT_ORG_WORD = '초등학교'
MAX_RECIPIENTS = 5
SETTINGS_MAX_LEVEL = 1           # [공고수집·메일 설정](인증키·수집 조건·메일)은 레벨 1만 연다(admin 계정 포함)
EMAIL_PATTERN = re.compile(r'^[^@\s]+@[^@\s]+\.[^@\s]{2,}$')

COLLECT_DAYS = g2b.MAX_RANGE_DAYS    # [최근 4주 공고 가져오기]가 보는 기간(28일)
RESULT_LIMIT = 100               # [낙찰 결과 가져오기] 요청 한 번에 확인하는 공고 수(화면이 남은 공고가 없을 때까지 이어 부른다)
BASE_AMOUNT_LIMIT = 60           # 한 번 수집할 때 기초금액을 확인할 공고 수
BASE_AMOUNT_AFTER_CLOSE_DAYS = 14
LIST_LIMIT = 500
MAX_ARCHIVES = 30                # 만들 수 있는 보관함(탭) 수
MAX_ARCHIVE_NAME = 30
PAGE_SIZES = (30, 60, 100)       # 한 쪽에 보여 줄 공고 수(서버가 쪽을 나눠 보낸다). per_page=all 이면 쪽을 나누지 않고 모두 보낸다
EXPORT_LIMIT = 5000              # 엑셀로 한 번에 내려받을 수 있는 최대 건수
MAX_RANGE_COLLECT_DAYS = 366     # 날짜를 정해 수집할 때 한 번에 볼 수 있는 기간
WEEK_DAYS = 7                    # [1주내 마감]: 마감일이 오늘부터 이 일수 안인 공고
SOON_DAYS = 3                    # [3일내 마감] 탭: 마감일이 오늘부터 이 일수 안인 공고
CANCEL_KEEP_DAYS = 7             # 마감일이 없는 취소 공고를 [진행 중] 탭에 남겨 두는 기간
NEW_MARK_HOURS = 24              # [공고가져오기]로 새로 기록된 공고에 반짝이는 N을 붙여 두는 시간
RUN_KEEP_DAYS = 60              # 수집 기록은 이 기간이 지나면 새 수집을 시작할 때 자동으로 지운다

# 담당자가 공고마다 고르는 진행 상태(비어 있으면 '상태선택')
WORK_STATUSES = ('입찰예정', '본부입찰', '지점입찰', '입찰제외', '서류접수완료', '가격투찰완료', '설명회완료', '개찰완료')

# 첨부파일은 서버에 두지 않는다. 예전 버전이 압축파일을 보관하던 폴더는 시작할 때 비운다.
LEGACY_FILES_DIR = DATA_ROOT / 'bid_notice_files'
ZIP_MEMORY_LIMIT = 20 * 1024 * 1024     # 이보다 큰 압축파일은 메모리 대신 임시 파일에 만든다
MAX_ATTACH_FILE_BYTES = 100 * 1024 * 1024
MAX_ATTACH_TOTAL_BYTES = 300 * 1024 * 1024
MAX_BUNDLE_NOTICES = 30                       # 체크한 공고를 한꺼번에 받을 수 있는 최대 건수
MAX_BUNDLE_TOTAL_BYTES = 600 * 1024 * 1024    # 한꺼번에 받는 압축파일 전체 크기 상한

EVENT_LABELS = {'new': '신규', 'changed': '정정', 'cancelled': '취소'}
DIFF_FIELDS = (
    ('title', '공고명'), ('bid_begin', '입찰 시작'), ('bid_close', '입찰 마감'),
    ('open_dt', '개찰'), ('budget', '배정예산'), ('est_price', '추정가격'),
)
NOTICE_COLUMNS = (
    'bid_ord', 'kind', 'title', 'org_name', 'demand_name', 'region', 'school', 'notice_dt',
    'bid_begin', 'bid_close', 'open_dt', 'budget', 'est_price', 'contract_method',
    'award_method', 'official', 'tel', 'url', 'change_reason', 'attachments',
)
# 처음 만든 뒤에 더한 열(이미 운영 중인 테이블에는 ALTER TABLE로 붙인다)
EXTRA_NOTICE_COLUMNS = {
    'attachments': 'TEXT',            # 첨부파일 목록 JSON [{name, url}] (None이면 아직 모름)
    'work_status': 'TEXT',
    'status_updated_by': 'TEXT',
    'status_updated_at': 'TEXT',
    'winner_name': 'TEXT',
    'winner_amount': 'INTEGER',
    'winner_rate': 'TEXT',            # 낙찰률(%) — 나라장터가 준 글자 그대로
    'result_state': 'TEXT',           # 최종낙찰 / 개찰완료(최종 낙찰 전) / 유찰 등
    'result_participants': 'INTEGER',
    'result_opened_at': 'TEXT',
    'result_final_date': 'TEXT',
    'result_updated_by': 'TEXT',
    'result_updated_at': 'TEXT',
    'result_checked_at': 'TEXT',      # 결과를 마지막으로 찾아본 때(결과가 없었어도 적는다)
    'archive_id': 'INTEGER',          # 보관함(탭) 번호. 비어 있으면 [공고 목록]에 보인다
    'archived_at': 'TEXT',
    'changed_at': 'TEXT',             # 마지막으로 정정·취소를 확인한 때(목록 맨 위로 올리는 기준)
    'change_summary': 'TEXT',         # 그때 바뀐 내용('입찰 마감: A → B')
    'bidders': 'TEXT',                # 참가 업체 전체 JSON [{rank, name, bizno, ceo, amount, rate, note}] (None이면 아직 안 읽음)
}

_collect_lock = threading.Lock()
_attach_lock = threading.Lock()


# ---------------------------------------------------------------------------
# 스키마
# ---------------------------------------------------------------------------
def init_bid_notice_schema(conn=None):
    owns_connection = conn is None
    if owns_connection:
        conn = get_db()
    try:
        conn.executescript('''
            CREATE TABLE IF NOT EXISTS admin_settings (
                key TEXT PRIMARY KEY,
                value TEXT,
                updated_at DATETIME DEFAULT CURRENT_TIMESTAMP
            );

            CREATE TABLE IF NOT EXISTS bid_notices (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                bid_no TEXT NOT NULL UNIQUE,
                bid_ord TEXT NOT NULL DEFAULT '000',
                kind TEXT,
                title TEXT NOT NULL,
                org_name TEXT,
                demand_name TEXT,
                region TEXT,
                school TEXT,
                notice_dt TEXT,
                bid_begin TEXT,
                bid_close TEXT,
                open_dt TEXT,
                budget INTEGER,
                est_price INTEGER,
                base_amount INTEGER,
                base_opened_at TEXT,
                base_checked_at TEXT,
                contract_method TEXT,
                award_method TEXT,
                official TEXT,
                tel TEXT,
                url TEXT,
                change_reason TEXT,
                matched_keyword TEXT,
                first_seen_at TEXT,
                updated_at TEXT
            );
            CREATE INDEX IF NOT EXISTS idx_bid_notices_close ON bid_notices(bid_close);

            CREATE TABLE IF NOT EXISTS bid_notice_events (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                notice_id INTEGER NOT NULL,
                bid_no TEXT NOT NULL,
                bid_ord TEXT NOT NULL,
                event_type TEXT NOT NULL,
                kind TEXT,
                summary TEXT,
                reason TEXT,
                created_at TEXT,
                mailed_at TEXT,
                UNIQUE (bid_no, bid_ord, event_type),
                FOREIGN KEY (notice_id) REFERENCES bid_notices(id)
            );
            CREATE INDEX IF NOT EXISTS idx_bid_notice_events_notice ON bid_notice_events(notice_id);

            CREATE TABLE IF NOT EXISTS bid_collect_runs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                slot_key TEXT UNIQUE,
                trigger_type TEXT NOT NULL,
                actor TEXT,
                started_at TEXT,
                finished_at TEXT,
                status TEXT NOT NULL DEFAULT 'running',
                range_begin TEXT,
                range_end TEXT,
                fetched INTEGER DEFAULT 0,
                new_count INTEGER DEFAULT 0,
                changed_count INTEGER DEFAULT 0,
                cancelled_count INTEGER DEFAULT 0,
                mail_status TEXT,
                message TEXT
            );

            CREATE TABLE IF NOT EXISTS bid_archives (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL UNIQUE COLLATE NOCASE,
                created_by TEXT,
                created_at TEXT
            );

            -- 직접 삭제한 공고의 번호. [공고가져오기]가 같은 공고를 다시 가져오지 않게 한다.
            CREATE TABLE IF NOT EXISTS bid_notice_deleted (
                bid_no TEXT PRIMARY KEY,
                deleted_at TEXT,
                deleted_by TEXT
            );
        ''')
        _add_missing_columns(conn)
        _ensure_indexes(conn)
        conn.commit()
    finally:
        if owns_connection:
            conn.close()
    _remove_legacy_files()


# 보관함 탭 · 목록 기본 정렬(게시일) · 현황 요약(오늘 새 소식, 안 보낸 소식)
INDEXES = (
    ('idx_bid_notices_archive', 'bid_notices', 'archive_id', ''),
    ('idx_bid_notices_notice_dt', 'bid_notices', 'notice_dt', ''),
    ('idx_bid_notice_events_created', 'bid_notice_events', 'created_at', ''),
    ('idx_bid_notice_events_unsent', 'bid_notice_events', 'mailed_at', ' WHERE mailed_at IS NULL'),
)


def _ensure_indexes(conn):
    """색인을 만든다(아주 예전 테이블에 열이 없으면 그 색인은 건너뛴다)."""
    columns = {}
    for name, table, column, where in INDEXES:
        if table not in columns:
            columns[table] = {row['name'] for row in conn.execute(f'PRAGMA table_info({table})')}
        if column in columns[table]:
            conn.execute(f'CREATE INDEX IF NOT EXISTS {name} ON {table}({column}){where}')


def _add_missing_columns(conn):
    columns = {row['name'] for row in conn.execute('PRAGMA table_info(bid_notices)')}
    for name, ddl in EXTRA_NOTICE_COLUMNS.items():
        if name not in columns:
            conn.execute(f'ALTER TABLE bid_notices ADD COLUMN {name} {ddl}')
    if 'changed_at' not in columns:
        # 열을 처음 붙일 때 예전에 쌓인 정정·취소 이력에서 마지막 소식을 채워 둔다.
        latest = ("FROM bid_notice_events e WHERE e.notice_id = bid_notices.id "
                  "AND e.event_type IN ('changed', 'cancelled')")
        conn.execute(f'''UPDATE bid_notices
                            SET changed_at = (SELECT MAX(e.created_at) {latest}),
                                change_summary = (SELECT e.summary {latest} ORDER BY e.id DESC LIMIT 1)
                          WHERE EXISTS (SELECT 1 {latest})''')


def _remove_legacy_files():
    """예전 버전이 서버에 보관하던 첨부 압축파일 폴더를 지운다(지금은 누를 때 바로 내려 준다)."""
    if LEGACY_FILES_DIR.is_dir():
        shutil.rmtree(LEGACY_FILES_DIR, ignore_errors=True)


# ---------------------------------------------------------------------------
# 도우미
# ---------------------------------------------------------------------------
def _now():
    """한국시간 지금(시간대 정보 없는 datetime — 나라장터 시각·저장 문자열과 바로 비교한다)."""
    return datetime.now(KST).replace(tzinfo=None, microsecond=0)


def _stamp(moment=None):
    return (moment or _now()).strftime('%Y-%m-%d %H:%M:%S')


def _json_error(message, code=400):
    return jsonify({'status': 'error', 'message': message}), code


def _fernet():
    digest = hashlib.sha256(load_credential_secret().encode('utf-8')).digest()
    return Fernet(base64.urlsafe_b64encode(digest))


def _won(value):
    return f'{int(value):,}원' if value else ''


def _school_label(row):
    """[학교] 칸에 보여 줄 이름. 학교가 아니면 수요기관 이름에서 앞의 시·도 교육청을 뺀다
    ('경기도교육청 경기도수원교육지원청' → '경기도수원교육지원청'). 뒤에 이름이 없으면 그대로 둔다."""
    if row['school']:
        return row['school']
    words = str(row['demand_name'] or '').split()
    if len(words) > 1 and words[0].endswith('교육청'):
        words = words[1:]
    return ' '.join(words)


def _short_dt(value):
    """'2025-12-22 10:00:00' → '2025-12-22 10:00'"""
    return str(value or '')[:16]


# ---------------------------------------------------------------------------
# 설정
# ---------------------------------------------------------------------------
def clean_recipients(values):
    """받는 메일 주소를 검사해 정리한다. 문제가 있으면 ValueError(안내 문구)."""
    if isinstance(values, str):
        values = re.split(r'[,;\s]+', values)
    result = []
    for value in values or []:
        email = str(value or '').strip()
        if not email:
            continue
        if not EMAIL_PATTERN.match(email) or len(email) > 120:
            raise ValueError(f'메일 주소를 확인해 주세요: {email}')
        if email.lower() not in [item.lower() for item in result]:
            result.append(email)
    if len(result) > MAX_RECIPIENTS:
        raise ValueError(f'받는 메일 주소는 {MAX_RECIPIENTS}개까지 등록할 수 있습니다.')
    return result


def _stored_settings(conn):
    row = conn.execute('SELECT value FROM admin_settings WHERE key=?', (SETTINGS_KEY,)).fetchone()
    try:
        data = json.loads(row['value']) if row and row['value'] else {}
    except (TypeError, ValueError):
        data = {}
    return data if isinstance(data, dict) else {}


def get_settings(conn=None):
    """저장값과 기본값을 합친 설정. 인증키는 저장값이 환경변수보다 우선한다."""
    owns_connection = conn is None
    if owns_connection:
        conn = get_db()
    try:
        stored = _stored_settings(conn)
    finally:
        if owns_connection:
            conn.close()

    service_key, key_source = '', 'none'
    token = str(stored.get('service_key_enc') or '').strip()
    if token:
        try:
            service_key = _fernet().decrypt(token.encode('ascii')).decode('utf-8')
            key_source = 'menu'
        except (InvalidToken, ValueError, TypeError):
            service_key = ''
    if not service_key:
        service_key = os.environ.get(ENV_SERVICE_KEY, '').strip()
        key_source = 'env' if service_key else 'none'

    def words(name, default):
        return g2b.clean_words(stored[name]) if isinstance(stored.get(name), list) else list(default)

    return {
        'service_key': service_key,
        'key_source': key_source,
        'keywords': words('keywords', DEFAULT_KEYWORDS) or list(DEFAULT_KEYWORDS),
        'exclude_words': words('exclude_words', DEFAULT_EXCLUDE_WORDS),
        'org_free_keywords': words('org_free_keywords', DEFAULT_ORG_FREE_KEYWORDS),
        'region': str(stored['region']).strip() if 'region' in stored else DEFAULT_REGION,
        'org_word': str(stored['org_word']).strip() if 'org_word' in stored else DEFAULT_ORG_WORD,
        'recipients': [str(item) for item in stored.get('recipients') or []][:MAX_RECIPIENTS],
        'updated_by': str(stored.get('updated_by') or ''),
        'updated_at': str(stored.get('updated_at') or ''),
    }


def public_settings(settings):
    """화면에 보여 줄 설정(인증키는 끝 4자리만)."""
    key = settings['service_key']
    sender = _mail_credentials()[0]
    result = {name: settings[name] for name in (
        'keywords', 'exclude_words', 'org_free_keywords', 'region', 'org_word', 'recipients',
        'updated_by', 'updated_at', 'key_source',
    )}
    result.update({
        'has_key': bool(key),
        'masked_key': f'••••••••{key[-4:]}' if key else '',
        'mail_sender': sender,
        'max_recipients': MAX_RECIPIENTS,
    })
    return result


def save_settings(payload, actor=''):
    """설정을 저장한다. 인증키 칸을 비워 두면 기존 키를 그대로 둔다."""
    keywords = g2b.clean_words(payload.get('keywords'))
    if not keywords:
        raise ValueError('검색어를 한 개 이상 입력해 주세요.')
    recipients = clean_recipients(payload.get('recipients'))
    new_key = str(payload.get('service_key') or '').strip()
    if new_key and (len(new_key) < 20 or len(new_key) > 400 or re.search(r'\s', new_key)):
        raise ValueError('인증키 형식을 확인해 주세요. 공공데이터포털의 [일반 인증키]를 그대로 붙여 넣습니다.')

    conn = get_db()
    try:
        stored = _stored_settings(conn)
        if payload.get('clear_key'):
            stored.pop('service_key_enc', None)
        elif new_key:
            stored['service_key_enc'] = _fernet().encrypt(new_key.encode('utf-8')).decode('ascii')
        for old_name in ('auto_enabled', 'mail_enabled'):      # 없앤 자동 수집·자동 메일 설정
            stored.pop(old_name, None)
        stored.update({
            'keywords': keywords,
            'exclude_words': g2b.clean_words(payload.get('exclude_words')),
            # 예전 화면이 보낸 저장 요청(이 항목이 없음)은 기존 값을 그대로 둔다.
            **({'org_free_keywords': g2b.clean_words(payload.get('org_free_keywords'))}
               if 'org_free_keywords' in payload else {}),
            'region': str(payload.get('region') or '').strip()[:20],
            'org_word': str(payload.get('org_word') or '').strip()[:20],
            'recipients': recipients,
            'updated_by': str(actor or ''),
            'updated_at': _stamp(),
        })
        conn.execute('''
            INSERT INTO admin_settings (key, value, updated_at)
            VALUES (?, ?, CURRENT_TIMESTAMP)
            ON CONFLICT(key) DO UPDATE SET value=excluded.value, updated_at=CURRENT_TIMESTAMP
        ''', (SETTINGS_KEY, json.dumps(stored, ensure_ascii=False)))
        conn.commit()
        return get_settings(conn)
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# 수집
# ---------------------------------------------------------------------------
def _diff_summary(before, after):
    """정정 전후로 달라진 항목을 '입찰 마감: A → B' 꼴로 적는다."""
    parts = []
    for field, label in DIFF_FIELDS:
        old, new = before[field], after[field]
        if (old or '') == (new or ''):
            continue
        if field in ('budget', 'est_price'):
            old, new = _won(old) or '없음', _won(new) or '없음'
        elif field != 'title':
            old, new = _short_dt(old) or '없음', _short_dt(new) or '없음'
        parts.append(f'{label}: {old} → {new}')
    return ' / '.join(parts)


def _add_event(conn, notice_id, notice, event_type, summary, now_text, mailed_at):
    cursor = conn.execute(
        '''INSERT OR IGNORE INTO bid_notice_events
               (notice_id, bid_no, bid_ord, event_type, kind, summary, reason, created_at, mailed_at)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)''',
        (notice_id, notice['bid_no'], notice['bid_ord'], event_type, notice['kind'],
         summary, notice['change_reason'], now_text, mailed_at),
    )
    return bool(cursor.rowcount)


def store_items(conn, items, settings, baseline=False, restore_deleted=False):
    """조회한 공고를 저장하고 새로 생긴 소식 수를 돌려준다.

    공고번호 하나가 공고 한 건이고, 차수(000 → 001 …)가 올라가면 정정·취소로 본다.
    ``baseline``(맨 처음 수집)이면 이미 지난 소식이라 메일 대상에서 뺀다.
    ``restore_deleted``면 직접 삭제한 공고도 다시 가져오고 삭제 기록을 지운다(이미 본 공고라 메일 대상에서 뺀다).
    """
    counts = {'new': 0, 'changed': 0, 'cancelled': 0, 'restored': 0}
    now_text = _stamp()
    mailed_at = now_text if baseline else None
    parsed = {}
    for item in items:
        notice = g2b.parse_notice(item)
        if notice['bid_no'] and notice['title']:
            parsed[(notice['bid_no'], notice['bid_ord'])] = notice

    deleted = {row['bid_no'] for row in conn.execute('SELECT bid_no FROM bid_notice_deleted')}
    for key in sorted(parsed):
        notice = parsed[key]
        restoring = notice['bid_no'] in deleted
        if restoring and not restore_deleted:
            continue
        existing = conn.execute('SELECT * FROM bid_notices WHERE bid_no=?', (notice['bid_no'],)).fetchone()
        cancelled = notice['kind'] == g2b.KIND_CANCEL
        if existing is None:
            keyword = g2b.passes_filters(
                notice, settings['keywords'], settings['exclude_words'],
                settings['region'], settings['org_word'], settings['org_free_keywords'],
            )
            if not keyword:
                continue
            if restoring:
                # 되살린 공고: 삭제 기록을 지우고, 예전에 이미 알린 공고라 메일 대상에서 뺀다.
                conn.execute('DELETE FROM bid_notice_deleted WHERE bid_no=?', (notice['bid_no'],))
                deleted.discard(notice['bid_no'])
                counts['restored'] += 1
            columns = ('bid_no',) + NOTICE_COLUMNS + ('matched_keyword', 'first_seen_at', 'updated_at')
            values = [notice['bid_no']] + [notice[name] for name in NOTICE_COLUMNS] + [keyword, now_text, now_text]
            cursor = conn.execute(
                f"INSERT INTO bid_notices ({', '.join(columns)}) VALUES ({', '.join('?' * len(columns))})",
                values,
            )
            event_type = 'cancelled' if cancelled else 'new'
            if _add_event(conn, cursor.lastrowid, notice, event_type, '', now_text,
                          now_text if restoring else mailed_at):
                counts[event_type] += 1
            continue

        # 이미 쌓아 둔 공고는 조건과 상관없이 끝까지 따라간다(정정으로 공고명이 바뀌어도 놓치지 않는다).
        if notice['bid_ord'] <= existing['bid_ord']:
            continue
        # 정정·취소는 제자리에서 갱신한다: 게시일은 처음 공고의 것을 지켜 목록 순서가 바뀌지 않게 한다.
        notice = {**notice, 'notice_dt': existing['notice_dt'] or notice['notice_dt']}
        if cancelled:
            # 취소 공고에는 일정이 비어 오기도 한다. 원래 마감일까지 [진행 중]에 보이도록 예전 일정을 남긴다.
            notice = {**notice, **{name: notice[name] or existing[name] for name in ('bid_close', 'open_dt')}}
        summary = _diff_summary(existing, notice)
        # 차수가 바뀌면 앞선 공고의 기초금액과 낙찰 결과는 새 차수의 값이 아니다.
        # 비워야 아래 후속 조회가 정정된 공고를 다시 확인한다.
        assignments = ', '.join(f'{name}=?' for name in NOTICE_COLUMNS)
        reset_fields = (
            'base_amount', 'base_opened_at', 'base_checked_at',
            'winner_name', 'winner_amount', 'winner_rate', 'result_state',
            'result_participants', 'result_opened_at', 'result_final_date',
            'result_updated_by', 'result_updated_at', 'result_checked_at', 'bidders',
        )
        assignments += ', ' + ', '.join(f'{name}=NULL' for name in reset_fields)
        conn.execute(
            f'UPDATE bid_notices SET {assignments}, updated_at=?, changed_at=?, change_summary=? WHERE id=?',
            [notice[name] for name in NOTICE_COLUMNS] + [now_text, now_text, summary or None, existing['id']],
        )
        event_type = 'cancelled' if cancelled else 'changed'
        if _add_event(conn, existing['id'], notice, event_type, summary, now_text, mailed_at):
            counts[event_type] += 1
    return counts


def _refresh_base_amounts(service_key, include_closed=False):
    """마감 전(또는 마감한 지 얼마 안 된) 공고 중 기초금액이 아직 없는 것을 확인한다.

    ``include_closed``(지난 기간을 정해 수집할 때)면 마감한 공고도 아직 한 번도 확인하지 않은 것은 본다.
    """
    now = _now()
    cutoff = _stamp(now - timedelta(days=BASE_AMOUNT_AFTER_CLOSE_DAYS))
    recent = "(COALESCE(bid_close, '') = '' OR bid_close >= ?)"
    if include_closed:
        recent = f'({recent} OR base_checked_at IS NULL)'
    conn = get_db()
    try:
        targets = conn.execute(
            f'''SELECT id, bid_no FROM bid_notices
                WHERE base_amount IS NULL AND COALESCE(kind, '') <> ? AND {recent}
                ORDER BY COALESCE(base_checked_at, '') ASC, id DESC LIMIT ?''',
            (g2b.KIND_CANCEL, cutoff, BASE_AMOUNT_LIMIT),
        ).fetchall()
    finally:
        conn.close()

    found, note = [], ''
    for target in targets:
        try:
            found.append((target['id'], g2b.fetch_base_amount(service_key, target['bid_no'])))
        except g2b.G2BError as error:
            note = f'기초금액 확인은 일부만 했습니다({error})'
            break
    if found:
        conn = get_db()
        try:
            for notice_id, amount in found:
                conn.execute(
                    'UPDATE bid_notices SET base_amount=?, base_opened_at=?, base_checked_at=? WHERE id=?',
                    (amount['amount'] if amount else None, amount['opened_at'] if amount else None,
                     _stamp(now), notice_id),
                )
            conn.commit()
        finally:
            conn.close()
    return sum(1 for _, amount in found if amount), note


def _finish_run(run_id, status, message='', **values):
    conn = get_db()
    try:
        fields = {'status': status, 'message': message[:500], 'finished_at': _stamp()}
        fields.update(values)
        assignments = ', '.join(f'{name}=?' for name in fields)
        conn.execute(f'UPDATE bid_collect_runs SET {assignments} WHERE id=?', [*fields.values(), run_id])
        conn.commit()
    finally:
        conn.close()


def _purge_old_runs(conn):
    """``RUN_KEEP_DAYS``일이 지난 수집 기록을 지운다. 마지막 성공 기록은 오래돼도 남긴다(맨 처음 수집인지 가리는 데 쓴다)."""
    cutoff = _stamp(_now() - timedelta(days=RUN_KEEP_DAYS))
    conn.execute(
        'DELETE FROM bid_collect_runs WHERE started_at < ? AND id NOT IN ('
        "SELECT id FROM bid_collect_runs WHERE status='success' AND trigger_type <> 'range' ORDER BY id DESC LIMIT 1)",
        (cutoff,),
    )


def _start_run(conn, trigger, actor=''):
    _purge_old_runs(conn)
    cursor = conn.execute(
        'INSERT INTO bid_collect_runs (trigger_type, actor, started_at, status) VALUES (?, ?, ?, ?)',
        (trigger, actor, _stamp(), 'running'),
    )
    conn.commit()
    return cursor.lastrowid


def _search_term_pages(settings, term, begin, end):
    """한 기간의 결과가 페이지 한도를 넘으면 기간을 반으로 나눠 빠짐없이 조회한다."""
    emitted = False
    try:
        for page in g2b.iter_search_notices(settings['service_key'], term, begin, end):
            emitted = True
            yield page
    except g2b.SearchPageLimit:
        if emitted or end - begin <= timedelta(minutes=1):
            raise
        middle = begin + (end - begin) / 2
        yield from _search_term_pages(settings, term, begin, middle)
        yield from _search_term_pages(settings, term, middle, end)


def _search_all(settings, begin, end):
    """기간을 나눠 한 페이지씩 전달해 전체 응답을 메모리에 쌓지 않는다."""
    chunk_begin = begin
    while chunk_begin < end:
        chunk_end = min(end, chunk_begin + timedelta(days=g2b.MAX_RANGE_DAYS))
        for term in g2b.search_terms(settings['keywords'] + settings['org_free_keywords']):
            yield from _search_term_pages(settings, term, chunk_begin, chunk_end)
        chunk_begin = chunk_end


def collect_notices(trigger='manual', actor='', date_range=None, restore_deleted=False):
    """나라장터에서 공고를 찾아 저장한다. 결과 요약(dict)을 돌려준다.

    ``date_range``가 없으면 항상 최근 4주(``COLLECT_DAYS``일)를 찾는다. 이미 저장된 공고는 다시 기록하지 않고
    정정·취소가 생긴 것만 이력으로 더한다. ``date_range``=(시작, 끝)을 주면 그 게시 기간만 찾는다. 지난 공고를
    채워 넣는 용도라 메일 대상에서 뺀다. ``restore_deleted``면 직접 삭제한 공고도 삭제 기록과 상관없이 다시 가져온다.

    {'status': 'success'|'error'|'busy', 'message', 'new', 'changed', 'cancelled', 'restored', 'fetched', 'run_id'}
    """
    if not _collect_lock.acquire(blocking=False):
        return {'status': 'busy', 'message': '지금 다른 수집이 진행 중입니다. 잠시 후 다시 눌러 주세요.'}
    run_id = None
    try:
        conn = get_db()
        try:
            settings = get_settings(conn)
            run_id = _start_run(conn, trigger, actor)
            has_earlier_run = conn.execute(
                "SELECT 1 FROM bid_collect_runs WHERE status='success' AND trigger_type <> 'range' LIMIT 1"
            ).fetchone() is not None
        finally:
            conn.close()

        if not settings['service_key']:
            message = '나라장터 인증키가 설정되지 않았습니다. [공고수집·메일 설정]에서 인증키를 저장해 주세요.'
            _finish_run(run_id, 'error', message)
            return {'status': 'error', 'message': message, 'run_id': run_id}

        if date_range:
            begin, end = date_range
            baseline = True
        else:
            end = _now()
            begin = end - timedelta(days=COLLECT_DAYS)
            baseline = not has_earlier_run

        counts = {'new': 0, 'changed': 0, 'cancelled': 0, 'restored': 0}
        fetched = 0
        try:
            for items in _search_all(settings, begin, end):
                conn = get_db()
                try:
                    batch_counts = store_items(conn, items, settings, baseline=baseline,
                                               restore_deleted=restore_deleted)
                    conn.commit()
                finally:
                    conn.close()
                fetched += len(items)
                for name in counts:
                    counts[name] += batch_counts[name]
        except g2b.G2BError as error:
            _finish_run(run_id, 'error', str(error), range_begin=_stamp(begin), range_end=_stamp(end))
            return {'status': 'error', 'message': str(error), 'run_id': run_id}

        base_found, note = _refresh_base_amounts(settings['service_key'], include_closed=bool(date_range))
        message = (f"신규 {counts['new']}건 · 정정 {counts['changed']}건 · 취소 {counts['cancelled']}건"
                   f" (기초금액 확인 {base_found}건)")
        if restore_deleted:
            message += f" · 삭제했던 공고 {counts['restored']}건을 다시 가져왔습니다(신규에 포함, 메일 대상 아님)"
        if date_range:
            message = f'{begin:%Y-%m-%d} ~ {end:%Y-%m-%d} · {message} · 기간을 정해 찾은 공고는 메일로 보내지 않습니다'
        elif baseline:
            message += ' · 처음 수집이라 메일 대상에서는 뺐습니다'
        if note:
            message += f' · {note}'
        _finish_run(
            run_id, 'success', message, range_begin=_stamp(begin), range_end=_stamp(end),
            fetched=fetched, new_count=counts['new'], changed_count=counts['changed'],
            cancelled_count=counts['cancelled'],
        )
        return {'status': 'success', 'message': message, 'fetched': fetched, 'run_id': run_id, **counts}
    except Exception as error:  # 예기치 못한 오류도 수집 기록에 남겨 화면에서 보이게 한다.
        if run_id is not None:
            _finish_run(run_id, 'error', f'수집 중 오류가 났습니다: {error}')
        raise
    finally:
        _collect_lock.release()


# ---------------------------------------------------------------------------
# 첨부파일 (누를 때 나라장터에서 받아 압축파일 하나로 바로 내려 준다)
# ---------------------------------------------------------------------------
def _member_name(name, used):
    """압축파일 안에 넣을 이름(폴더 기호 제거, 같은 이름은 (2), (3)…을 붙인다)."""
    clean = re.sub(r'[\\/:*?"<>|\x00-\x1f]', '_', str(name or '')).strip(' .') or '첨부파일'
    stem, ext = os.path.splitext(clean[:150])
    candidate, number = f'{stem}{ext}', 2
    while candidate.lower() in used:
        candidate = f'{stem} ({number}){ext}'
        number += 1
    used.add(candidate.lower())
    return candidate


def _attachment_list(raw):
    try:
        files = json.loads(raw or '')
    except (TypeError, ValueError):
        return []
    if not isinstance(files, list):
        return []
    return [item for item in files if isinstance(item, dict) and item.get('url')]


def _notice_files(row, settings):
    """공고의 첨부 목록 [{'name', 'url'}]. 예전에 쌓은 공고(목록이 NULL)는 공고번호로 나라장터에서 다시 읽는다."""
    files = _attachment_list(row['attachments'])
    if row['attachments'] is None:
        if not settings['service_key']:
            raise ValueError('나라장터 인증키가 없어 첨부파일 목록을 읽을 수 없습니다.')
        item = g2b.fetch_notice(settings['service_key'], row['bid_no'])
        files = g2b.attachments_of(item) if item else []
        conn = get_db()
        try:
            conn.execute('UPDATE bid_notices SET attachments=? WHERE id=?',
                         (json.dumps(files, ensure_ascii=False), row['id']))
            conn.commit()
        finally:
            conn.close()
    return files


def build_attachment_zip(notice_id):
    """공고의 첨부파일을 모두 받아 압축파일(열려 있는 임시 파일)로 묶는다. 서버에 남기지 않는다.

    돌려주는 값: {'file', 'name', 'count', 'failed': [못 받은 파일 이름]}.
    공고가 없으면 LookupError, 받을 파일이 없거나 실패하면 ValueError / G2BError(안내 문구).
    """
    conn = get_db()
    try:
        row = conn.execute('SELECT * FROM bid_notices WHERE id=?', (notice_id,)).fetchone()
        settings = get_settings(conn)
    finally:
        conn.close()
    if not row:
        raise LookupError('공고를 찾을 수 없습니다.')

    files = _notice_files(row, settings)
    if not files:
        raise ValueError('이 공고에는 나라장터에 올라온 첨부파일이 없습니다.')

    buffer = tempfile.SpooledTemporaryFile(max_size=ZIP_MEMORY_LIMIT)
    saved, total, failed, used = 0, 0, [], set()
    try:
        with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as archive:
            for item in files:
                label = str(item.get('name') or '첨부파일')
                try:
                    data = g2b.download_attachment(item['url'], MAX_ATTACH_FILE_BYTES)
                except g2b.G2BError:
                    failed.append(label)
                    continue
                if total + len(data) > MAX_ATTACH_TOTAL_BYTES:
                    failed.append(label)
                    continue
                total += len(data)
                archive.writestr(_member_name(label, used), data)
                saved += 1
        if not saved:
            raise g2b.G2BError('첨부파일을 받지 못했습니다. 잠시 후 다시 시도하거나 공고명을 눌러 나라장터에서 받아 주세요.')
    except Exception:
        buffer.close()
        raise
    buffer.seek(0)
    label = re.sub(r'[\\/:*?"<>|\s]+', '_', row['school'] or '').strip('_')
    closing = str(row['bid_close'] or '')[:10]            # 파일명: 학교명_마감날짜 (예: 푸른초등학교_2026-04-24)
    name = '_'.join(part for part in (label, closing) if part) or f"{row['bid_no']}_첨부파일"
    name += '.zip'
    return {'file': buffer, 'name': name, 'count': saved, 'failed': failed}


def _clean_part(text):
    """압축파일 안 폴더 이름으로 쓸 수 있게 경로 기호를 없앤다."""
    return re.sub(r'[\\/:*?"<>|\x00-\x1f]', '_', str(text or '')).strip(' .')[:80]


def build_bundle_zip(notice_ids):
    """체크한 공고들의 첨부파일을 모두 받아 압축파일 하나로 묶는다(공고마다 폴더). 서버에 남기지 않는다.

    돌려주는 값: {'file', 'name', 'files': 받은 파일 수, 'failed': [못 받은 파일], 'notices': 파일이 든 공고 수,
    'skipped': 첨부가 없거나 목록을 못 읽은 공고 수}.
    고른 공고를 하나도 찾지 못하면 LookupError, 받을 파일이 없거나 모두 실패하면 ValueError / G2BError.
    """
    conn = get_db()
    try:
        marks = ', '.join('?' * len(notice_ids))
        found = {row['id']: row for row in conn.execute(
            f'SELECT * FROM bid_notices WHERE id IN ({marks})', list(notice_ids)).fetchall()}
        settings = get_settings(conn)
    finally:
        conn.close()
    rows = [found[notice_id] for notice_id in notice_ids if notice_id in found]
    if not rows:
        raise LookupError('선택한 공고를 찾을 수 없습니다. 목록을 새로 고친 뒤 다시 골라 주세요.')

    buffer = tempfile.SpooledTemporaryFile(max_size=ZIP_MEMORY_LIMIT)
    saved, total, failed, skipped, notices = 0, 0, [], 0, 0
    try:
        with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as archive:
            for row in rows:
                try:
                    files = _notice_files(row, settings)
                except (ValueError, g2b.G2BError):
                    files = []
                if not files:
                    skipped += 1
                    continue
                folder = _clean_part(f"{row['school'] or ''}_{row['bid_no']}".strip('_')) or row['bid_no']
                used, got = set(), 0
                notice_total = 0
                for item in files:
                    label = str(item.get('name') or '첨부파일')
                    try:
                        data = g2b.download_attachment(item['url'], MAX_ATTACH_FILE_BYTES)
                    except g2b.G2BError:
                        failed.append(f'{folder}/{label}')
                        continue
                    if notice_total + len(data) > MAX_ATTACH_TOTAL_BYTES or total + len(data) > MAX_BUNDLE_TOTAL_BYTES:
                        failed.append(f'{folder}/{label}')
                        continue
                    notice_total += len(data)
                    total += len(data)
                    archive.writestr(f'{folder}/{_member_name(label, used)}', data)
                    got += 1
                saved += got
                notices += 1 if got else 0
                skipped += 0 if got else 1
        if not saved:
            if not failed:
                raise ValueError('선택한 공고에는 나라장터에 올라온 첨부파일이 없습니다.')
            raise g2b.G2BError('첨부파일을 받지 못했습니다. 잠시 후 다시 시도하거나 공고명을 눌러 나라장터에서 받아 주세요.')
    except Exception:
        buffer.close()
        raise
    buffer.seek(0)
    name = f"입찰공고_첨부파일_{_now():%Y%m%d_%H%M}.zip"
    return {'file': buffer, 'name': name, 'files': saved, 'failed': failed, 'notices': notices, 'skipped': skipped}


# ---------------------------------------------------------------------------
# 낙찰 결과 (나라장터 낙찰정보서비스) — [낙찰 결과 가져오기]를 누를 때만 가져온다
# ---------------------------------------------------------------------------
# 확인 대상: 취소되지 않았고 개찰 예정 시각이 지났으며, 아직 '최종낙찰'이 아니거나 참가 업체를 읽지 않은 공고
RESULT_ELIGIBLE_SQL = ("COALESCE(kind, '') <> ? AND COALESCE(open_dt, '') <> '' AND open_dt <= ? "
                       "AND (COALESCE(result_state, '') <> ? OR bidders IS NULL)")


def _save_award(conn, notice_id, award, actor, now_text):
    """나라장터에서 찾은 낙찰 결과를 공고에 저장한다(결과를 찾아본 시각도 함께 적는다)."""
    conn.execute(
        """UPDATE bid_notices SET winner_name=?, winner_amount=?, winner_rate=?, result_state=?,
                  result_participants=?, result_opened_at=?, result_final_date=?,
                  result_updated_by=?, result_updated_at=?, result_checked_at=? WHERE id=?""",
        (award['winner_name'] or None, award['winner_amount'], award['winner_rate'] or None,
         award['state'], award['participants'], award['opened_at'] or None,
         award['final_date'] or None, actor, now_text, now_text, notice_id),
    )


def refresh_results(service_key, since, limit=None, actor=''):
    """개찰일시가 지난 공고의 낙찰 결과를 ``limit``건 가져와 저장한다.

    결과가 아직 없는 공고부터, 오래 확인하지 않은 순으로 본다. ``since``(첫 요청 시각) 이후에 이미 확인한 공고는
    건너뛰므로 화면이 같은 ``since``로 이어 부르면 같은 공고를 두 번 묻지 않고 끝까지 받는다.
    돌려주는 값: {'checked', 'found', 'note', 'remaining'(아직 확인하지 않은 공고 수)}.
    """
    limit = limit or RESULT_LIMIT
    now_text = _stamp()
    base_args = [g2b.KIND_CANCEL, now_text, g2b.STATE_AWARDED, since]
    pending_sql = f"FROM bid_notices WHERE {RESULT_ELIGIBLE_SQL} AND COALESCE(result_checked_at, '') < ?"
    conn = get_db()
    try:
        targets = conn.execute(
            f"SELECT id, bid_no, result_state {pending_sql} ORDER BY CASE WHEN COALESCE(result_state, '') = '' THEN 0 ELSE 1 END, "
            "COALESCE(result_checked_at, '') ASC, open_dt DESC LIMIT ?", [*base_args, limit]).fetchall()
    finally:
        conn.close()

    checked, note = [], ''
    for target in targets:
        final = target['result_state'] == g2b.STATE_AWARDED      # 이미 최종 낙찰: 참가 업체만 읽는다
        try:
            award = None if final else g2b.fetch_award(service_key, target['bid_no'])
            bidders = g2b.fetch_bidders(service_key, target['bid_no']) if final or award else None
        except g2b.G2BError as error:
            note = f'낙찰 결과 확인은 일부만 했습니다({error})'
            break
        checked.append((target['id'], award, bidders))
    conn = get_db()
    try:
        for notice_id, award, bidders in checked:
            if award:
                _save_award(conn, notice_id, award, actor, now_text)
            else:
                # 아직 결과가 안 올라온 공고는 예전 값을 그대로 두고 찾아본 시각만 남긴다.
                conn.execute('UPDATE bid_notices SET result_checked_at=? WHERE id=?', (now_text, notice_id))
            if bidders is not None:
                conn.execute('UPDATE bid_notices SET bidders=? WHERE id=?',
                             (json.dumps(bidders, ensure_ascii=False), notice_id))
        conn.commit()
        remaining = conn.execute(f'SELECT COUNT(*) {pending_sql}', base_args).fetchone()[0]
    finally:
        conn.close()
    return {'checked': len(checked), 'found': sum(1 for _, award, _bidders in checked if award), 'note': note,
            'remaining': remaining}


def pending_result_count():
    """개찰은 지났는데 아직 결과가 없는 공고 수(취소 공고 제외)."""
    conn = get_db()
    try:
        return conn.execute(
            """SELECT COUNT(*) FROM bid_notices
                WHERE COALESCE(kind, '') <> ? AND COALESCE(open_dt, '') <> '' AND open_dt <= ?
                  AND COALESCE(result_state, '') = '' AND COALESCE(winner_name, '') = ''""",
            (g2b.KIND_CANCEL, _stamp())).fetchone()[0]
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# 메일
# ---------------------------------------------------------------------------
def _amount_text(row):
    if row['base_amount']:
        return f"{_won(row['base_amount'])} (기초금액)"
    if row['budget']:
        return f"{_won(row['budget'])} (배정예산)"
    return '-'


def _mail_rows(conn, test=False):
    """메일에 실을 공고. 공고 한 건은 한 줄만(가장 마지막 소식 기준)."""
    if test:
        rows = conn.execute(
            "SELECT n.*, 'new' AS event_type, '' AS summary, 0 AS event_id FROM bid_notices n ORDER BY n.id DESC LIMIT 5"
        ).fetchall()
        return rows, []
    rows = conn.execute(
        '''SELECT n.*, e.event_type, e.summary, e.id AS event_id
             FROM bid_notice_events e JOIN bid_notices n ON n.id = e.notice_id
            WHERE e.mailed_at IS NULL ORDER BY e.id'''
    ).fetchall()
    latest = {}
    for row in rows:
        latest[row['id']] = row
    return list(latest.values()), [row['event_id'] for row in rows]


def build_mail(rows, test=False):
    """(제목, HTML 본문, 글자 본문)"""
    groups = {name: [row for row in rows if row['event_type'] == name] for name in EVENT_LABELS}
    counts = ' · '.join(f'{EVENT_LABELS[name]} {len(items)}건' for name, items in groups.items() if items)
    today = _now().strftime('%m월 %d일')
    subject = f"[새담 입찰공고] {today} {counts or '새 소식'}"
    if test:
        subject = f'[테스트] {subject}'
    colors = {'new': '#004ea2', 'changed': '#b45309', 'cancelled': '#b91c1c'}
    cell = 'padding:9px 10px;border:1px solid #dbe3ef;font-size:13px;vertical-align:top;'
    head = 'padding:8px 10px;border:1px solid #dbe3ef;background:#f8fafc;font-size:12px;text-align:left;white-space:nowrap;'
    html_parts, text_parts = [], []
    for name, items in groups.items():
        if not items:
            continue
        label = EVENT_LABELS[name]
        html_parts.append(
            f'<h3 style="margin:22px 0 8px;color:{colors[name]};font-size:16px;">{label} 공고 {len(items)}건</h3>'
            '<table cellpadding="0" cellspacing="0" style="border-collapse:collapse;width:100%;">'
            f'<tr><th style="{head}">공고명</th><th style="{head}">지역</th><th style="{head}">학교</th>'
            f'<th style="{head}">입찰 마감</th><th style="{head}">금액</th></tr>'
        )
        text_parts.append(f'■ {label} 공고 {len(items)}건')
        for row in items:
            detail = f'<div style="margin-top:4px;color:#b45309;font-size:12px;">{escape(row["summary"])}</div>' \
                if row['summary'] else ''
            title = escape(row['title'])
            if row['url']:
                title = f'<a href="{escape(row["url"])}" style="color:#1d4ed8;">{title}</a>'
            html_parts.append(
                f'<tr><td style="{cell}">{title}'
                f'<div style="margin-top:3px;color:#64748b;font-size:12px;">{escape(row["bid_no"])}-{escape(row["bid_ord"])}</div>'
                f'{detail}</td>'
                f'<td style="{cell}white-space:nowrap;">{escape(row["region"] or "-")}</td>'
                f'<td style="{cell}white-space:nowrap;">{escape(_school_label(row) or "-")}</td>'
                f'<td style="{cell}white-space:nowrap;">{escape(_short_dt(row["bid_close"]) or "-")}</td>'
                f'<td style="{cell}white-space:nowrap;text-align:right;">{escape(_amount_text(row))}</td></tr>'
            )
            text_parts.append(
                f"- {row['title']} | {row['region'] or '-'} | {row['school'] or '-'} | "
                f"마감 {_short_dt(row['bid_close']) or '-'} | {_amount_text(row)}"
                + (f" | {row['summary']}" if row['summary'] else '')
                + (f"\n  {row['url']}" if row['url'] else '')
            )
        html_parts.append('</table>')
    note = '설정 확인용 테스트 메일입니다. 최근에 수집한 공고를 예시로 보여 드립니다.' if test \
        else '나라장터에서 새로 확인한 입찰공고 소식입니다.'
    html = (
        '<div style="font-family:\'Malgun Gothic\',\'Apple SD Gothic Neo\',sans-serif;color:#0f172a;line-height:1.6;max-width:900px;">'
        f'<h2 style="margin:0 0 6px;color:#004ea2;">새담 입찰공고 알림</h2>'
        f'<p style="margin:0;color:#475569;font-size:13px;">{note}</p>'
        + ''.join(html_parts) +
        f'<p style="margin:22px 0 0;font-size:13px;">전체 목록 보기: <a href="{INTRANET_URL}">{INTRANET_URL}</a></p>'
        '<p style="margin:6px 0 0;color:#94a3b8;font-size:12px;">공고명을 누르면 나라장터 공고 화면이 열립니다. '
        '금액은 기초금액이 공개되기 전에는 배정예산으로 보여 드립니다.</p></div>'
    )
    text = '\n'.join([note, '', *text_parts, '', f'전체 목록: {INTRANET_URL}'])
    return subject, html, text


def _send_mail(recipients, subject, html, text):
    """인트라넷 공용 메일 계정(Gmail)으로 보낸다. 실패하면 예외."""
    sender, password = _mail_credentials()
    if not sender or not password:
        raise RuntimeError('보내는 메일 계정 설정이 없습니다(MAIL_USERNAME / MAIL_PASSWORD).')
    message = MIMEMultipart('alternative')
    message['From'] = formataddr((str(Header('새담 인트라넷', 'utf-8')), sender))
    message['To'] = ', '.join(recipients)
    message['Subject'] = str(Header(subject, 'utf-8'))
    message.attach(MIMEText(text, 'plain', 'utf-8'))
    message.attach(MIMEText(html, 'html', 'utf-8'))
    with smtplib.SMTP('smtp.gmail.com', 587, timeout=30) as server:
        server.starttls()
        server.login(sender, password)
        server.sendmail(sender, recipients, message.as_string())


def send_update_mail(test=False):
    """아직 메일로 알리지 않은 소식을 등록된 주소로 보낸다.

    {'status': 'sent'|'none'|'no_recipient'|'error', 'message', 'count'}
    """
    conn = get_db()
    try:
        settings = get_settings(conn)
        rows, event_ids = _mail_rows(conn, test=test)
    finally:
        conn.close()
    recipients = settings['recipients']
    if not recipients:
        return {'status': 'no_recipient', 'message': '받는 메일 주소가 등록되지 않았습니다.', 'count': 0}
    if not rows:
        message = '예시로 보낼 공고가 아직 없습니다. 먼저 [공고가져오기]를 눌러 주세요.' if test \
            else '메일로 보낼 새 소식이 없습니다.'
        return {'status': 'none', 'message': message, 'count': 0}
    subject, html, text = build_mail(rows, test=test)
    try:
        _send_mail(recipients, subject, html, text)
    except Exception as error:
        return {'status': 'error', 'message': f'메일을 보내지 못했습니다: {error}', 'count': 0}
    if event_ids:
        conn = get_db()
        try:
            marks = ', '.join('?' * len(event_ids))
            conn.execute(f'UPDATE bid_notice_events SET mailed_at=? WHERE id IN ({marks})', [_stamp(), *event_ids])
            conn.commit()
        finally:
            conn.close()
    return {'status': 'sent', 'count': len(rows),
            'message': f'{len(recipients)}곳에 공고 {len(rows)}건을 보냈습니다.'}


# ---------------------------------------------------------------------------
# 화면 · API
# ---------------------------------------------------------------------------
def _notice_status(row, now_text):
    if row['kind'] == g2b.KIND_CANCEL:
        return 'cancelled'
    if row['bid_close'] and row['bid_close'] < now_text:
        return 'closed'
    return 'open'


def _notice_dict(row, now):
    now_text = _stamp(now)
    status = _notice_status(row, now_text)
    close = g2b.parse_datetime(row['bid_close'])
    return {
        'id': row['id'],
        'bid_no': row['bid_no'],
        'bid_ord': row['bid_ord'],
        'kind': row['kind'] or '',
        'status': status,
        'changed': row['bid_ord'] != '000' and status != 'cancelled',
        # 반짝이는 N: [공고가져오기]로 새로 기록된 지 NEW_MARK_HOURS시간 안인 공고.
        # 기간을 정해 가져온 4주보다 오래된 공고(지난 공고 채워 넣기)에는 붙이지 않는다.
        'is_new': bool(row['first_seen_at']) and row['first_seen_at'] >= _stamp(now - timedelta(hours=NEW_MARK_HOURS))
                  and (not row['notice_dt'] or row['notice_dt'] >= _stamp(now - timedelta(days=COLLECT_DAYS))),
        'title': row['title'],
        'region': row['region'] or '',
        'school': row['school'] or '',
        'demand_name': row['demand_name'] or '',
        'school_label': _school_label(row),           # [학교] 칸(교육지원청은 앞의 시·도 교육청을 뺀 이름)
        'notice_dt': _short_dt(row['notice_dt']),
        'bid_close': _short_dt(row['bid_close']),
        'days_left': (close.date() - now.date()).days if close and status == 'open' else None,
        'open_dt': _short_dt(row['open_dt']),
        'base_amount': row['base_amount'],
        'budget': row['budget'],
        'est_price': row['est_price'],
        'contract_method': row['contract_method'] or '',
        'official': row['official'] or '',
        'tel': row['tel'] or '',
        'url': row['url'] or '',
        'matched_keyword': row['matched_keyword'] or '',
        'event_count': row['event_count'] if 'event_count' in row.keys() else 0,
        'changed_at': _short_dt(row['changed_at']),      # 정정·취소를 확인한 때
        'work_status': row['work_status'] or '',
        'archive_id': row['archive_id'],
        'result': {
            'state': row['result_state'] or ('직접입력' if row['winner_name'] else ''),
            'winner_name': row['winner_name'] or '',
            'winner_amount': row['winner_amount'],
            'winner_rate': row['winner_rate'] or '',
            'participants': row['result_participants'],
            'opened_at': _short_dt(row['result_opened_at']),
            'final_date': row['result_final_date'] or '',
            'checked_at': _short_dt(row['result_updated_at']),
        },
        # 첨부가 없다고 확인된 공고(목록이 빈 값)는 받기 버튼을 숨긴다.
        'no_files': row['attachments'] is not None and not _attachment_list(row['attachments']),
    }


def _run_dict(row):
    return {
        'trigger': {'auto': '자동', 'range': '기간 수집'}.get(row['trigger_type'], '직접'),
        'actor': row['actor'] or '',
        'started_at': _short_dt(row['started_at']),
        'status': row['status'],
        'message': row['message'] or '',
        'mail_status': row['mail_status'] or '',
    }


def _parse_day(value):
    """'2026-12-01' → datetime(그날 0시). 비었거나 잘못된 값이면 None."""
    try:
        return datetime.strptime(str(value or '').strip(), '%Y-%m-%d')
    except ValueError:
        return None


@bid_notice_bp.route('/')
def index():
    return render_template('bid_notice/bid_notice.html', is_admin=is_admin_session(),
                           can_settings=is_admin_session(max_level=SETTINGS_MAX_LEVEL))


HAS_RESULT_SQL = "(COALESCE(n.result_state, '') <> '' OR COALESCE(n.winner_name, '') <> '')"
NO_WINNER_SQL = "COALESCE(n.winner_name, '') = ''"


def _to_int(value):
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _list_conditions(get, now, with_region_result=True):
    """목록 조건(보관함·상태 탭·검색어·게시일·진행 상태·지역·결과) → (where 목록, params 목록, status).

    ``get``은 ``request.args.get``처럼 이름으로 값을 돌려주는 함수다. ``with_region_result=False``이면
    지역·결과 조건을 뺀다(머리글 선택 상자에 보여 줄 지역·업체 목록을 만들 때 쓴다).
    """
    now_text = _stamp(now)
    status = str(get('status') or 'open')
    where, params = [], []
    archive_id = _to_int(get('archive'))
    if archive_id:
        # 보관함 안의 공고는 상태 탭과 상관없이 모두 보여 주고, [공고 목록]에는 나오지 않는다.
        status = 'all'
        where.append('n.archive_id = ?')
        params.append(archive_id)
    else:
        where.append('n.archive_id IS NULL')
    if status == 'cancelled':
        where.append('n.kind = ?')
        params.append(g2b.KIND_CANCEL)
    elif status == 'closed':
        where.append("COALESCE(n.kind, '') <> ? AND n.bid_close <> '' AND n.bid_close < ?")
        params.extend([g2b.KIND_CANCEL, now_text])
    elif status in ('soon', 'tomorrow', 'week'):
        # 마감 임박: 3일내(오늘부터 D-3까지) / 내일 마감(마감일이 내일) / 1주내(오늘부터 D-7까지). 화면의 D-day와 같은 날짜 기준
        today = now.date()
        if status == 'tomorrow':
            start, end = f'{today + timedelta(days=1):%Y-%m-%d}', f'{today + timedelta(days=2):%Y-%m-%d}'
        else:
            start, end = now_text, f'{today + timedelta(days=(SOON_DAYS if status == "soon" else WEEK_DAYS) + 1):%Y-%m-%d}'
        where.append("COALESCE(n.kind, '') <> ? AND n.bid_close >= ? AND n.bid_close < ?")
        params.extend([g2b.KIND_CANCEL, start, end])
    elif status != 'all':
        # 진행 중: 취소된 공고도 원래 마감일까지(마감일이 없으면 취소 확인 후 CANCEL_KEEP_DAYS일) 취소 표시와 함께 남긴다.
        where.append("((COALESCE(n.kind, '') <> ? AND (COALESCE(n.bid_close, '') = '' OR n.bid_close >= ?))"
                     " OR (n.kind = ? AND (n.bid_close >= ? OR (COALESCE(n.bid_close, '') = '' AND n.changed_at >= ?))))")
        params.extend([g2b.KIND_CANCEL, now_text, g2b.KIND_CANCEL, now_text,
                       _stamp(now - timedelta(days=CANCEL_KEEP_DAYS))])
    keyword = str(get('q') or '').strip()
    if keyword:
        like = f'%{keyword}%'
        where.append('(n.title LIKE ? OR n.school LIKE ? OR n.region LIKE ? OR n.demand_name LIKE ? '
                     'OR n.bid_no LIKE ? OR n.winner_name LIKE ?)')
        params.extend([like] * 6)
    # 게시일 기간
    date_from = _parse_day(get('date_from'))
    date_to = _parse_day(get('date_to'))
    if date_from:
        where.append('n.notice_dt >= ?')
        params.append(f'{date_from:%Y-%m-%d}')
    if date_to:
        where.append('n.notice_dt < ?')
        params.append(f'{date_to + timedelta(days=1):%Y-%m-%d}')
    work_status = str(get('work_status') or '').strip()
    if work_status == 'none':
        where.append("COALESCE(n.work_status, '') = ''")
    elif work_status in WORK_STATUSES:
        where.append('n.work_status = ?')
        params.append(work_status)
    if with_region_result:
        region = str(get('region') or '').strip()
        if region:
            where.append('n.region = ?')
            params.append(region)
        result = str(get('result') or '').strip()
        if result == '__has':
            where.append(HAS_RESULT_SQL)
        elif result == '__none':
            where.append(f'NOT {HAS_RESULT_SQL}')
        elif result == '__etc':
            where.append(f'{HAS_RESULT_SQL} AND {NO_WINNER_SQL}')
        elif result.startswith('w:') and result[2:]:
            where.append('n.winner_name = ?')
            params.append(result[2:])
    return where, params, status


def _order_sql(status, sort, direction):
    """정렬: [입찰 마감]·[결과] 머리글(빈 값은 항상 맨 뒤), 기본은 진행 중·3일내 마감은 마감 가까운 순, 그 밖은 최근 공고 순.

    정정·취소된 공고도 따로 끌어올리지 않는다(같은 줄이 제자리에서 정정·취소 표시로 바뀐다).
    """
    descending = 'DESC' if str(direction or '').lower() == 'desc' else 'ASC'
    tail = 'n.notice_dt DESC, n.id DESC'
    if sort == 'close':
        return f"CASE WHEN COALESCE(n.bid_close, '') = '' THEN 1 ELSE 0 END, n.bid_close {descending}, {tail}"
    if sort == 'posted':
        return f"CASE WHEN COALESCE(n.notice_dt, '') = '' THEN 1 ELSE 0 END, n.notice_dt {descending}, n.id {descending}"
    if sort == 'winner':
        return f"CASE WHEN COALESCE(n.winner_name, '') = '' THEN 1 ELSE 0 END, n.winner_name {descending}, {tail}"
    if status in ('open', 'soon', 'tomorrow', 'week'):
        return "CASE WHEN COALESCE(n.bid_close, '') = '' THEN 1 ELSE 0 END, n.bid_close ASC, n.id DESC"
    return tail


@bid_notice_bp.route('/api/notices')
def api_notices():
    """공고 목록 한 쪽(서버가 조건·정렬·쪽 나누기를 모두 처리한다)."""
    now = _now()
    now_text = _stamp(now)
    get = request.args.get
    where, params, status = _list_conditions(get, now)
    show_all = str(get('per_page') or '').strip().lower() in ('all', '0')       # 모두 보기: 쪽을 나누지 않는다
    per_page = _to_int(get('per_page'))
    per_page = per_page if per_page in PAGE_SIZES else PAGE_SIZES[0]
    order = _order_sql(status, get('sort'), get('dir'))
    where_sql = ' WHERE ' + ' AND '.join(where)
    base_where, base_params, _ = _list_conditions(get, now, with_region_result=False)
    base_sql = ' WHERE ' + ' AND '.join(base_where)

    conn = get_db()
    try:
        total = conn.execute(f'SELECT COUNT(*) FROM bid_notices n{where_sql}', params).fetchone()[0]
        pages = 1 if show_all else max(1, -(-total // per_page))
        page = 1 if show_all else min(max(1, _to_int(get('page')) or 1), pages)
        limit_sql, limit_args = ('', []) if show_all else (' LIMIT ? OFFSET ?', [per_page, (page - 1) * per_page])
        rows = conn.execute(
            f"""SELECT n.*, (SELECT COUNT(*) FROM bid_notice_events e WHERE e.notice_id = n.id) AS event_count
                  FROM bid_notices n{where_sql} ORDER BY {order}{limit_sql}""",
            [*params, *limit_args]).fetchall()
        # 머리글 선택 상자: 지역·결과 조건만 뺀 같은 범위의 지역·낙찰 업체·결과 건수
        regions = [{'name': row['name'], 'count': row['cnt']} for row in conn.execute(
            f"SELECT n.region AS name, COUNT(*) AS cnt FROM bid_notices n{base_sql} "
            "AND COALESCE(n.region, '') <> '' GROUP BY n.region ORDER BY n.region", base_params)]
        winners = [{'name': row['name'], 'count': row['cnt']} for row in conn.execute(
            f"SELECT n.winner_name AS name, COUNT(*) AS cnt FROM bid_notices n{base_sql} "
            "AND COALESCE(n.winner_name, '') <> '' GROUP BY n.winner_name ORDER BY n.winner_name", base_params)]
        counts = conn.execute(
            f"""SELECT COUNT(*) AS total,
                       SUM(CASE WHEN {HAS_RESULT_SQL} THEN 1 ELSE 0 END) AS has_result,
                       SUM(CASE WHEN {HAS_RESULT_SQL} AND {NO_WINNER_SQL} THEN 1 ELSE 0 END) AS etc
                  FROM bid_notices n{base_sql}""", base_params).fetchone()
        has_result = counts['has_result'] or 0
        result_counts = {'has': has_result, 'none': counts['total'] - has_result, 'etc': counts['etc'] or 0}
        today = now.strftime('%Y-%m-%d')
        week_ago = _stamp(now - timedelta(days=7))
        summary = {
            'open': conn.execute(
                "SELECT COUNT(*) FROM bid_notices WHERE archive_id IS NULL AND COALESCE(kind, '') <> ? "
                "AND (COALESCE(bid_close, '') = '' OR bid_close >= ?)",
                (g2b.KIND_CANCEL, now_text)).fetchone()[0],
            'total': conn.execute('SELECT COUNT(*) FROM bid_notices WHERE archive_id IS NULL').fetchone()[0],
            'new_today': conn.execute(
                "SELECT COUNT(*) FROM bid_notice_events WHERE event_type='new' AND created_at >= ?",
                (today,)).fetchone()[0],
            'changed_week': conn.execute(
                "SELECT COUNT(*) FROM bid_notice_events WHERE event_type IN ('changed', 'cancelled') AND created_at >= ?",
                (week_ago,)).fetchone()[0],
            'unsent': conn.execute(
                'SELECT COUNT(DISTINCT notice_id) FROM bid_notice_events WHERE mailed_at IS NULL').fetchone()[0],
        }
        runs = conn.execute('SELECT * FROM bid_collect_runs ORDER BY id DESC LIMIT 10').fetchall()
        settings = get_settings(conn)
        archives = _archive_list(conn)
    finally:
        conn.close()
    return jsonify({
        'status': 'success',
        'notices': [_notice_dict(row, now) for row in rows],
        'total': total, 'page': page, 'pages': pages, 'per_page': 0 if show_all else per_page,
        'regions': regions, 'winners': winners, 'result_counts': result_counts,
        'archives': archives,
        'summary': summary,
        'runs': [_run_dict(row) for row in runs],
        'has_key': bool(settings['service_key']),
        'work_statuses': list(WORK_STATUSES),
    })


@bid_notice_bp.route('/api/notices/ids')
def api_notice_ids():
    """지금 조건에 맞는 공고 전체의 번호(체크박스 '조건에 맞는 전체 선택'용). 한 번에 ``LIST_LIMIT``건까지."""
    now = _now()
    get = request.args.get
    where, params, status = _list_conditions(get, now)
    order = _order_sql(status, get('sort'), get('dir'))
    where_sql = ' WHERE ' + ' AND '.join(where)
    conn = get_db()
    try:
        total = conn.execute(f'SELECT COUNT(*) FROM bid_notices n{where_sql}', params).fetchone()[0]
        rows = conn.execute(f'SELECT n.id FROM bid_notices n{where_sql} ORDER BY {order} LIMIT ?',
                            [*params, LIST_LIMIT]).fetchall()
    finally:
        conn.close()
    return jsonify({'status': 'success', 'ids': [row['id'] for row in rows], 'total': total, 'capped': total > LIST_LIMIT})


@bid_notice_bp.route('/api/notices/<int:notice_id>/events')
def api_events(notice_id):
    conn = get_db()
    try:
        rows = conn.execute(
            'SELECT * FROM bid_notice_events WHERE notice_id=? ORDER BY id', (notice_id,)).fetchall()
    finally:
        conn.close()
    return jsonify({'status': 'success', 'events': [{
        'label': EVENT_LABELS.get(row['event_type'], row['event_type']),
        'type': row['event_type'],
        'bid_ord': row['bid_ord'],
        'summary': row['summary'] or '',
        'reason': row['reason'] or '',
        'created_at': _short_dt(row['created_at']),
        'mailed': bool(row['mailed_at']),
    } for row in rows]})


def _restore_deleted_flag(payload):
    """[삭제한 공고도 다시 가져오기] 체크 여부. 삭제와 같이 관리자만 쓸 수 있다(아니면 PermissionError)."""
    if not payload.get('restore_deleted'):
        return False
    if not is_admin_session():
        raise PermissionError('삭제한 공고를 다시 가져오는 것은 관리자만 할 수 있습니다.')
    return True


@bid_notice_bp.route('/api/collect', methods=['POST'])
def api_collect():
    try:
        restore_deleted = _restore_deleted_flag(request.get_json(silent=True) or {})
    except PermissionError as error:
        return _json_error(str(error), 403)
    result = collect_notices(trigger='manual', actor=session.get('user_name') or '',
                             restore_deleted=restore_deleted)
    if result['status'] == 'busy':
        return _json_error(result['message'], 409)
    if result['status'] == 'error':
        return _json_error(result['message'], 502)
    return jsonify(result)


@bid_notice_bp.route('/api/collect-range', methods=['POST'])
def api_collect_range():
    """게시일 기간을 정해 나라장터에서 찾는다(지난 공고 채워 넣기). 메일은 보내지 않는다."""
    payload = request.get_json(silent=True) or {}
    try:
        restore_deleted = _restore_deleted_flag(payload)
    except PermissionError as error:
        return _json_error(str(error), 403)
    begin = _parse_day(payload.get('date_from'))
    last_day = _parse_day(payload.get('date_to'))
    if not begin or not last_day:
        return _json_error('시작일과 종료일을 모두 골라 주세요.')
    if begin > last_day:
        return _json_error('시작일이 종료일보다 늦습니다.')
    now = _now()
    if begin > now:
        return _json_error('시작일이 오늘보다 뒤입니다.')
    end = min(last_day + timedelta(days=1) - timedelta(minutes=1), now)
    if end - begin > timedelta(days=MAX_RANGE_COLLECT_DAYS):
        return _json_error('한 번에 1년까지만 찾을 수 있습니다. 기간을 나눠서 찾아 주세요.')
    result = collect_notices(trigger='range', actor=session.get('user_name') or '', date_range=(begin, end),
                             restore_deleted=restore_deleted)
    if result['status'] == 'busy':
        return _json_error(result['message'], 409)
    if result['status'] == 'error':
        return _json_error(result['message'], 502)
    return jsonify(result)


@bid_notice_bp.route('/api/notices/<int:notice_id>/bidders')
def api_bidders(notice_id):
    """낙찰 업체를 누르면 펼쳐지는 상세에 보여 줄 참가 업체 전체(저장된 값만, 나라장터에 묻지 않는다)."""
    conn = get_db()
    try:
        row = conn.execute('SELECT bidders FROM bid_notices WHERE id=?', (notice_id,)).fetchone()
    finally:
        conn.close()
    if not row:
        return _json_error('공고를 찾을 수 없습니다.', 404)
    try:
        bidders = json.loads(row['bidders']) if row['bidders'] else None
    except (TypeError, ValueError):
        bidders = None
    return jsonify({'status': 'success', 'loaded': isinstance(bidders, list),
                    'bidders': bidders if isinstance(bidders, list) else []})


@bid_notice_bp.route('/api/notices/<int:notice_id>/status', methods=['POST'])
def api_work_status(notice_id):
    payload = request.get_json(silent=True) or {}
    work_status = str(payload.get('work_status') or '').strip()
    if work_status and work_status not in WORK_STATUSES:
        return _json_error('상태를 목록에서 골라 주세요.')
    conn = get_db()
    try:
        cursor = conn.execute(
            'UPDATE bid_notices SET work_status=?, status_updated_by=?, status_updated_at=? WHERE id=?',
            (work_status or None, session.get('user_name') or '', _stamp(), notice_id),
        )
        conn.commit()
        if not cursor.rowcount:
            return _json_error('공고를 찾을 수 없습니다.', 404)
    finally:
        conn.close()
    return jsonify({'status': 'success', 'work_status': work_status})


@bid_notice_bp.route('/api/notices/<int:notice_id>/files')
def api_files_download(notice_id):
    """[첨부 받기]: 첨부파일을 지금 나라장터에서 받아 압축파일 하나로 내려 준다(서버에 남기지 않는다)."""
    if not _attach_lock.acquire(blocking=False):
        return _json_error('다른 공고의 첨부파일을 받는 중입니다. 잠시 후 다시 눌러 주세요.', 409)
    try:
        bundle = build_attachment_zip(notice_id)
    except LookupError as error:
        return _json_error(str(error), 404)
    except ValueError as error:
        return _json_error(str(error))
    except g2b.G2BError as error:
        return _json_error(str(error), 502)
    finally:
        _attach_lock.release()
    response = send_file(bundle['file'], mimetype='application/zip', as_attachment=True,
                         download_name=bundle['name'])
    response.headers['X-Bid-File-Count'] = str(bundle['count'])
    response.headers['X-Bid-File-Failed'] = str(len(bundle['failed']))
    response.headers['Cache-Control'] = 'no-store'
    return response


@bid_notice_bp.route('/api/files/bundle', methods=['POST'])
def api_files_bundle():
    """체크한 공고들의 첨부파일을 한꺼번에 받아 압축파일 하나(공고마다 폴더)로 내려 준다(서버에 남기지 않는다)."""
    payload = request.get_json(silent=True) or {}
    try:
        ids = _parse_ids(payload, limit=MAX_BUNDLE_NOTICES,
                         too_many=f'한 번에 {MAX_BUNDLE_NOTICES}건까지 받을 수 있습니다. 나눠서 받아 주세요.')
    except ValueError as error:
        return _json_error(str(error))
    if not _attach_lock.acquire(blocking=False):
        return _json_error('다른 공고의 첨부파일을 받는 중입니다. 잠시 후 다시 눌러 주세요.', 409)
    try:
        bundle = build_bundle_zip(ids)
    except LookupError as error:
        return _json_error(str(error), 404)
    except ValueError as error:
        return _json_error(str(error))
    except g2b.G2BError as error:
        return _json_error(str(error), 502)
    finally:
        _attach_lock.release()
    response = send_file(bundle['file'], mimetype='application/zip', as_attachment=True,
                         download_name=bundle['name'])
    response.headers['X-Bid-File-Count'] = str(bundle['files'])
    response.headers['X-Bid-File-Failed'] = str(len(bundle['failed']))
    response.headers['X-Bid-Notice-Count'] = str(bundle['notices'])
    response.headers['X-Bid-Notice-Skipped'] = str(bundle['skipped'])
    response.headers['Cache-Control'] = 'no-store'
    return response


# ---------------------------------------------------------------------------
# 보관함(탭) · 선택 공고 삭제 · 엑셀 내보내기
# ---------------------------------------------------------------------------
def _parse_ids(payload, limit=LIST_LIMIT, too_many=''):
    """JSON의 ids를 중복 없는 정수 목록(고른 순서 유지)으로 만든다. 잘못되면 ValueError(안내 문구)."""
    raw = payload.get('ids')
    if not isinstance(raw, list):
        raise ValueError('공고를 선택해 주세요.')
    try:
        ids = list(dict.fromkeys(int(value) for value in raw))
    except (TypeError, ValueError):
        raise ValueError('선택한 공고 번호가 올바르지 않습니다.') from None
    if not ids:
        raise ValueError('공고를 선택해 주세요.')
    if len(ids) > limit:
        raise ValueError(too_many or f'한 번에 {limit}건까지 처리할 수 있습니다.')
    return ids


def _archive_list(conn):
    rows = conn.execute(
        """SELECT a.id, a.name, COUNT(n.id) AS cnt
             FROM bid_archives a LEFT JOIN bid_notices n ON n.archive_id = a.id
            GROUP BY a.id, a.name ORDER BY a.id"""
    ).fetchall()
    return [{'id': row['id'], 'name': row['name'], 'count': row['cnt']} for row in rows]


def _clean_archive_name(value):
    name = re.sub(r'\s+', ' ', str(value or '')).strip()
    if not name:
        raise ValueError('보관함 이름을 입력해 주세요.')
    if len(name) > MAX_ARCHIVE_NAME:
        raise ValueError(f'보관함 이름은 {MAX_ARCHIVE_NAME}자 이내로 입력해 주세요.')
    return name


def _archive_by_name(conn, name):
    return conn.execute('SELECT id, name FROM bid_archives WHERE name = ? COLLATE NOCASE', (name,)).fetchone()


def _create_archive(conn, name, actor):
    if conn.execute('SELECT COUNT(*) FROM bid_archives').fetchone()[0] >= MAX_ARCHIVES:
        raise ValueError(f'보관함은 {MAX_ARCHIVES}개까지 만들 수 있습니다. 안 쓰는 보관함을 삭제해 주세요.')
    cursor = conn.execute('INSERT INTO bid_archives (name, created_by, created_at) VALUES (?, ?, ?)',
                          (name, actor, _stamp()))
    return cursor.lastrowid


@bid_notice_bp.route('/api/archives', methods=['POST'])
def api_archive_create():
    """보관함(탭)을 만든다. 예) '2025년 공고'"""
    payload = request.get_json(silent=True) or {}
    try:
        name = _clean_archive_name(payload.get('name'))
    except ValueError as error:
        return _json_error(str(error))
    conn = get_db()
    try:
        if _archive_by_name(conn, name):
            return _json_error('같은 이름의 보관함이 이미 있습니다.', 409)
        try:
            archive_id = _create_archive(conn, name, session.get('user_name') or '')
        except ValueError as error:
            return _json_error(str(error))
        conn.commit()
    finally:
        conn.close()
    return jsonify({'status': 'success', 'archive': {'id': archive_id, 'name': name, 'count': 0}})


@bid_notice_bp.route('/api/archives/<int:archive_id>/rename', methods=['POST'])
def api_archive_rename(archive_id):
    payload = request.get_json(silent=True) or {}
    try:
        name = _clean_archive_name(payload.get('name'))
    except ValueError as error:
        return _json_error(str(error))
    conn = get_db()
    try:
        if not conn.execute('SELECT 1 FROM bid_archives WHERE id=?', (archive_id,)).fetchone():
            return _json_error('보관함을 찾을 수 없습니다.', 404)
        other = _archive_by_name(conn, name)
        if other and other['id'] != archive_id:
            return _json_error('같은 이름의 보관함이 이미 있습니다.', 409)
        conn.execute('UPDATE bid_archives SET name=? WHERE id=?', (name, archive_id))
        conn.commit()
    finally:
        conn.close()
    return jsonify({'status': 'success', 'name': name})


@bid_notice_bp.route('/api/archives/<int:archive_id>/delete', methods=['POST'])
def api_archive_delete(archive_id):
    """보관함(탭)을 지운다. 안에 있던 공고는 지워지지 않고 [공고 목록]으로 돌아온다."""
    conn = get_db()
    try:
        if not conn.execute('SELECT 1 FROM bid_archives WHERE id=?', (archive_id,)).fetchone():
            return _json_error('보관함을 찾을 수 없습니다.', 404)
        cursor = conn.execute(
            'UPDATE bid_notices SET archive_id=NULL, archived_at=NULL WHERE archive_id=?', (archive_id,))
        conn.execute('DELETE FROM bid_archives WHERE id=?', (archive_id,))
        conn.commit()
        moved_back = cursor.rowcount
    finally:
        conn.close()
    return jsonify({'status': 'success', 'moved_back': moved_back})


@bid_notice_bp.route('/api/notices/archive', methods=['POST'])
def api_notices_archive():
    """선택한 공고를 보관함으로 옮긴다. ``archive_id``(있는 보관함) 또는 ``new_name``(새로 만들거나 같은 이름 사용)."""
    payload = request.get_json(silent=True) or {}
    try:
        ids = _parse_ids(payload)
    except ValueError as error:
        return _json_error(str(error))
    conn = get_db()
    try:
        if str(payload.get('new_name') or '').strip():
            try:
                name = _clean_archive_name(payload.get('new_name'))
                found = _archive_by_name(conn, name)
                archive_id = found['id'] if found else _create_archive(conn, name, session.get('user_name') or '')
                name = found['name'] if found else name
            except ValueError as error:
                return _json_error(str(error))
        else:
            try:
                archive_id = int(payload.get('archive_id'))
            except (TypeError, ValueError):
                return _json_error('옮길 보관함을 골라 주세요.')
            row = conn.execute('SELECT id, name FROM bid_archives WHERE id=?', (archive_id,)).fetchone()
            if not row:
                return _json_error('보관함을 찾을 수 없습니다.', 404)
            name = row['name']
        marks = ', '.join('?' * len(ids))
        cursor = conn.execute(
            f'UPDATE bid_notices SET archive_id=?, archived_at=? WHERE id IN ({marks})',
            [archive_id, _stamp(), *ids])
        conn.commit()
        moved = cursor.rowcount
    finally:
        conn.close()
    return jsonify({'status': 'success', 'moved': moved, 'archive': {'id': archive_id, 'name': name}})


@bid_notice_bp.route('/api/notices/unarchive', methods=['POST'])
def api_notices_unarchive():
    """선택한 공고를 보관함에서 꺼내 [공고 목록]으로 되돌린다."""
    payload = request.get_json(silent=True) or {}
    try:
        ids = _parse_ids(payload)
    except ValueError as error:
        return _json_error(str(error))
    conn = get_db()
    try:
        marks = ', '.join('?' * len(ids))
        cursor = conn.execute(
            f'UPDATE bid_notices SET archive_id=NULL, archived_at=NULL WHERE id IN ({marks})', ids)
        conn.commit()
        moved = cursor.rowcount
    finally:
        conn.close()
    return jsonify({'status': 'success', 'moved': moved})


@bid_notice_bp.route('/api/notices/delete', methods=['POST'])
@admin_required
def api_notices_delete():
    """선택한 공고를 완전히 지운다(관리자 전용, 되돌릴 수 없음). [공고가져오기]가 다시 가져오지 않도록 공고번호를 기억한다."""
    payload = request.get_json(silent=True) or {}
    if not payload.get('confirm'):
        return _json_error('삭제를 확인해 주세요.')
    try:
        ids = _parse_ids(payload)
    except ValueError as error:
        return _json_error(str(error))
    conn = get_db()
    try:
        marks = ', '.join('?' * len(ids))
        rows = conn.execute(f'SELECT id, bid_no FROM bid_notices WHERE id IN ({marks})', ids).fetchall()
        if not rows:
            return _json_error('삭제할 공고를 찾을 수 없습니다. 목록을 새로 고친 뒤 다시 골라 주세요.', 404)
        found = [row['id'] for row in rows]
        found_marks = ', '.join('?' * len(found))
        conn.execute(f'DELETE FROM bid_notice_events WHERE notice_id IN ({found_marks})', found)
        conn.execute(f'DELETE FROM bid_notices WHERE id IN ({found_marks})', found)
        actor, now_text = session.get('user_name') or '', _stamp()
        conn.executemany('INSERT OR IGNORE INTO bid_notice_deleted (bid_no, deleted_at, deleted_by) VALUES (?, ?, ?)',
                         [(row['bid_no'], now_text, actor) for row in rows])
        conn.commit()
    finally:
        conn.close()
    return jsonify({'status': 'success', 'deleted': len(rows)})


EXPORT_COLUMNS = (
    # (머리글, 너비, 숫자 서식)
    ('번호', 7, None), ('공고 구분', 10, None), ('상태', 8, None),
    ('공고명', 62, None), ('지역', 14, None), ('학교', 20, None), ('수요기관', 36, None),
    ('게시일시', 17, None), ('입찰 마감', 17, None), ('개찰일시', 17, None),
    ('기초금액', 14, '#,##0'), ('금액 구분', 10, None),
    ('낙찰 업체', 28, None), ('낙찰금액', 14, '#,##0'), ('낙찰률(%)', 10, '0.000'), ('참가 업체 수', 11, '0'),
    ('결과', 16, None), ('공고 링크', 12, None),
)
RESULT_STATE_LABELS = {'최종낙찰': '최종 낙찰', '개찰완료': '개찰 완료(최종 낙찰 전)'}
NOTICE_STATUS_LABELS = {'open': '진행 중', 'closed': '마감', 'cancelled': '취소'}


def _rate_number(value):
    """낙찰률('80.606')을 엑셀에서 계산할 수 있게 숫자로. 숫자가 아니면 글자 그대로."""
    text = str(value or '').strip()
    try:
        return float(text)
    except ValueError:
        return text


def build_export_workbook(rows):
    """공고 목록(보이는 순서)을 엑셀 파일(bytes)로 만든다."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    now_text = _stamp()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '입찰공고'
    header_fill = PatternFill('solid', fgColor='1F3A5F')
    for column, (title, width, _fmt) in enumerate(EXPORT_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=column, value=title)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        sheet.column_dimensions[get_column_letter(column)].width = width

    for index, row in enumerate(rows):
        state = row['result_state'] or ('직접 입력' if row['winner_name'] else '')
        # 화면과 같이 기초금액이 아직 없으면 배정예산을 넣고, 어느 금액인지 옆 칸에 적는다.
        if row['base_amount']:
            amount, amount_kind = row['base_amount'], '기초금액'
        elif row['budget']:
            amount, amount_kind = row['budget'], '배정예산'
        else:
            amount, amount_kind = None, ''
        values = (
            len(rows) - index, row['kind'] or '',
            NOTICE_STATUS_LABELS.get(_notice_status(row, now_text), ''),
            row['title'], row['region'] or '', _school_label(row), row['demand_name'] or '',
            _short_dt(row['notice_dt']), _short_dt(row['bid_close']), _short_dt(row['open_dt']),
            amount, amount_kind,
            row['winner_name'] or '', row['winner_amount'], _rate_number(row['winner_rate']), row['result_participants'],
            RESULT_STATE_LABELS.get(state, state), '',
        )
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=index + 2, column=column, value=value)
            if isinstance(value, str):
                cell.data_type = 's'          # '='로 시작하는 공고명도 수식이 아니라 글자로 저장한다
            number_format = EXPORT_COLUMNS[column - 1][2]
            if number_format and value not in (None, ''):
                cell.number_format = number_format
        url = str(row['url'] or '')
        if url.startswith('https://'):
            link = sheet.cell(row=index + 2, column=len(EXPORT_COLUMNS), value='공고 보기')
            link.hyperlink = url
            link.font = Font(color='1D4ED8', underline='single')
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@bid_notice_bp.route('/api/export', methods=['POST'])
def api_export():
    """공고 목록을 엑셀로 내려 준다. ``ids``(체크한 공고)나 ``query``(지금 목록의 조건)가 필요하다.

    ``query``만 있으면 조건에 맞는 공고 전체(모든 쪽, 최대 ``EXPORT_LIMIT``건)를 목록과 같은 순서로,
    ``ids``만 있으면 그 순서대로, 둘 다 있으면 ``ids``의 공고를 ``query``의 정렬로 내려 준다.
    """
    payload = request.get_json(silent=True) or {}
    query = payload.get('query') if isinstance(payload.get('query'), dict) else None
    now = _now()
    capped = False
    conn = get_db()
    try:
        if payload.get('ids') is not None or query is None:
            try:
                ids = _parse_ids(payload)
            except ValueError as error:
                return _json_error(str(error))
            marks = ', '.join('?' * len(ids))
            if query is not None:
                _, _, status = _list_conditions(query.get, now)
                order = _order_sql(status, query.get('sort'), query.get('dir'))
                rows = conn.execute(f'SELECT n.* FROM bid_notices n WHERE n.id IN ({marks}) ORDER BY {order}', ids).fetchall()
            else:
                found = {row['id']: row for row in conn.execute(
                    f'SELECT * FROM bid_notices WHERE id IN ({marks})', ids).fetchall()}
                rows = [found[notice_id] for notice_id in ids if notice_id in found]
        else:
            where, params, status = _list_conditions(query.get, now)
            order = _order_sql(status, query.get('sort'), query.get('dir'))
            rows = conn.execute(
                f"SELECT n.* FROM bid_notices n WHERE {' AND '.join(where)} ORDER BY {order} LIMIT ?",
                [*params, EXPORT_LIMIT + 1]).fetchall()
            if len(rows) > EXPORT_LIMIT:
                rows, capped = rows[:EXPORT_LIMIT], True
    finally:
        conn.close()
    if not rows:
        return _json_error('내려받을 공고가 없습니다. 목록을 새로 고친 뒤 다시 눌러 주세요.', 404)
    data = build_export_workbook(rows)
    name = f'입찰공고_목록_{_now():%Y%m%d_%H%M}.xlsx'
    response = send_file(BytesIO(data), as_attachment=True, download_name=name,
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    if capped:
        response.headers['X-Bid-Export-Capped'] = '1'
    return response


@bid_notice_bp.route('/api/results/refresh', methods=['POST'])
def api_results_refresh():
    """[낙찰 결과 가져오기]: 개찰일시가 지난 공고의 낙찰 결과를 ``RESULT_LIMIT``건씩 가져온다.

    ``remaining``(아직 안 본 공고 수)과 ``since``를 돌려주고, 화면은 ``remaining``이 0이 될 때까지
    ``since``를 그대로 실어 이어서 부른다(한 요청이 오래 걸리지 않고, 같은 공고를 두 번 묻지 않는다).
    """
    payload = request.get_json(silent=True) or {}
    since = str(payload.get('since') or '').strip()
    if since and not g2b.parse_datetime(since):
        return _json_error('이어받기 정보가 올바르지 않습니다. 다시 눌러 주세요.')
    since = since or _stamp()
    settings = get_settings()
    if not settings['service_key']:
        return _json_error('나라장터 인증키가 설정되지 않았습니다.')
    if not _collect_lock.acquire(blocking=False):
        return _json_error('지금 공고를 가져오는 중입니다. 끝난 뒤 다시 눌러 주세요.', 409)
    try:
        outcome = refresh_results(settings['service_key'], since, actor=session.get('user_name') or '')
    finally:
        _collect_lock.release()
    pending = pending_result_count()
    if outcome['checked']:
        message = f"{outcome['checked']}건을 확인해 {outcome['found']}건의 결과를 가져왔습니다."
    else:
        message = '확인할 공고가 없습니다(개찰이 지난 공고의 결과를 모두 확인했습니다).'
    if pending:
        message += f' 아직 결과가 없는 공고가 {pending}건 남아 있습니다(나라장터에 아직 올라오지 않은 공고 포함).'
    if outcome['note']:
        message += f" ({outcome['note']})"
    return jsonify({'status': 'success', 'message': message, 'since': since, 'checked': outcome['checked'],
                    'found': outcome['found'], 'pending': pending, 'remaining': outcome['remaining'],
                    'note': outcome['note']})


@bid_notice_bp.route('/api/settings')
@admin_required(max_level=SETTINGS_MAX_LEVEL)
def api_settings():
    return jsonify({'status': 'success', 'settings': public_settings(get_settings())})


@bid_notice_bp.route('/api/settings', methods=['POST'])
@admin_required(max_level=SETTINGS_MAX_LEVEL)
def api_settings_save():
    payload = request.get_json(silent=True) or {}
    try:
        settings = save_settings(payload, actor=session.get('user_name') or '')
    except ValueError as error:
        return _json_error(str(error))
    return jsonify({'status': 'success', 'settings': public_settings(settings)})


@bid_notice_bp.route('/api/send-mail', methods=['POST'])
@admin_required(max_level=SETTINGS_MAX_LEVEL)
def api_send_mail():
    """미발송 소식(또는 테스트 메일)을 지금 보낸다. 화면의 확인 창에서 동의했을 때만 부른다."""
    payload = request.get_json(silent=True) or {}
    if not payload.get('confirm'):
        return _json_error('메일 발송을 확인해 주세요.')
    result = send_update_mail(test=bool(payload.get('test')))
    if result['status'] != 'sent':
        return _json_error(result['message'], 502 if result['status'] == 'error' else 400)
    return jsonify(result)
