"""조직관리 > 강사출결시스템.

방과후강사가 학교에 비치된 '오늘의 QR'을 자기 휴대폰으로 찍으면 출석이 기록된다.

- 강사는 학교별로 미리 등록한다(이름·휴대폰번호).
- QR은 학교×날짜마다 새 토큰으로 발급되어 어제 QR 사진으로는 출석할 수 없다.
- 휴대폰 인식: 강사가 처음 QR을 찍을 때 이름·휴대폰번호를 확인하고 그 기기에
  무작위 기기키(쿠키)를 심는다. 이후에는 입력 없이 찍기만 하면 출석된다.
  한 번호에는 기기 하나만 연결되며, 기기를 바꾸면 관리자가 '기기 초기화'를 한다.
- 출결은 강사×날짜 한 건으로 저장해 날짜별·월별로 조회한다.
"""

import hashlib
import io
import os
import re
import secrets
import socket
import time
from calendar import monthrange
from datetime import date, datetime, timedelta, timezone
from functools import lru_cache
from urllib.parse import urlsplit, urlunsplit

import segno
from flask import (
    Blueprint, Response, jsonify, render_template, request, send_file, session, url_for,
)

from .database import get_db
from .security import menu_permission_required

instructor_attendance_bp = Blueprint('instructor_attendance', __name__)

MENU_KEY = 'instructor_attendance'
KST = timezone(timedelta(hours=9))
DEVICE_COOKIE = 'saedam_lecturer_device'
DEVICE_COOKIE_MAX_AGE = 60 * 60 * 24 * 365 * 5
WEEKDAY_LABELS = ('월', '화', '수', '목', '금', '토', '일')

REGISTER_WINDOW_SECONDS = 10 * 60
REGISTER_MAX_ATTEMPTS = 10
_register_attempts = {}


# ---------------------------------------------------------------- 스키마

def ensure_instructor_attendance_schema(conn):
    conn.executescript('''
        CREATE TABLE IF NOT EXISTS lecturer_instructors (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            school_id INTEGER NOT NULL,
            name TEXT NOT NULL,
            phone TEXT NOT NULL,
            subject TEXT,
            memo TEXT,
            is_active INTEGER NOT NULL DEFAULT 1,
            created_by TEXT,
            created_at TEXT DEFAULT (datetime('now', 'localtime'))
        );
        CREATE INDEX IF NOT EXISTS idx_lecturer_instructors_school
            ON lecturer_instructors(school_id, is_active);
        CREATE INDEX IF NOT EXISTS idx_lecturer_instructors_phone
            ON lecturer_instructors(phone);

        -- 휴대폰번호 하나에 기기 하나. 쿠키 원문은 저장하지 않고 해시만 둔다.
        CREATE TABLE IF NOT EXISTS lecturer_devices (
            phone TEXT PRIMARY KEY,
            token_hash TEXT NOT NULL UNIQUE,
            user_agent TEXT,
            bound_at TEXT,
            last_seen_at TEXT
        );

        -- 학교×날짜별 QR 토큰. 재발급하면 같은 날의 토큰을 교체한다.
        CREATE TABLE IF NOT EXISTS lecturer_qr_codes (
            school_id INTEGER NOT NULL,
            qr_date TEXT NOT NULL,
            token TEXT NOT NULL,
            issued_by TEXT,
            issued_at TEXT,
            PRIMARY KEY (school_id, qr_date)
        );

        -- 학교에 둔 태블릿이 매일 새 QR을 스스로 띄우기 위한 표시용 링크 키.
        CREATE TABLE IF NOT EXISTS lecturer_qr_displays (
            school_id INTEGER PRIMARY KEY,
            display_key TEXT NOT NULL UNIQUE,
            created_at TEXT
        );

        CREATE TABLE IF NOT EXISTS lecturer_attendance (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            instructor_id INTEGER NOT NULL,
            school_id INTEGER NOT NULL,
            att_date TEXT NOT NULL,
            check_in_at TEXT NOT NULL,
            method TEXT NOT NULL DEFAULT 'QR',
            note TEXT,
            recorded_by TEXT,
            UNIQUE (instructor_id, att_date)
        );
        CREATE INDEX IF NOT EXISTS idx_lecturer_attendance_date
            ON lecturer_attendance(att_date, school_id);
    ''')
    conn.commit()


def init_instructor_attendance_schema():
    conn = get_db()
    try:
        ensure_instructor_attendance_schema(conn)
    finally:
        conn.close()


# ---------------------------------------------------------------- 공통

def _now():
    return datetime.now(KST)


def _today():
    return _now().date().isoformat()


def _now_text():
    return _now().strftime('%Y-%m-%d %H:%M:%S')


def _json_error(message, status=400, **extra):
    payload = {'status': 'error', 'message': message}
    payload.update(extra)
    return jsonify(payload), status


def normalize_phone(value):
    digits = re.sub(r'\D', '', str(value or ''))
    if digits.startswith('82') and len(digits) >= 11:
        digits = '0' + digits[2:]
    return digits


def format_phone(digits):
    digits = normalize_phone(digits)
    if len(digits) == 11:
        return f'{digits[:3]}-{digits[3:7]}-{digits[7:]}'
    if len(digits) == 10:
        return f'{digits[:3]}-{digits[3:6]}-{digits[6:]}'
    return digits


def _valid_phone(digits):
    return bool(re.fullmatch(r'01\d{8,9}', digits))


def _parse_date(value, default=None):
    try:
        return date.fromisoformat(str(value or '').strip()).isoformat()
    except ValueError:
        return default


def _hash_token(token):
    return hashlib.sha256(str(token).encode('utf-8')).hexdigest()


def _int_or_none(value):
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _school_label(row):
    year = str(row['year'] or '').strip()
    return f"{row['school_name']} ({year})" if year else row['school_name']


def _list_schools(conn):
    columns = {row[1] for row in conn.execute('PRAGMA table_info(schools)').fetchall()}
    where = 'WHERE COALESCE(is_active, 1) = 1' if 'is_active' in columns else ''
    rows = conn.execute(f'''
        SELECT id, school_name, year FROM schools {where}
        ORDER BY year DESC, school_name COLLATE NOCASE, id
    ''').fetchall()
    return [{'id': row['id'], 'name': row['school_name'], 'label': _school_label(row)} for row in rows]


def _get_school(conn, school_id):
    return conn.execute(
        'SELECT id, school_name, year FROM schools WHERE id=?', (school_id,)
    ).fetchone()


LOOPBACK_HOSTS = {'localhost', '127.0.0.1', '0.0.0.0', '::1'}


@lru_cache(maxsize=1)
def _lan_ip():
    """같은 와이파이의 휴대폰이 접속할 수 있는 이 PC의 주소."""
    try:
        probe = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        try:
            # UDP라 실제로 패킷을 보내지는 않고 사용할 랜카드 주소만 확인한다.
            probe.connect(('8.8.8.8', 53))
            return probe.getsockname()[0]
        finally:
            probe.close()
    except OSError:
        return ''


def _swap_loopback_host(root):
    """관리자가 127.0.0.1로 접속해 발급한 QR은 강사 휴대폰에서 열리지 않는다.
    개발 환경에서라도 찍어볼 수 있도록 같은 와이파이에서 통하는 주소로 바꾼다."""
    parts = urlsplit(root)
    if (parts.hostname or '') not in LOOPBACK_HOSTS:
        return root
    ip = _lan_ip()
    if not ip:
        return root
    netloc = f'{ip}:{parts.port}' if parts.port else ip
    return urlunsplit((parts.scheme, netloc, parts.path, '', ''))


def _external_root():
    """QR 안에 넣을 주소. 강사 휴대폰이 바깥에서 열 수 있는 주소여야 한다."""
    # 운영 주소를 환경변수로 지정했으면 어떤 경로로 접속했든 그 주소를 쓴다.
    configured = str(os.getenv('PUBLIC_BASE_URL') or '').strip()
    if configured:
        return configured.rstrip('/')
    root = request.url_root
    proto = request.headers.get('X-Forwarded-Proto', '').split(',')[0].strip()
    if proto == 'https' and root.startswith('http://'):
        root = 'https://' + root[len('http://'):]
    return _swap_loopback_host(root).rstrip('/')


def _scan_url_warning(scan_url):
    """QR 주소가 강사 휴대폰에서 열리지 않을 것 같으면 담당자에게 알린다."""
    host = urlsplit(scan_url).hostname or ''
    if host in LOOPBACK_HOSTS:
        return (f'이 QR에는 이 PC에서만 열리는 주소({host})가 들어 있어 강사 휴대폰으로 찍으면 '
                '아무 화면도 뜨지 않습니다. PUBLIC_BASE_URL 환경변수에 실제 운영 주소를 설정한 뒤 '
                'QR을 다시 발급해주세요.')
    if re.fullmatch(r'(10|127)\.\d+\.\d+\.\d+|172\.(1[6-9]|2\d|3[01])\.\d+\.\d+|192\.168\.\d+\.\d+', host):
        return ('이 QR은 같은 와이파이에 연결된 휴대폰에서만 열립니다. 강사가 휴대데이터로 찍으면 '
                '열리지 않으니, 실제 운영 주소로 접속해 발급한 QR을 비치해주세요.')
    return ''


def get_daily_token(conn, school_id, qr_date=None, create=True):
    qr_date = qr_date or _today()
    row = conn.execute(
        'SELECT token FROM lecturer_qr_codes WHERE school_id=? AND qr_date=?',
        (school_id, qr_date),
    ).fetchone()
    if row or not create:
        return row['token'] if row else None
    token = secrets.token_urlsafe(18)
    conn.execute('''
        INSERT OR IGNORE INTO lecturer_qr_codes (school_id, qr_date, token, issued_by, issued_at)
        VALUES (?, ?, ?, ?, ?)
    ''', (school_id, qr_date, token, session.get('user_name') or '자동발급', _now_text()))
    conn.commit()
    # 동시에 두 요청이 발급했더라도 먼저 저장된 토큰 하나로 통일한다.
    return conn.execute(
        'SELECT token FROM lecturer_qr_codes WHERE school_id=? AND qr_date=?',
        (school_id, qr_date),
    ).fetchone()['token']


def _get_display_key(conn, school_id, create=True):
    row = conn.execute(
        'SELECT display_key FROM lecturer_qr_displays WHERE school_id=?', (school_id,)
    ).fetchone()
    if row or not create:
        return row['display_key'] if row else None
    key = secrets.token_urlsafe(24)
    conn.execute(
        'INSERT OR IGNORE INTO lecturer_qr_displays (school_id, display_key, created_at) VALUES (?, ?, ?)',
        (school_id, key, _now_text()),
    )
    conn.commit()
    return _get_display_key(conn, school_id, create=False)


def _scan_url(school_id, token):
    return _external_root() + url_for('instructor_attendance.scan_page', school_id=school_id, token=token)


def _qr_svg(data, scale=10):
    buffer = io.BytesIO()
    segno.make(data, error='m').save(buffer, kind='svg', scale=scale, border=2, dark='#111827')
    return buffer.getvalue()


def _svg_response(data):
    response = Response(_qr_svg(data), mimetype='image/svg+xml')
    response.headers['Cache-Control'] = 'no-store'
    return response


def _register_rate_limited(client_key):
    now = time.monotonic()
    attempts = [t for t in _register_attempts.get(client_key, []) if now - t < REGISTER_WINDOW_SECONDS]
    limited = len(attempts) >= REGISTER_MAX_ATTEMPTS
    if not limited:
        attempts.append(now)
    _register_attempts[client_key] = attempts
    return limited


# ---------------------------------------------------------------- 관리자 화면

@instructor_attendance_bp.route('/instructor-attendance')
@menu_permission_required(MENU_KEY)
def admin_page():
    conn = get_db()
    try:
        schools = _list_schools(conn)
    finally:
        conn.close()
    return render_template('instructor_attendance/admin.html', schools=schools, today=_today())


@instructor_attendance_bp.route('/instructor-attendance/api/instructors')
@menu_permission_required(MENU_KEY)
def list_instructors():
    school_id = _int_or_none(request.args.get('school_id'))
    if not school_id:
        return _json_error('학교를 선택해주세요.')
    conn = get_db()
    try:
        rows = conn.execute('''
            SELECT i.id, i.name, i.phone, i.subject, i.memo, i.created_at,
                   d.bound_at, d.last_seen_at
            FROM lecturer_instructors i
            LEFT JOIN lecturer_devices d ON d.phone = i.phone
            WHERE i.school_id=? AND i.is_active=1
            ORDER BY i.name COLLATE NOCASE, i.id
        ''', (school_id,)).fetchall()
    finally:
        conn.close()
    return jsonify({'status': 'success', 'instructors': [{
        'id': row['id'],
        'name': row['name'],
        'phone': format_phone(row['phone']),
        'subject': row['subject'] or '',
        'memo': row['memo'] or '',
        'device_bound_at': row['bound_at'] or '',
        'device_last_seen_at': row['last_seen_at'] or '',
    } for row in rows]})


def _save_instructor(conn, school_id, name, phone, subject, memo, instructor_id=None):
    """등록·수정 공통 검증. 오류 메시지 문자열 또는 None을 돌려준다."""
    if not name:
        return '강사 이름을 입력해주세요.'
    if not _valid_phone(phone):
        return f'{name}: 휴대폰번호(010으로 시작하는 10~11자리)를 확인해주세요.'
    duplicate = conn.execute('''
        SELECT id FROM lecturer_instructors
        WHERE school_id=? AND phone=? AND is_active=1 AND id<>?
    ''', (school_id, phone, instructor_id or 0)).fetchone()
    if duplicate:
        return f'{name}: 이 학교에 같은 휴대폰번호({format_phone(phone)})의 강사가 이미 있습니다.'
    if instructor_id:
        conn.execute('''
            UPDATE lecturer_instructors SET name=?, phone=?, subject=?, memo=?
            WHERE id=? AND school_id=?
        ''', (name, phone, subject, memo, instructor_id, school_id))
    else:
        conn.execute('''
            INSERT INTO lecturer_instructors (school_id, name, phone, subject, memo, created_by)
            VALUES (?, ?, ?, ?, ?, ?)
        ''', (school_id, name, phone, subject, memo, session.get('user_name')))
    return None


@instructor_attendance_bp.route('/instructor-attendance/api/instructors', methods=['POST'])
@menu_permission_required(MENU_KEY)
def save_instructor():
    data = request.get_json(silent=True) or {}
    school_id = _int_or_none(data.get('school_id'))
    conn = get_db()
    try:
        if not school_id or not _get_school(conn, school_id):
            return _json_error('학교를 선택해주세요.')
        error = _save_instructor(
            conn, school_id,
            str(data.get('name') or '').strip()[:40],
            normalize_phone(data.get('phone')),
            str(data.get('subject') or '').strip()[:60],
            str(data.get('memo') or '').strip()[:200],
            _int_or_none(data.get('id')),
        )
        if error:
            conn.rollback()
            return _json_error(error)
        conn.commit()
    finally:
        conn.close()
    return jsonify({'status': 'success', 'message': '강사 정보를 저장했습니다.'})


@instructor_attendance_bp.route('/instructor-attendance/api/instructors/bulk', methods=['POST'])
@menu_permission_required(MENU_KEY)
def bulk_add_instructors():
    """'이름, 휴대폰번호, 과목' 형식의 여러 줄을 한 번에 등록한다."""
    data = request.get_json(silent=True) or {}
    school_id = _int_or_none(data.get('school_id'))
    lines = [line.strip() for line in str(data.get('text') or '').splitlines() if line.strip()]
    if not lines:
        return _json_error('등록할 강사 명단을 입력해주세요.')
    if len(lines) > 200:
        return _json_error('한 번에 200명까지 등록할 수 있습니다.')
    conn = get_db()
    try:
        if not school_id or not _get_school(conn, school_id):
            return _json_error('학교를 선택해주세요.')
        errors = []
        for number, line in enumerate(lines, start=1):
            parts = [part.strip() for part in re.split(r'[,\t]', line)]
            name = parts[0][:40] if parts else ''
            phone = normalize_phone(parts[1] if len(parts) > 1 else '')
            subject = parts[2][:60] if len(parts) > 2 else ''
            error = _save_instructor(conn, school_id, name, phone, subject, '')
            if error:
                errors.append(f'{number}번째 줄 - {error}')
        if errors:
            conn.rollback()
            return _json_error('명단을 확인해주세요. 아무것도 등록하지 않았습니다.', errors=errors)
        conn.commit()
    finally:
        conn.close()
    return jsonify({'status': 'success', 'message': f'강사 {len(lines)}명을 등록했습니다.'})


@instructor_attendance_bp.route('/instructor-attendance/api/instructors/<int:instructor_id>/delete', methods=['POST'])
@menu_permission_required(MENU_KEY)
def delete_instructor(instructor_id):
    # 지난 출결 기록이 남아 있어야 하므로 실제 삭제 대신 명단에서만 뺀다.
    conn = get_db()
    try:
        conn.execute('UPDATE lecturer_instructors SET is_active=0 WHERE id=?', (instructor_id,))
        conn.commit()
    finally:
        conn.close()
    return jsonify({'status': 'success', 'message': '강사를 명단에서 제외했습니다. 지난 출결 기록은 보존됩니다.'})


@instructor_attendance_bp.route('/instructor-attendance/api/instructors/<int:instructor_id>/reset-device', methods=['POST'])
@menu_permission_required(MENU_KEY)
def reset_device(instructor_id):
    conn = get_db()
    try:
        row = conn.execute('SELECT phone FROM lecturer_instructors WHERE id=?', (instructor_id,)).fetchone()
        if not row:
            return _json_error('강사를 찾을 수 없습니다.', 404)
        conn.execute('DELETE FROM lecturer_devices WHERE phone=?', (row['phone'],))
        conn.commit()
    finally:
        conn.close()
    return jsonify({'status': 'success', 'message': '휴대폰 연결을 초기화했습니다. 새 휴대폰으로 QR을 찍으면 다시 연결됩니다.'})


def _daily_rows(conn, att_date, school_id=None):
    params = [att_date]
    school_filter = ''
    if school_id:
        school_filter = 'AND i.school_id=?'
        params.append(school_id)
    params.append(att_date)
    # 명단에서 뺀 강사라도 그날 출석 기록이 있으면 함께 보여준다.
    return conn.execute(f'''
        SELECT i.id, i.name, i.phone, i.subject, i.school_id, i.is_active,
               s.school_name, a.check_in_at, a.method, a.note, a.recorded_by,
               d.phone AS device_phone
        FROM lecturer_instructors i
        JOIN schools s ON s.id = i.school_id
        LEFT JOIN lecturer_attendance a ON a.instructor_id = i.id AND a.att_date = ?
        LEFT JOIN lecturer_devices d ON d.phone = i.phone
        WHERE (i.is_active = 1 OR a.id IS NOT NULL) {school_filter}
          AND date(i.created_at) <= ?
        ORDER BY s.school_name COLLATE NOCASE, (a.check_in_at IS NULL), a.check_in_at, i.name COLLATE NOCASE
    ''', params).fetchall()


@instructor_attendance_bp.route('/instructor-attendance/api/daily')
@menu_permission_required(MENU_KEY)
def daily_status():
    att_date = _parse_date(request.args.get('date'), _today())
    school_id = _int_or_none(request.args.get('school_id'))
    conn = get_db()
    try:
        rows = _daily_rows(conn, att_date, school_id)
    finally:
        conn.close()
    items = [{
        'instructor_id': row['id'],
        'name': row['name'],
        'phone': format_phone(row['phone']),
        'subject': row['subject'] or '',
        'school_id': row['school_id'],
        'school_name': row['school_name'],
        'present': bool(row['check_in_at']),
        'check_in_time': (row['check_in_at'] or '')[11:16],
        'method': row['method'] or '',
        'note': row['note'] or '',
        'recorded_by': row['recorded_by'] or '',
        'device_bound': bool(row['device_phone']),
        'removed': not row['is_active'],
    } for row in rows]
    present = sum(1 for item in items if item['present'])
    weekday = WEEKDAY_LABELS[date.fromisoformat(att_date).weekday()]
    return jsonify({
        'status': 'success',
        'date': att_date,
        'weekday': weekday,
        'summary': {'total': len(items), 'present': present, 'absent': len(items) - present},
        'items': items,
    })


@instructor_attendance_bp.route('/instructor-attendance/api/monthly')
@menu_permission_required(MENU_KEY)
def monthly_status():
    data, error = _monthly_data(request.args.get('month'), _int_or_none(request.args.get('school_id')))
    if error:
        return _json_error(error)
    return jsonify({'status': 'success', **data})


def _monthly_data(month_text, school_id):
    match = re.fullmatch(r'(\d{4})-(\d{2})', str(month_text or '').strip())
    if not match:
        today = _now().date()
        year, month = today.year, today.month
    else:
        year, month = int(match.group(1)), int(match.group(2))
        if not 1 <= month <= 12:
            return None, '조회할 월을 확인해주세요.'
    if not school_id:
        return None, '학교를 선택해주세요.'
    last_day = monthrange(year, month)[1]
    first, last = f'{year:04d}-{month:02d}-01', f'{year:04d}-{month:02d}-{last_day:02d}'
    conn = get_db()
    try:
        school = _get_school(conn, school_id)
        if not school:
            return None, '학교를 찾을 수 없습니다.'
        records = conn.execute('''
            SELECT instructor_id, att_date, check_in_at, method
            FROM lecturer_attendance
            WHERE school_id=? AND att_date BETWEEN ? AND ?
        ''', (school_id, first, last)).fetchall()
        recorded_ids = {row['instructor_id'] for row in records}
        instructors = conn.execute('''
            SELECT id, name, subject, is_active FROM lecturer_instructors
            WHERE school_id=? ORDER BY name COLLATE NOCASE, id
        ''', (school_id,)).fetchall()
    finally:
        conn.close()

    by_instructor = {}
    for row in records:
        by_instructor.setdefault(row['instructor_id'], {})[int(row['att_date'][8:10])] = {
            'time': (row['check_in_at'] or '')[11:16],
            'method': row['method'],
        }
    days = []
    for day in range(1, last_day + 1):
        weekday = date(year, month, day).weekday()
        days.append({'day': day, 'weekday': WEEKDAY_LABELS[weekday], 'weekend': weekday >= 5})
    rows = []
    for instructor in instructors:
        if not instructor['is_active'] and instructor['id'] not in recorded_ids:
            continue
        marks = by_instructor.get(instructor['id'], {})
        rows.append({
            'instructor_id': instructor['id'],
            'name': instructor['name'],
            'subject': instructor['subject'] or '',
            'removed': not instructor['is_active'],
            'days': {str(day): mark for day, mark in marks.items()},
            'total': len(marks),
        })
    day_totals = {str(day['day']): 0 for day in days}
    for row in rows:
        for day in row['days']:
            day_totals[day] += 1
    return {
        'month': f'{year:04d}-{month:02d}',
        'school_name': school['school_name'],
        'days': days,
        'rows': rows,
        'day_totals': day_totals,
    }, None


@instructor_attendance_bp.route('/instructor-attendance/monthly.xlsx')
@menu_permission_required(MENU_KEY)
def monthly_excel():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    data, error = _monthly_data(request.args.get('month'), _int_or_none(request.args.get('school_id')))
    if error:
        return error, 400
    wb = Workbook()
    ws = wb.active
    ws.title = data['month']
    ws.append([f"{data['school_name']} 강사 출결부 ({data['month']})"])
    ws['A1'].font = Font(bold=True, size=14)
    header = ['강사명', '과목'] + [f"{d['day']}\n{d['weekday']}" for d in data['days']] + ['출석일수']
    ws.append(header)
    weekend_fill = PatternFill('solid', fgColor='F3F4F6')
    present_fill = PatternFill('solid', fgColor='DCFCE7')
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for row in data['rows']:
        values = [row['name'] + (' (제외)' if row['removed'] else ''), row['subject']]
        for d in data['days']:
            mark = row['days'].get(str(d['day']))
            values.append(mark['time'] + (' 수동' if mark['method'] != 'QR' else '') if mark else '')
        values.append(row['total'])
        ws.append(values)
    ws.append(['일별 출석 인원', ''] + [data['day_totals'][str(d['day'])] for d in data['days']]
              + [sum(r['total'] for r in data['rows'])])
    for row_cells in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row_cells:
            cell.alignment = center
            col = cell.column - 3
            if 0 <= col < len(data['days']):
                if cell.row > 2 and cell.value not in (None, '') and cell.row < ws.max_row:
                    cell.fill = present_fill
                elif data['days'][col]['weekend']:
                    cell.fill = weekend_fill
    for cell in ws[2]:
        cell.font = Font(bold=True)
    ws.row_dimensions[2].height = 30
    ws.column_dimensions['A'].width = 14
    ws.column_dimensions['B'].width = 12
    for index in range(len(data['days'])):
        ws.column_dimensions[get_column_letter(index + 3)].width = 7
    ws.freeze_panes = 'C3'
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f"강사출결_{data['school_name']}_{data['month']}.xlsx",
    )


@instructor_attendance_bp.route('/instructor-attendance/api/attendance/manual', methods=['POST'])
@menu_permission_required(MENU_KEY)
def manual_attendance():
    """휴대폰 분실·배터리 방전 등으로 QR을 못 찍은 경우 관리자가 직접 처리한다."""
    data = request.get_json(silent=True) or {}
    instructor_id = _int_or_none(data.get('instructor_id'))
    att_date = _parse_date(data.get('date'))
    action = data.get('action')
    note = str(data.get('note') or '').strip()[:200]
    if not instructor_id or not att_date or action not in {'check', 'cancel'}:
        return _json_error('요청 값을 확인해주세요.')
    conn = get_db()
    try:
        instructor = conn.execute(
            'SELECT id, school_id, name FROM lecturer_instructors WHERE id=?', (instructor_id,)
        ).fetchone()
        if not instructor:
            return _json_error('강사를 찾을 수 없습니다.', 404)
        if action == 'cancel':
            conn.execute(
                'DELETE FROM lecturer_attendance WHERE instructor_id=? AND att_date=?',
                (instructor_id, att_date),
            )
            message = f"{instructor['name']} 강사의 {att_date} 출석을 취소했습니다."
        else:
            time_text = str(data.get('time') or '').strip()
            if not re.fullmatch(r'([01]\d|2[0-3]):[0-5]\d', time_text):
                time_text = _now().strftime('%H:%M') if att_date == _today() else '00:00'
            conn.execute('''
                INSERT INTO lecturer_attendance
                    (instructor_id, school_id, att_date, check_in_at, method, note, recorded_by)
                VALUES (?, ?, ?, ?, '수동', ?, ?)
                ON CONFLICT(instructor_id, att_date) DO UPDATE SET
                    check_in_at=excluded.check_in_at, method='수동',
                    note=excluded.note, recorded_by=excluded.recorded_by
            ''', (instructor_id, instructor['school_id'], att_date, f'{att_date} {time_text}:00',
                  note, session.get('user_name')))
            message = f"{instructor['name']} 강사를 {att_date} {time_text} 출석으로 처리했습니다."
        conn.commit()
    finally:
        conn.close()
    return jsonify({'status': 'success', 'message': message})


@instructor_attendance_bp.route('/instructor-attendance/api/qr')
@menu_permission_required(MENU_KEY)
def qr_info():
    school_id = _int_or_none(request.args.get('school_id'))
    conn = get_db()
    try:
        school = _get_school(conn, school_id) if school_id else None
        if not school:
            return _json_error('학교를 선택해주세요.')
        token = get_daily_token(conn, school_id)
        display_key = _get_display_key(conn, school_id)
        issued = conn.execute(
            'SELECT issued_by, issued_at FROM lecturer_qr_codes WHERE school_id=? AND qr_date=?',
            (school_id, _today()),
        ).fetchone()
    finally:
        conn.close()
    today = _now().date()
    scan_url = _scan_url(school_id, token)
    return jsonify({
        'status': 'success',
        'school_name': school['school_name'],
        'date': today.isoformat(),
        'weekday': WEEKDAY_LABELS[today.weekday()],
        'scan_url': scan_url,
        'scan_url_warning': _scan_url_warning(scan_url),
        'qr_svg_url': url_for('instructor_attendance.admin_qr_svg', school_id=school_id, t=token[:6]),
        'display_url': _external_root() + url_for('instructor_attendance.display_page', display_key=display_key),
        'issued_by': issued['issued_by'] if issued else '',
        'issued_at': issued['issued_at'] if issued else '',
    })


@instructor_attendance_bp.route('/instructor-attendance/qr.svg')
@menu_permission_required(MENU_KEY)
def admin_qr_svg():
    school_id = _int_or_none(request.args.get('school_id'))
    conn = get_db()
    try:
        if not school_id or not _get_school(conn, school_id):
            return '학교를 선택해주세요.', 400
        token = get_daily_token(conn, school_id)
    finally:
        conn.close()
    return _svg_response(_scan_url(school_id, token))


@instructor_attendance_bp.route('/instructor-attendance/api/qr/reissue', methods=['POST'])
@menu_permission_required(MENU_KEY)
def reissue_qr():
    data = request.get_json(silent=True) or {}
    school_id = _int_or_none(data.get('school_id'))
    target = data.get('target')
    conn = get_db()
    try:
        if not school_id or not _get_school(conn, school_id):
            return _json_error('학교를 선택해주세요.')
        if target == 'display':
            conn.execute('DELETE FROM lecturer_qr_displays WHERE school_id=?', (school_id,))
            message = '태블릿 표시용 링크를 새로 만들었습니다. 기존 링크는 더 이상 열리지 않습니다.'
        else:
            conn.execute('DELETE FROM lecturer_qr_codes WHERE school_id=? AND qr_date=?', (school_id, _today()))
            message = '오늘 QR을 재발급했습니다. 이전 QR로는 출석할 수 없습니다.'
        conn.commit()
    finally:
        conn.close()
    return jsonify({'status': 'success', 'message': message})


# ---------------------------------------------------------------- 학교 비치용 QR 표시(로그인 없음)

@instructor_attendance_bp.route('/instructor-attendance/display/<display_key>')
def display_page(display_key):
    conn = get_db()
    try:
        row = conn.execute('''
            SELECT s.id, s.school_name FROM lecturer_qr_displays q
            JOIN schools s ON s.id = q.school_id WHERE q.display_key=?
        ''', (display_key,)).fetchone()
    finally:
        conn.close()
    if not row:
        return render_template('instructor_attendance/display.html', invalid=True), 404
    today = _now().date()
    return render_template(
        'instructor_attendance/display.html',
        invalid=False,
        school_name=row['school_name'],
        display_key=display_key,
        today=today.isoformat(),
        weekday=WEEKDAY_LABELS[today.weekday()],
    )


@instructor_attendance_bp.route('/instructor-attendance/display/<display_key>/qr.svg')
def display_qr_svg(display_key):
    conn = get_db()
    try:
        row = conn.execute(
            'SELECT school_id FROM lecturer_qr_displays WHERE display_key=?', (display_key,)
        ).fetchone()
        if not row:
            return '사용할 수 없는 링크입니다.', 404
        token = get_daily_token(conn, row['school_id'])
    finally:
        conn.close()
    return _svg_response(_scan_url(row['school_id'], token))


# ---------------------------------------------------------------- 강사 휴대폰 스캔(로그인 없음)

def _scan_context(conn, school_id, token):
    """QR이 오늘 발급된 유효한 것인지 확인한다. (학교행, 오류메시지)"""
    school = _get_school(conn, school_id)
    if not school:
        return None, '등록되지 않은 학교 QR입니다.'
    today_token = get_daily_token(conn, school_id, create=False)
    if not today_token or not secrets.compare_digest(today_token, str(token)):
        return school, '오늘 발급된 QR이 아닙니다. 학교에 비치된 오늘의 QR을 다시 찍어주세요.'
    return school, None


@instructor_attendance_bp.route('/instructor-attendance/scan/<int:school_id>/<token>')
def scan_page(school_id, token):
    conn = get_db()
    try:
        school, error = _scan_context(conn, school_id, token)
    finally:
        conn.close()
    today = _now().date()
    response = Response(render_template(
        'instructor_attendance/scan.html',
        school_name=school['school_name'] if school else '',
        error=error,
        school_id=school_id,
        token=token,
        today=today.isoformat(),
        weekday=WEEKDAY_LABELS[today.weekday()],
    ))
    response.headers['Cache-Control'] = 'no-store'
    return response


@instructor_attendance_bp.route('/instructor-attendance/scan/<int:school_id>/<token>', methods=['POST'])
def scan_check_in(school_id, token):
    data = request.get_json(silent=True) or {}
    device_token = request.cookies.get(DEVICE_COOKIE, '')
    new_device_token = None
    conn = get_db()
    try:
        school, error = _scan_context(conn, school_id, token)
        if error:
            return _json_error(error, 410)

        phone = None
        if device_token:
            device = conn.execute(
                'SELECT phone FROM lecturer_devices WHERE token_hash=?', (_hash_token(device_token),)
            ).fetchone()
            if device:
                phone = device['phone']

        if phone:
            instructor = conn.execute('''
                SELECT id, name FROM lecturer_instructors
                WHERE school_id=? AND phone=? AND is_active=1 ORDER BY id LIMIT 1
            ''', (school_id, phone)).fetchone()
            if not instructor:
                return _json_error(
                    f"이 휴대폰({format_phone(phone)})은 {school['school_name']} 강사로 등록되어 있지 않습니다. "
                    '담당자에게 강사 등록을 요청해주세요.', 403)
        else:
            name = str(data.get('name') or '').strip()
            entered_phone = normalize_phone(data.get('phone'))
            if not name and not entered_phone:
                return jsonify({'status': 'need_register', 'message': '처음 사용하는 휴대폰입니다. 본인 확인을 해주세요.'})
            client_key = request.headers.get('X-Forwarded-For', request.remote_addr or 'unknown').split(',')[0].strip()
            if _register_rate_limited(client_key):
                return _json_error('입력 시도가 너무 많습니다. 10분 뒤 다시 시도해주세요.', 429)
            instructor = conn.execute('''
                SELECT id, name, phone FROM lecturer_instructors
                WHERE school_id=? AND phone=? AND REPLACE(name, ' ', '')=? AND is_active=1
                ORDER BY id LIMIT 1
            ''', (school_id, entered_phone, name.replace(' ', ''))).fetchone()
            if not instructor:
                return _json_error(
                    f"{school['school_name']}에 등록된 강사 정보와 일치하지 않습니다. "
                    '이름과 휴대폰번호를 확인하거나 담당자에게 문의해주세요.', 404)
            if conn.execute('SELECT 1 FROM lecturer_devices WHERE phone=?', (instructor['phone'],)).fetchone():
                return _json_error(
                    '이 번호는 이미 다른 휴대폰에 연결되어 있습니다. '
                    '휴대폰을 바꾸셨다면 담당자에게 [기기 초기화]를 요청해주세요.', 409)
            new_device_token = secrets.token_urlsafe(32)
            conn.execute('''
                INSERT INTO lecturer_devices (phone, token_hash, user_agent, bound_at, last_seen_at)
                VALUES (?, ?, ?, ?, ?)
            ''', (instructor['phone'], _hash_token(new_device_token),
                  (request.headers.get('User-Agent') or '')[:200], _now_text(), _now_text()))
            phone = instructor['phone']

        now_text = _now_text()
        conn.execute('UPDATE lecturer_devices SET last_seen_at=? WHERE phone=?', (now_text, phone))
        existing = conn.execute(
            'SELECT check_in_at FROM lecturer_attendance WHERE instructor_id=? AND att_date=?',
            (instructor['id'], _today()),
        ).fetchone()
        if existing:
            already, check_in_at = True, existing['check_in_at']
        else:
            conn.execute('''
                INSERT INTO lecturer_attendance (instructor_id, school_id, att_date, check_in_at, method)
                VALUES (?, ?, ?, ?, 'QR')
            ''', (instructor['id'], school_id, _today(), now_text))
            already, check_in_at = False, now_text
        conn.commit()
    finally:
        conn.close()

    response = jsonify({
        'status': 'success',
        'already': already,
        'name': instructor['name'],
        'school_name': school['school_name'],
        'check_in_time': check_in_at[11:16],
        'device_registered': bool(new_device_token),
        'message': ('이미 출석 처리되었습니다.' if already else '출석이 완료되었습니다.'),
    })
    if new_device_token:
        response.set_cookie(
            DEVICE_COOKIE, new_device_token, max_age=DEVICE_COOKIE_MAX_AGE,
            httponly=True, samesite='Lax', path='/instructor-attendance/scan',
            secure=request.is_secure or request.headers.get('X-Forwarded-Proto', '').startswith('https'),
        )
    return response
