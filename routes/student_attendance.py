"""학교관리 > 학생출석관리.

방과후 수업에 온 학생이 교실의 QR을 자기 휴대폰으로 찍으면 출석이 기록되고,
보호자에게 무료 웹푸시(브라우저 알림)로 바로 전달된다. 문자·알림톡은 쓰지 않는다.

- 과목(강좌)·학생·보호자·푸시구독은 [학부모알림전송]과 같은 테이블을 함께 쓴다.
  (parent_classes / parent_students / parent_guardians / parent_push_subscriptions)
  그래서 보호자는 한 번만 알림을 등록하면 출결 알림과 일반 안내를 모두 받는다.
- 한 학교에 과목을 여러 개 만들 수 있고, 과목마다 QR·강사용 실시간 화면·
  학부모 등록 링크가 따로 발급된다.
- 학생 휴대폰 인식: 처음 QR을 찍을 때 이름과 보호자 휴대폰 뒷 4자리를 확인하고
  그 기기에 무작위 기기키(쿠키)를 심는다. 학생 한 명에는 기기 하나만 연결되며,
  형제자매는 한 기기를 함께 쓸 수 있다.
- QR 방식: 'daily'는 과목×날짜별 고정 QR(인쇄 가능), 'live'는 강사 화면에
  30초마다 바뀌는 QR만 인정해 사진을 돌려 찍는 대리출석을 막는다.
- 강사는 로그인 없이 과목별 비밀 링크로 실시간 출석 화면(조작 가능)을 열고,
  교실 TV·태블릿에는 조작 기능이 없는 '교실 QR 화면' 링크를 따로 띄운다.
  출석 변동은 Socket.IO로 즉시 반영된다(연결이 끊겨도 주기적으로 다시 조회한다).
"""

import base64
import hashlib
import hmac
import logging
import re
import secrets
import threading
import time
from calendar import monthrange
from datetime import date, datetime, timedelta, timezone

from flask import (
    Blueprint, Response, current_app, jsonify, redirect, render_template, request,
    send_file, session, url_for,
)
from flask_socketio import join_room
from itsdangerous import BadSignature, SignatureExpired, URLSafeTimedSerializer

from . import parent_notifications as pn
from .database import get_db
from .instructor_attendance import (
    _external_root, _qr_svg, _scan_url_warning, format_phone, normalize_phone,
)
from .security import has_menu_permission, menu_permission_required
from .socketio_ext import socketio

student_attendance_bp = Blueprint('student_attendance', __name__)
logger = logging.getLogger(__name__)

MENU_KEY = 'school_student_attendance'
KST = timezone(timedelta(hours=9))
WEEKDAY_LABELS = ('월', '화', '수', '목', '금', '토', '일')
SOCKET_NS = '/student-attendance'

DEVICE_COOKIE = 'saedam_student_device'
DEVICE_COOKIE_MAX_AGE = 60 * 60 * 24 * 365 * 5
LIVE_QR_WINDOW_SECONDS = 30
LIVE_QR_ACCEPT_WINDOWS = 3          # 현재 포함 최근 3개 창(최대 90초 전 QR)까지 인정
SCAN_TICKET_MAX_AGE = 15 * 60       # QR을 연 뒤 본인확인까지 15분 안에 끝내면 된다
CHECK_STATUSES = ('출석', '지각')
MARK_STATUSES = ('출석', '지각', '결석')
RESET_DEVICE_MESSAGE = '휴대폰 등록을 풀었습니다. 학생이 QR을 찍고 이름과 부모님 휴대폰 뒷 4자리를 다시 입력하면 됩니다.'

# 본인확인 '실패'만 센다. 한 반 학생들이 같은 학교 와이파이(같은 IP)로 첫날 한꺼번에
# 등록해도 막히지 않게 하고, 이름·번호를 바꿔 가며 맞히려는 시도만 과목×IP 단위로 제한한다.
REGISTER_WINDOW_SECONDS = 10 * 60
REGISTER_MAX_FAILURES = 20
_register_attempts = {}


# ---------------------------------------------------------------- 스키마

def ensure_student_attendance_schema(conn):
    pn.ensure_parent_notification_schema(conn)
    conn.executescript('''
        -- parent_classes 한 줄(과목)에 붙는 출석 설정.
        CREATE TABLE IF NOT EXISTS student_att_settings (
            class_id INTEGER PRIMARY KEY,
            weekdays TEXT,
            start_time TEXT,
            end_time TEXT,
            late_minutes INTEGER NOT NULL DEFAULT 10,
            qr_mode TEXT NOT NULL DEFAULT 'daily',
            notify_parents INTEGER NOT NULL DEFAULT 1,
            live_key TEXT NOT NULL UNIQUE,
            parent_key TEXT NOT NULL UNIQUE,
            display_key TEXT,
            updated_at TEXT
        );

        -- 과목×날짜별 QR 토큰. 실시간 QR도 이 값을 비밀키로 삼아 30초마다 파생한다.
        CREATE TABLE IF NOT EXISTS student_att_qr (
            class_id INTEGER NOT NULL,
            qr_date TEXT NOT NULL,
            token TEXT NOT NULL,
            issued_at TEXT,
            PRIMARY KEY (class_id, qr_date)
        );

        -- 학생 한 명에 기기 하나. 형제자매는 같은 기기키를 함께 쓴다. 쿠키 원문은 저장하지 않는다.
        CREATE TABLE IF NOT EXISTS student_att_devices (
            student_id INTEGER PRIMARY KEY,
            token_hash TEXT NOT NULL,
            user_agent TEXT,
            bound_at TEXT,
            last_seen_at TEXT
        );
        CREATE INDEX IF NOT EXISTS idx_student_att_devices_token
            ON student_att_devices(token_hash);

        CREATE TABLE IF NOT EXISTS student_att_records (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            class_id INTEGER NOT NULL,
            student_id INTEGER NOT NULL,
            att_date TEXT NOT NULL,
            status TEXT NOT NULL,
            check_in_at TEXT,
            method TEXT NOT NULL DEFAULT 'QR',
            note TEXT,
            recorded_by TEXT,
            dismissed_at TEXT,
            notified_at TEXT,
            notify_result TEXT,
            notification_id INTEGER,
            updated_at TEXT,
            UNIQUE (class_id, student_id, att_date)
        );
        CREATE INDEX IF NOT EXISTS idx_student_att_records_date
            ON student_att_records(att_date, class_id);
    ''')
    columns = {row[1] for row in conn.execute('PRAGMA table_info(student_att_settings)').fetchall()}
    if 'display_key' not in columns:
        # 교실 QR 화면(학생용·조작 불가) 링크는 나중에 추가된 열이라 기존 DB에 채워 넣는다.
        conn.execute('ALTER TABLE student_att_settings ADD COLUMN display_key TEXT')
    conn.execute('CREATE UNIQUE INDEX IF NOT EXISTS idx_student_att_settings_display '
                 'ON student_att_settings(display_key)')
    # 출결 기록마다 마지막으로 보낸 보호자 알림을 연결해 수신 확인을 보여준다.
    pn._add_missing_columns(conn, 'student_att_records', {'notification_id': 'INTEGER'})
    conn.commit()


def init_student_attendance_schema():
    conn = get_db()
    try:
        ensure_student_attendance_schema(conn)
    finally:
        conn.close()


# ---------------------------------------------------------------- 공통

def _now():
    return datetime.now(KST)


def _today():
    return _now().date().isoformat()


def _now_text():
    return _now().strftime('%Y-%m-%d %H:%M:%S')


def _json_error(message, status=400, **extra):
    payload = {'status': 'error', 'message': message}
    payload.update(extra)
    return jsonify(payload), status


def _text(value, limit=200):
    return str(value or '').strip()[:limit]


def _int_or_none(value):
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _parse_date(value, default=None):
    try:
        return date.fromisoformat(str(value or '').strip()).isoformat()
    except ValueError:
        return default


def _valid_time(value):
    value = _text(value, 5)
    return value if re.fullmatch(r'([01]\d|2[0-3]):[0-5]\d', value) else ''


def _hash_token(token):
    return hashlib.sha256(str(token).encode('utf-8')).hexdigest()


def _actor():
    return session.get('user_name') or '관리자'


def _recent_failures(client_key):
    now = time.monotonic()
    attempts = [t for t in _register_attempts.get(client_key, []) if now - t < REGISTER_WINDOW_SECONDS]
    if attempts:
        _register_attempts[client_key] = attempts
    else:
        _register_attempts.pop(client_key, None)
    return attempts


def _too_many_failures(client_key):
    return len(_recent_failures(client_key)) >= REGISTER_MAX_FAILURES


def _record_failure(client_key):
    if len(_register_attempts) > 5000:
        # 오래된 IP 기록이 끝없이 쌓이지 않게 정리한다.
        for key in list(_register_attempts):
            _recent_failures(key)
    _register_attempts.setdefault(client_key, []).append(time.monotonic())


def _client_key(prefix):
    ip = request.headers.get('X-Forwarded-For', request.remote_addr or 'unknown').split(',')[0].strip()
    return f'{prefix}:{ip}'


def _list_schools(conn):
    columns = {row[1] for row in conn.execute('PRAGMA table_info(schools)').fetchall()}
    where = 'WHERE COALESCE(is_active, 1) = 1' if 'is_active' in columns else ''
    rows = conn.execute(f'''
        SELECT id, school_name, year FROM schools {where}
        ORDER BY year DESC, school_name COLLATE NOCASE, id
    ''').fetchall()
    result = []
    for row in rows:
        year = str(row['year'] or '').strip()
        label = f"{row['school_name']} ({year})" if year else row['school_name']
        result.append({'id': row['id'], 'name': row['school_name'], 'label': label})
    return result


def _weekday_list(text):
    days = set()
    for part in str(text or '').split(','):
        value = _int_or_none(part)
        if value is not None and 0 <= value <= 6:
            days.add(value)
    return sorted(days)


def _schedule_label(settings):
    days = ''.join(WEEKDAY_LABELS[d] for d in _weekday_list(settings['weekdays']))
    times = settings['start_time'] or ''
    if times and settings['end_time']:
        times += f"~{settings['end_time']}"
    return ' '.join(part for part in (days, times) if part)


def _ensure_settings(conn, class_id):
    row = conn.execute('SELECT * FROM student_att_settings WHERE class_id=?', (class_id,)).fetchone()
    if row and row['display_key']:
        return row
    if row:
        conn.execute('UPDATE student_att_settings SET display_key=? WHERE class_id=? AND display_key IS NULL',
                     (secrets.token_urlsafe(24), class_id))
    else:
        conn.execute('''
            INSERT OR IGNORE INTO student_att_settings (class_id, live_key, parent_key, display_key, updated_at)
            VALUES (?, ?, ?, ?, ?)
        ''', (class_id, secrets.token_urlsafe(24), secrets.token_urlsafe(18), secrets.token_urlsafe(24),
              _now_text()))
    conn.commit()
    return conn.execute('SELECT * FROM student_att_settings WHERE class_id=?', (class_id,)).fetchone()


def _load_class(conn, class_id):
    """활성 과목과 출석 설정을 함께 읽는다. 없으면 None."""
    row = conn.execute(
        'SELECT * FROM parent_classes WHERE id=? AND is_active=1', (class_id,)
    ).fetchone()
    if not row:
        return None
    settings = _ensure_settings(conn, class_id)
    info = dict(row)
    info.update({key: settings[key] for key in settings.keys() if key != 'class_id'})
    return info


def _class_by_key(conn, column, key):
    if column not in {'live_key', 'parent_key', 'display_key'} or not key:
        return None
    row = conn.execute(f'''
        SELECT c.id FROM student_att_settings s
        JOIN parent_classes c ON c.id=s.class_id AND c.is_active=1
        WHERE s.{column}=?
    ''', (key,)).fetchone()
    return _load_class(conn, row['id']) if row else None


# ---------------------------------------------------------------- QR

def get_daily_token(conn, class_id, qr_date=None, create=True):
    qr_date = qr_date or _today()
    row = conn.execute(
        'SELECT token FROM student_att_qr WHERE class_id=? AND qr_date=?', (class_id, qr_date)
    ).fetchone()
    if row or not create:
        return row['token'] if row else None
    conn.execute('''
        INSERT OR IGNORE INTO student_att_qr (class_id, qr_date, token, issued_at)
        VALUES (?, ?, ?, ?)
    ''', (class_id, qr_date, secrets.token_urlsafe(18), _now_text()))
    conn.commit()
    return conn.execute(
        'SELECT token FROM student_att_qr WHERE class_id=? AND qr_date=?', (class_id, qr_date)
    ).fetchone()['token']


def _live_window():
    return int(_now().timestamp() // LIVE_QR_WINDOW_SECONDS)


def _live_token(daily_token, window):
    digest = hmac.new(daily_token.encode('utf-8'), str(window).encode('ascii'), hashlib.sha256).digest()
    return base64.urlsafe_b64encode(digest[:12]).decode('ascii')


def _current_scan_token(conn, class_info):
    daily = get_daily_token(conn, class_info['id'])
    if class_info['qr_mode'] == 'live':
        return _live_token(daily, _live_window())
    return daily


def _scan_token_valid(conn, class_info, token):
    daily = get_daily_token(conn, class_info['id'], create=False)
    if not daily or not token:
        return False
    token = str(token)
    if class_info['qr_mode'] != 'live':
        return secrets.compare_digest(daily, token)
    window = _live_window()
    return any(
        secrets.compare_digest(_live_token(daily, window - offset), token)
        for offset in range(LIVE_QR_ACCEPT_WINDOWS)
    )


def _scan_url(class_id, token):
    return _external_root() + url_for('student_attendance.scan_page', class_id=class_id, token=token)


def _live_url(class_info):
    return _external_root() + url_for('student_attendance.live_page', live_key=class_info['live_key'])


def _display_url(class_info):
    return _external_root() + url_for('student_attendance.display_page', display_key=class_info['display_key'])


def _parent_url(class_info):
    return _external_root() + url_for('student_attendance.parent_lookup', parent_key=class_info['parent_key'])


def _svg_response(data):
    response = Response(_qr_svg(data), mimetype='image/svg+xml')
    response.headers['Cache-Control'] = 'no-store'
    return response


def _ticket_serializer():
    return URLSafeTimedSerializer(current_app.secret_key, salt='student-attendance-scan')


def _make_ticket(class_id):
    return _ticket_serializer().dumps({'c': class_id, 'd': _today()})


def _ticket_valid(ticket, class_id):
    try:
        data = _ticket_serializer().loads(str(ticket or ''), max_age=SCAN_TICKET_MAX_AGE)
    except (BadSignature, SignatureExpired):
        return False
    return data.get('c') == class_id and data.get('d') == _today()


# ---------------------------------------------------------------- 학생·보호자 명단

def _find_or_create_guardian(conn, name, phone):
    row = conn.execute('''
        SELECT id FROM parent_guardians WHERE phone=? AND is_active=1
        ORDER BY (name=?) DESC, id LIMIT 1
    ''', (phone, name)).fetchone()
    if row:
        return row['id']
    return conn.execute(
        'INSERT INTO parent_guardians (name, phone) VALUES (?, ?)', (name, phone)
    ).lastrowid


def _save_student(conn, class_info, data, student_id=None):
    """학생 1명을 과목 명단에 넣거나 고친다. (오류메시지 또는 None)"""
    name = _text(data.get('name'), 40)
    grade = _text(data.get('grade'), 10)
    classroom = _text(data.get('classroom'), 10)
    guardian_name = _text(data.get('guardian_name'), 40) or (f'{name} 보호자' if name else '')
    phone = normalize_phone(data.get('guardian_phone'))
    if not name:
        return '학생 이름을 입력해주세요.'
    if not re.fullmatch(r'01\d{8,9}', phone):
        return f'{name}: 보호자 휴대폰번호(010으로 시작하는 10~11자리)를 확인해주세요.'

    school_id, school_name = class_info['school_id'], class_info['school_name']
    if student_id:
        enrolled = conn.execute(
            'SELECT 1 FROM parent_class_students WHERE class_id=? AND student_id=?',
            (class_info['id'], student_id),
        ).fetchone()
        if not enrolled:
            return '이 과목의 수강생이 아닙니다.'
        conn.execute('''
            UPDATE parent_students SET name=?, grade=?, classroom=?, updated_at=CURRENT_TIMESTAMP
            WHERE id=?
        ''', (name, grade, classroom, student_id))
    else:
        # 같은 학교에서 이미 다른 과목을 듣는 학생이면 그 학생에 연결한다.
        row = conn.execute('''
            SELECT id FROM parent_students
            WHERE name=? AND COALESCE(school_id, 0)=COALESCE(?, 0)
              AND COALESCE(grade, '')=? AND COALESCE(classroom, '')=? AND is_active=1
            ORDER BY id LIMIT 1
        ''', (name, school_id, grade, classroom)).fetchone()
        if row:
            student_id = row['id']
        else:
            student_id = conn.execute('''
                INSERT INTO parent_students (name, school_id, school_name, grade, classroom)
                VALUES (?, ?, ?, ?, ?)
            ''', (name, school_id, school_name, grade, classroom)).lastrowid
        conn.execute(
            'INSERT OR IGNORE INTO parent_class_students (class_id, student_id) VALUES (?, ?)',
            (class_info['id'], student_id),
        )

    guardian_id = _find_or_create_guardian(conn, guardian_name, phone)
    previous = _int_or_none(data.get('guardian_id'))
    if previous and previous != guardian_id:
        conn.execute(
            'DELETE FROM parent_guardian_students WHERE guardian_id=? AND student_id=?',
            (previous, student_id),
        )
    conn.execute(
        'INSERT OR IGNORE INTO parent_guardian_students (guardian_id, student_id) VALUES (?, ?)',
        (guardian_id, student_id),
    )
    # pn._active_invite는 바로 commit하므로, 일괄등록을 통째로 되돌릴 수 있게 여기서는 직접 만든다.
    if not conn.execute('SELECT 1 FROM parent_invites WHERE guardian_id=? AND is_active=1',
                        (guardian_id,)).fetchone():
        conn.execute('INSERT INTO parent_invites (guardian_id, token) VALUES (?, ?)',
                     (guardian_id, secrets.token_urlsafe(32)))
    return None


def _roster(conn, class_id):
    students = conn.execute('''
        SELECT s.id, s.name, s.grade, s.classroom, d.bound_at, d.last_seen_at
        FROM parent_class_students cs
        JOIN parent_students s ON s.id=cs.student_id AND s.is_active=1
        LEFT JOIN student_att_devices d ON d.student_id=s.id
        WHERE cs.class_id=?
        ORDER BY s.name COLLATE NOCASE, s.id
    ''', (class_id,)).fetchall()
    ids = [row['id'] for row in students]
    guardians = {}
    if ids:
        marks = ','.join('?' * len(ids))
        for row in conn.execute(f'''
            SELECT gs.student_id, g.id, g.name, g.phone,
                   (SELECT COUNT(*) FROM parent_push_subscriptions ps
                    WHERE ps.guardian_id=g.id AND ps.is_active=1) AS push_count,
                   (SELECT token FROM parent_invites i
                    WHERE i.guardian_id=g.id AND i.is_active=1 ORDER BY i.id DESC LIMIT 1) AS invite_token
            FROM parent_guardian_students gs
            JOIN parent_guardians g ON g.id=gs.guardian_id AND g.is_active=1
            WHERE gs.student_id IN ({marks})
            ORDER BY g.id
        ''', ids).fetchall():
            guardians.setdefault(row['student_id'], []).append({
                'id': row['id'],
                'name': row['name'],
                'phone': format_phone(row['phone']),
                'push': row['push_count'] > 0,
                'invite_path': f"/parent/register/{row['invite_token']}" if row['invite_token'] else '',
            })
        health = pn.receipt_health(conn, [g['id'] for items in guardians.values() for g in items])
        for items in guardians.values():
            for guardian in items:
                receipt = health.get(guardian['id'], {})
                guardian['last_received_at'] = receipt.get('last_received_at', '')
                guardian['receipt_warn'] = bool(receipt.get('warn'))
    return [{
        'id': row['id'],
        'name': row['name'],
        'grade': row['grade'] or '',
        'classroom': row['classroom'] or '',
        'device_bound_at': row['bound_at'] or '',
        'device_last_seen_at': row['last_seen_at'] or '',
        'guardians': guardians.get(row['id'], []),
    } for row in students]


# ---------------------------------------------------------------- 출석 기록·알림

def _day_items(conn, class_id, att_date):
    """그날 과목 명단 + 출석기록. 명단에서 빠진 학생도 기록이 있으면 함께 보여준다."""
    roster = {item['id']: item for item in _roster(conn, class_id)}
    records = {row['student_id']: row for row in conn.execute('''
        SELECT r.*, s.name, s.grade, s.classroom, n.kind AS notice_kind,
               (SELECT COUNT(*) FROM parent_notification_recipients x
                WHERE x.notification_id=r.notification_id AND x.status='발송') AS notice_sent,
               (SELECT COUNT(*) FROM parent_notification_recipients x
                WHERE x.notification_id=r.notification_id AND x.received_at IS NOT NULL) AS notice_received,
               (SELECT MIN(x.received_at) FROM parent_notification_recipients x
                WHERE x.notification_id=r.notification_id) AS notice_received_at
        FROM student_att_records r
        JOIN parent_students s ON s.id=r.student_id
        LEFT JOIN parent_notifications n ON n.id=r.notification_id
        WHERE r.class_id=? AND r.att_date=?
    ''', (class_id, att_date)).fetchall()}
    items = []
    for student_id in list(roster) + [sid for sid in records if sid not in roster]:
        base = roster.get(student_id)
        record = records.get(student_id)
        items.append({
            'student_id': student_id,
            'name': base['name'] if base else record['name'],
            'grade': (base or record)['grade'] or '',
            'classroom': (base or record)['classroom'] or '',
            'removed': base is None,
            'status': record['status'] if record else '',
            'check_in_time': (record['check_in_at'] or '')[11:16] if record else '',
            'method': record['method'] if record else '',
            'note': (record['note'] or '') if record else '',
            'dismissed_time': (record['dismissed_at'] or '')[11:16] if record else '',
            'notify_result': (record['notify_result'] or '') if record else '',
            # 마지막으로 보낸 보호자 알림과, 그 알림이 휴대폰에 실제로 뜬(수신 확인) 보호자 수
            'notice_kind': (record['notice_kind'] or '') if record else '',
            'notice_sent': record['notice_sent'] if record else 0,
            'notice_received': record['notice_received'] if record else 0,
            'notice_received_time': (pn._kst_text(record['notice_received_at']) or '')[11:16]
                                    if record and record['notice_received_at'] else '',
            'device_bound': bool(base and base['device_bound_at']),
            'push_ready': bool(base and any(g['push'] for g in base['guardians'])),
            'receipt_warn': bool(base and any(g['receipt_warn'] for g in base['guardians'])),
            'guardian_count': len(base['guardians']) if base else 0,
        })
    order = {'출석': 0, '지각': 0, '결석': 2, '': 1}
    items.sort(key=lambda it: (order.get(it['status'], 1), it['check_in_time'] or '99', it['name']))
    return items


def _summary(items):
    summary = {'total': len(items), 'present': 0, 'late': 0, 'absent': 0, 'unchecked': 0, 'dismissed': 0}
    for item in items:
        if item['status'] == '출석':
            summary['present'] += 1
        elif item['status'] == '지각':
            summary['late'] += 1
        elif item['status'] == '결석':
            summary['absent'] += 1
        else:
            summary['unchecked'] += 1
        if item['dismissed_time']:
            summary['dismissed'] += 1
    return summary


def _day_payload(conn, class_info, att_date):
    items = _day_items(conn, class_info['id'], att_date)
    return {
        'status': 'success',
        'class_id': class_info['id'],
        'class_name': class_info['class_name'],
        'school_name': class_info['school_name'],
        'date': att_date,
        'weekday': WEEKDAY_LABELS[date.fromisoformat(att_date).weekday()],
        'is_today': att_date == _today(),
        'qr_mode': class_info['qr_mode'],
        'summary': _summary(items),
        'items': items,
    }


def _auto_status(class_info, when):
    """시작시각 + 지각기준(분)이 지나서 찍으면 지각."""
    start = class_info.get('start_time') or ''
    if not re.fullmatch(r'\d{2}:\d{2}', start):
        return '출석'
    hour, minute = map(int, start.split(':'))
    limit = when.replace(hour=hour, minute=minute, second=59, microsecond=0) \
        + timedelta(minutes=int(class_info.get('late_minutes') or 0))
    return '지각' if when > limit else '출석'


def _emit_change(class_info, student_id=None, status='', name=''):
    payload = {'class_id': class_info['id'], 'student_id': student_id, 'status': status, 'name': name}
    try:
        socketio.emit('attendance_changed', payload, namespace=SOCKET_NS, to=f"class:{class_info['id']}")
        if class_info.get('school_id'):
            socketio.emit('attendance_changed', payload, namespace=SOCKET_NS,
                          to=f"school:{class_info['school_id']}")
    except Exception:
        # 소켓이 없어도 화면은 주기 조회로 따라온다.
        pass


def _on_parent_receipt(conn, notification):
    """학부모 휴대폰에 출결 알림이 뜨면(수신 확인) 그 과목 화면에 '보호자 확인'을 바로 반영한다."""
    if not notification or notification['target_type'] != '학생출석' or not notification['class_id']:
        return
    row = conn.execute('SELECT id, school_id FROM parent_classes WHERE id=?',
                       (notification['class_id'],)).fetchone()
    if row:
        _emit_change({'id': row['id'], 'school_id': row['school_id']}, notification['student_id'], 'receipt')


if _on_parent_receipt not in pn.RECEIPT_LISTENERS:
    pn.RECEIPT_LISTENERS.append(_on_parent_receipt)


def _notice_text(class_info, student_name, kind, time_text):
    where = f"{class_info['school_name']} {class_info['class_name']}"
    body = {
        '출석': f'{student_name} 학생이 {time_text}에 {where} 수업에 출석했습니다.',
        '지각': f'{student_name} 학생이 {time_text}에 {where} 수업에 도착했습니다(지각).',
        '결석': f'{student_name} 학생이 오늘 {where} 수업에 출석하지 않았습니다.',
        '하원': f'{student_name} 학생이 {time_text}에 {where} 수업을 마치고 하원했습니다.',
    }[kind]
    return f'[{kind}] {student_name} 학생', body


def _send_parent_notice(class_id, student_id, att_date, kind, time_text, actor):
    """보호자 웹푸시를 보낸다. 요청 스레드 밖에서도 돌 수 있도록 연결을 새로 연다."""
    conn = get_db()
    try:
        class_info = _load_class(conn, class_id)
        student = conn.execute('SELECT name FROM parent_students WHERE id=?', (student_id,)).fetchone()
        if not class_info or not student:
            return
        guardian_ids = [row['guardian_id'] for row in conn.execute('''
            SELECT gs.guardian_id FROM parent_guardian_students gs
            JOIN parent_guardians g ON g.id=gs.guardian_id AND g.is_active=1
            WHERE gs.student_id=?
        ''', (student_id,)).fetchall()]
        notification_id = None
        if not guardian_ids:
            result = '보호자 미등록'
        else:
            title, body = _notice_text(class_info, student['name'], kind, time_text)
            notification_id, sent, failed = pn._send_notification(
                conn, guardian_ids, kind, title, body, class_id=class_id,
                student_id=student_id, target_type='학생출석', created_by=actor,
            )
            result = f'{kind} 알림 {sent}/{len(guardian_ids)}명' if sent else f'{kind} 알림 미수신(푸시 미등록)'
        conn.execute('''
            UPDATE student_att_records SET notified_at=?, notify_result=?, notification_id=?
            WHERE class_id=? AND student_id=? AND att_date=?
        ''', (_now_text(), result, notification_id, class_id, student_id, att_date))
        conn.commit()
    except Exception:
        logger.exception('학생출석 보호자 알림 실패')
    finally:
        conn.close()


def _dispatch(func, *args):
    """푸시 발송은 수 초 걸릴 수 있어 학생 화면을 기다리게 하지 않는다."""
    threading.Thread(target=func, args=args, daemon=True).start()


def _send_parent_notices(class_id, att_date, kind, targets, actor):
    for student_id, time_text in targets:
        _send_parent_notice(class_id, student_id, att_date, kind, time_text, actor)


def _notify_many(class_info, targets, att_date, kind, actor):
    """targets = [(학생ID, 'HH:MM'), ...]. 일괄 처리도 스레드 하나로 차례차례 보내 DB 잠김을 피한다."""
    if not targets or not class_info.get('notify_parents'):
        return
    # 지난 날짜를 고칠 때는 보호자에게 뒤늦은 알림을 보내지 않는다.
    if att_date != _today():
        return
    _dispatch(_send_parent_notices, class_info['id'], att_date, kind, list(targets), actor)


def _notify(class_info, student_id, att_date, kind, time_text, actor):
    _notify_many(class_info, [(student_id, time_text)], att_date, kind, actor)


def _mark(conn, class_info, student_id, status, att_date, actor, *, time_text='', note='', notify=True,
          emit=True):
    """관리자·강사가 출석 상태를 직접 정한다. (메시지, 오류)"""
    student = conn.execute('SELECT id, name FROM parent_students WHERE id=?', (student_id,)).fetchone()
    enrolled = conn.execute(
        'SELECT 1 FROM parent_class_students WHERE class_id=? AND student_id=?', (class_info['id'], student_id)
    ).fetchone()
    existing = conn.execute(
        'SELECT * FROM student_att_records WHERE class_id=? AND student_id=? AND att_date=?',
        (class_info['id'], student_id, att_date),
    ).fetchone()
    if not student or not (enrolled or existing):
        return None, '이 과목의 학생을 찾을 수 없습니다.'
    now_text = _now_text()

    if status == 'cancel':
        if existing:
            conn.execute('DELETE FROM student_att_records WHERE id=?', (existing['id'],))
            conn.commit()
        if emit:
            _emit_change(class_info, student_id, '', student['name'])
        return f"{student['name']} 학생의 {att_date} 출석 기록을 지웠습니다.", None

    if status == '하원':
        if not existing or existing['status'] not in CHECK_STATUSES:
            return None, '출석한 학생만 하원 처리할 수 있습니다.'
        if existing['dismissed_at']:
            return f"{student['name']} 학생은 이미 하원 처리되었습니다.", None
        conn.execute('UPDATE student_att_records SET dismissed_at=?, updated_at=? WHERE id=?',
                     (now_text, now_text, existing['id']))
        conn.commit()
        if emit:
            _emit_change(class_info, student_id, '하원', student['name'])
        if notify:
            _notify(class_info, student_id, att_date, '하원', now_text[11:16], actor)
        return f"{student['name']} 학생을 하원 처리했습니다.", None

    if status not in MARK_STATUSES:
        return None, '출석·지각·결석 중에서 선택해주세요.'
    time_text = _valid_time(time_text)
    if status == '결석':
        check_in_at = None
    elif time_text:
        check_in_at = f'{att_date} {time_text}:00'
    elif existing and existing['check_in_at']:
        check_in_at = existing['check_in_at']
    else:
        check_in_at = f'{att_date} {now_text[11:16]}:00' if att_date == _today() else f'{att_date} 00:00:00'
    conn.execute('''
        INSERT INTO student_att_records
            (class_id, student_id, att_date, status, check_in_at, method, note, recorded_by, updated_at)
        VALUES (?, ?, ?, ?, ?, '수동', ?, ?, ?)
        ON CONFLICT(class_id, student_id, att_date) DO UPDATE SET
            status=excluded.status, check_in_at=excluded.check_in_at, method='수동',
            note=excluded.note, recorded_by=excluded.recorded_by,
            dismissed_at=CASE WHEN excluded.status='결석' THEN NULL ELSE dismissed_at END,
            updated_at=excluded.updated_at
    ''', (class_info['id'], student_id, att_date, status, check_in_at, _text(note), actor, now_text))
    conn.commit()
    if emit:
        _emit_change(class_info, student_id, status, student['name'])
    changed = not existing or existing['status'] != status
    if notify and changed:
        _notify(class_info, student_id, att_date, status, (check_in_at or now_text)[11:16], actor)
    return f"{student['name']} 학생을 {att_date} {status}(으)로 처리했습니다.", None


def _bulk_mark(conn, class_info, student_ids, status, actor):
    """여러 학생을 같은 상태로 처리하고, 화면 갱신과 보호자 알림은 마지막에 한 번에 보낸다."""
    today, now_hhmm = _today(), _now().strftime('%H:%M')
    done = []
    for student_id in student_ids:
        _, error = _mark(conn, class_info, student_id, status, today, actor, notify=False, emit=False)
        if not error:
            done.append((student_id, now_hhmm))
    if done:
        _emit_change(class_info, status='bulk')
        _notify_many(class_info, done, today, status, actor)
    return len(done)


def _bulk_message(class_info, count, status):
    if not count:
        return f'{status} 처리할 학생이 없습니다.'
    if not class_info.get('notify_parents'):
        return f'{count}명을 {status} 처리했습니다. (이 과목은 보호자 알림이 꺼져 있습니다)'
    return f'{count}명을 {status} 처리했습니다. 보호자에게 {status} 알림을 보냅니다.'


def _dismiss_all(conn, class_info, actor):
    rows = conn.execute(f'''
        SELECT student_id FROM student_att_records
        WHERE class_id=? AND att_date=? AND status IN ({','.join('?' * len(CHECK_STATUSES))})
          AND dismissed_at IS NULL
    ''', (class_info['id'], _today(), *CHECK_STATUSES)).fetchall()
    return _bulk_mark(conn, class_info, [row['student_id'] for row in rows], '하원', actor)


def _absent_rest(conn, class_info, actor):
    """오늘 아직 출석 확인이 안 된 수강생을 모두 결석으로 처리한다."""
    targets = [item['student_id'] for item in _day_items(conn, class_info['id'], _today())
               if not item['status'] and not item['removed']]
    return _bulk_mark(conn, class_info, targets, '결석', actor)


def _monthly_data(conn, class_info, month_text):
    match = re.fullmatch(r'(\d{4})-(\d{2})', str(month_text or '').strip())
    if match and 1 <= int(match.group(2)) <= 12:
        year, month = int(match.group(1)), int(match.group(2))
    else:
        today = _now().date()
        year, month = today.year, today.month
    last_day = monthrange(year, month)[1]
    first, last = f'{year:04d}-{month:02d}-01', f'{year:04d}-{month:02d}-{last_day:02d}'
    records = conn.execute('''
        SELECT r.student_id, r.att_date, r.status, r.check_in_at, r.method, s.name
        FROM student_att_records r JOIN parent_students s ON s.id=r.student_id
        WHERE r.class_id=? AND r.att_date BETWEEN ? AND ?
    ''', (class_info['id'], first, last)).fetchall()
    class_days = set(_weekday_list(class_info['weekdays']))
    days = []
    for day in range(1, last_day + 1):
        weekday = date(year, month, day).weekday()
        days.append({'day': day, 'weekday': WEEKDAY_LABELS[weekday], 'weekend': weekday >= 5,
                     'class_day': weekday in class_days})
    rows, by_student = {}, {}
    for item in _roster(conn, class_info['id']):
        rows[item['id']] = {'student_id': item['id'], 'name': item['name'], 'removed': False}
    for record in records:
        rows.setdefault(record['student_id'], {
            'student_id': record['student_id'], 'name': record['name'], 'removed': True,
        })
        by_student.setdefault(record['student_id'], {})[str(int(record['att_date'][8:10]))] = {
            'status': record['status'],
            'time': (record['check_in_at'] or '')[11:16],
            'method': record['method'],
        }
    result = []
    for student_id, row in rows.items():
        marks = by_student.get(student_id, {})
        row['days'] = marks
        row['present'] = sum(1 for m in marks.values() if m['status'] == '출석')
        row['late'] = sum(1 for m in marks.values() if m['status'] == '지각')
        row['absent'] = sum(1 for m in marks.values() if m['status'] == '결석')
        result.append(row)
    result.sort(key=lambda r: (r['removed'], r['name']))
    return {
        'month': f'{year:04d}-{month:02d}',
        'class_name': class_info['class_name'],
        'school_name': class_info['school_name'],
        'days': days,
        'rows': result,
    }


# ---------------------------------------------------------------- 관리 화면(인트라넷)

def _class_payload(conn, row):
    info = _load_class(conn, row['id'])
    counts = conn.execute('''
        SELECT
          (SELECT COUNT(*) FROM parent_class_students cs
           JOIN parent_students s ON s.id=cs.student_id AND s.is_active=1 WHERE cs.class_id=?) AS students,
          (SELECT COUNT(*) FROM student_att_records WHERE class_id=? AND att_date=? AND status IN ('출석','지각')) AS present,
          (SELECT COUNT(*) FROM student_att_records WHERE class_id=? AND att_date=? AND status='지각') AS late,
          (SELECT COUNT(*) FROM student_att_records WHERE class_id=? AND att_date=? AND status='결석') AS absent
    ''', (row['id'], row['id'], _today(), row['id'], _today(), row['id'], _today())).fetchone()
    today_weekday = _now().weekday()
    return {
        'id': info['id'],
        'class_name': info['class_name'],
        'department': info['department'] or '',
        'instructor_name': info['instructor_name'] or '',
        'weekdays': _weekday_list(info['weekdays']),
        'start_time': info['start_time'] or '',
        'end_time': info['end_time'] or '',
        'late_minutes': info['late_minutes'],
        'qr_mode': info['qr_mode'],
        'notify_parents': bool(info['notify_parents']),
        'schedule_label': _schedule_label(info),
        'today_is_class_day': today_weekday in _weekday_list(info['weekdays']),
        'live_url': _live_url(info),
        'parent_url': _parent_url(info),
        'counts': dict(counts),
    }


def _admin_class(conn, class_id):
    return _load_class(conn, class_id) if class_id else None


@student_attendance_bp.route('/student-attendance')
@menu_permission_required(MENU_KEY)
def admin_page():
    conn = get_db()
    try:
        schools = _list_schools(conn)
    finally:
        conn.close()
    return render_template('student_attendance/admin.html', schools=schools, today=_today(),
                           push_configured=pn._push_ready())


@student_attendance_bp.route('/student-attendance/api/classes')
@menu_permission_required(MENU_KEY)
def list_classes():
    school_id = _int_or_none(request.args.get('school_id'))
    if not school_id:
        return _json_error('학교를 선택해주세요.')
    conn = get_db()
    try:
        rows = conn.execute('''
            SELECT id FROM parent_classes WHERE school_id=? AND is_active=1
            ORDER BY class_name COLLATE NOCASE, id
        ''', (school_id,)).fetchall()
        classes = [_class_payload(conn, row) for row in rows]
    finally:
        conn.close()
    return jsonify({'status': 'success', 'date': _today(), 'classes': classes})


@student_attendance_bp.route('/student-attendance/api/classes', methods=['POST'])
@menu_permission_required(MENU_KEY)
def save_class():
    data = request.get_json(silent=True) or {}
    school_id = _int_or_none(data.get('school_id'))
    class_id = _int_or_none(data.get('id'))
    class_name = _text(data.get('class_name'), 60)
    if not class_name:
        return _json_error('과목명을 입력해주세요.')
    weekdays = ','.join(str(d) for d in _weekday_list(','.join(str(v) for v in data.get('weekdays') or [])))
    start_time, end_time = _valid_time(data.get('start_time')), _valid_time(data.get('end_time'))
    late_minutes = max(0, min(120, _int_or_none(data.get('late_minutes')) or 0))
    qr_mode = 'live' if data.get('qr_mode') == 'live' else 'daily'
    notify_parents = 0 if data.get('notify_parents') is False else 1
    conn = get_db()
    try:
        school = conn.execute('SELECT id, school_name FROM schools WHERE id=?', (school_id,)).fetchone() \
            if school_id else None
        if not school:
            return _json_error('학교를 선택해주세요.')
        duplicate = conn.execute('''
            SELECT id FROM parent_classes
            WHERE school_id=? AND class_name=? AND is_active=1 AND id<>?
        ''', (school_id, class_name, class_id or 0)).fetchone()
        if duplicate:
            return _json_error(f'이 학교에 같은 이름의 과목({class_name})이 이미 있습니다.')
        values = (_text(data.get('department'), 40), class_name, _text(data.get('instructor_name'), 40))
        if class_id:
            if not conn.execute('SELECT 1 FROM parent_classes WHERE id=? AND school_id=? AND is_active=1',
                                (class_id, school_id)).fetchone():
                return _json_error('과목을 찾을 수 없습니다.', 404)
            conn.execute('''
                UPDATE parent_classes SET department=?, class_name=?, instructor_name=?,
                    updated_at=CURRENT_TIMESTAMP WHERE id=?
            ''', (*values, class_id))
        else:
            class_id = conn.execute('''
                INSERT INTO parent_classes (school_id, school_name, department, class_name,
                    instructor_name, access_token)
                VALUES (?, ?, ?, ?, ?, ?)
            ''', (school_id, school['school_name'], *values, secrets.token_urlsafe(32))).lastrowid
        conn.commit()
        _ensure_settings(conn, class_id)
        conn.execute('''
            UPDATE student_att_settings SET weekdays=?, start_time=?, end_time=?, late_minutes=?,
                qr_mode=?, notify_parents=?, updated_at=? WHERE class_id=?
        ''', (weekdays, start_time, end_time, late_minutes, qr_mode, notify_parents, _now_text(), class_id))
        conn.commit()
    finally:
        conn.close()
    return jsonify({'status': 'success', 'id': class_id, 'message': f'{class_name} 과목을 저장했습니다.'})


@student_attendance_bp.route('/student-attendance/api/classes/<int:class_id>/delete', methods=['POST'])
@menu_permission_required(MENU_KEY)
def delete_class(class_id):
    # 지난 출석부와 학부모 알림 기록이 남아야 하므로 사용 중지만 한다.
    conn = get_db()
    try:
        conn.execute('UPDATE parent_classes SET is_active=0, updated_at=CURRENT_TIMESTAMP WHERE id=?', (class_id,))
        conn.commit()
    finally:
        conn.close()
    return jsonify({'status': 'success', 'message': '과목을 삭제했습니다. 지난 출석 기록은 보존됩니다.'})


@student_attendance_bp.route('/student-attendance/api/classes/<int:class_id>/reissue', methods=['POST'])
@menu_permission_required(MENU_KEY)
def reissue_link(class_id):
    target = (request.get_json(silent=True) or {}).get('target')
    conn = get_db()
    try:
        info = _admin_class(conn, class_id)
        if not info:
            return _json_error('과목을 찾을 수 없습니다.', 404)
        if target == 'live':
            conn.execute('UPDATE student_att_settings SET live_key=? WHERE class_id=?',
                         (secrets.token_urlsafe(24), class_id))
            message = '강사용 화면 링크를 새로 만들었습니다. 기존 링크는 더 이상 열리지 않습니다.'
        elif target == 'display':
            conn.execute('UPDATE student_att_settings SET display_key=? WHERE class_id=?',
                         (secrets.token_urlsafe(24), class_id))
            message = '교실 QR 화면 링크를 새로 만들었습니다. 기존에 띄워 둔 화면은 더 이상 열리지 않습니다.'
        elif target == 'parent':
            conn.execute('UPDATE student_att_settings SET parent_key=? WHERE class_id=?',
                         (secrets.token_urlsafe(18), class_id))
            message = '학부모 등록 링크를 새로 만들었습니다. 기존 안내문 QR은 더 이상 쓸 수 없습니다.'
        else:
            conn.execute('DELETE FROM student_att_qr WHERE class_id=? AND qr_date=?', (class_id, _today()))
            message = '오늘 출석 QR을 재발급했습니다. 이전 QR로는 출석할 수 없습니다.'
        conn.commit()
    finally:
        conn.close()
    return jsonify({'status': 'success', 'message': message})


@student_attendance_bp.route('/student-attendance/api/classes/<int:class_id>/qr')
@menu_permission_required(MENU_KEY)
def qr_info(class_id):
    conn = get_db()
    try:
        info = _admin_class(conn, class_id)
        if not info:
            return _json_error('과목을 찾을 수 없습니다.', 404)
        token = get_daily_token(conn, class_id)
    finally:
        conn.close()
    today = _now().date()
    scan_url = _scan_url(class_id, token)
    return jsonify({
        'status': 'success',
        'class_name': info['class_name'],
        'school_name': info['school_name'],
        'qr_mode': info['qr_mode'],
        'date': today.isoformat(),
        'weekday': WEEKDAY_LABELS[today.weekday()],
        'scan_url': scan_url,
        'scan_url_warning': _scan_url_warning(scan_url),
        'qr_svg_url': url_for('student_attendance.admin_qr_svg', class_id=class_id, t=token[:6]),
        'parent_qr_svg_url': url_for('student_attendance.parent_qr_svg', class_id=class_id,
                                     t=info['parent_key'][:6]),
        'live_url': _live_url(info),
        'display_url': _display_url(info),
        'parent_url': _parent_url(info),
    })


@student_attendance_bp.route('/student-attendance/classes/<int:class_id>/qr.svg')
@menu_permission_required(MENU_KEY)
def admin_qr_svg(class_id):
    conn = get_db()
    try:
        info = _admin_class(conn, class_id)
        if not info:
            return '과목을 찾을 수 없습니다.', 404
        # 인쇄는 하루 고정 QR만 가능하다. 실시간 모드는 강사 화면에서만 QR이 뜬다.
        token = get_daily_token(conn, class_id)
    finally:
        conn.close()
    return _svg_response(_scan_url(class_id, token))


@student_attendance_bp.route('/student-attendance/classes/<int:class_id>/parent-qr.svg')
@menu_permission_required(MENU_KEY)
def parent_qr_svg(class_id):
    conn = get_db()
    try:
        info = _admin_class(conn, class_id)
    finally:
        conn.close()
    if not info:
        return '과목을 찾을 수 없습니다.', 404
    return _svg_response(_parent_url(info))


@student_attendance_bp.route('/student-attendance/api/classes/<int:class_id>/students')
@menu_permission_required(MENU_KEY)
def list_students(class_id):
    conn = get_db()
    try:
        info = _admin_class(conn, class_id)
        if not info:
            return _json_error('과목을 찾을 수 없습니다.', 404)
        students = _roster(conn, class_id)
    finally:
        conn.close()
    base = _external_root()
    for student in students:
        for guardian in student['guardians']:
            guardian['invite_url'] = base + guardian.pop('invite_path') if guardian['invite_path'] else ''
    return jsonify({'status': 'success', 'class_name': info['class_name'], 'students': students})


@student_attendance_bp.route('/student-attendance/api/classes/<int:class_id>/students', methods=['POST'])
@menu_permission_required(MENU_KEY)
def save_student(class_id):
    data = request.get_json(silent=True) or {}
    conn = get_db()
    try:
        info = _admin_class(conn, class_id)
        if not info:
            return _json_error('과목을 찾을 수 없습니다.', 404)
        error = _save_student(conn, info, data, _int_or_none(data.get('id')))
        if error:
            conn.rollback()
            return _json_error(error)
        conn.commit()
    finally:
        conn.close()
    return jsonify({'status': 'success', 'message': '학생 정보를 저장했습니다.'})


@student_attendance_bp.route('/student-attendance/api/classes/<int:class_id>/students/bulk', methods=['POST'])
@menu_permission_required(MENU_KEY)
def bulk_students(class_id):
    """'학생명, 학년, 반, 보호자명, 보호자연락처' 여러 줄을 한 번에 등록한다."""
    lines = [line.strip() for line in str((request.get_json(silent=True) or {}).get('text') or '').splitlines()
             if line.strip()]
    if not lines:
        return _json_error('등록할 학생 명단을 입력해주세요.')
    if len(lines) > 300:
        return _json_error('한 번에 300명까지 등록할 수 있습니다.')
    conn = get_db()
    try:
        info = _admin_class(conn, class_id)
        if not info:
            return _json_error('과목을 찾을 수 없습니다.', 404)
        errors = []
        for number, line in enumerate(lines, start=1):
            parts = [part.strip() for part in re.split(r'[,\t]', line)] + [''] * 5
            error = _save_student(conn, info, {
                'name': parts[0], 'grade': re.sub(r'학년$', '', parts[1]), 'classroom': re.sub(r'반$', '', parts[2]),
                'guardian_name': parts[3], 'guardian_phone': parts[4],
            })
            if error:
                errors.append(f'{number}번째 줄 - {error}')
        if errors:
            conn.rollback()
            return _json_error('명단을 확인해주세요. 아무것도 등록하지 않았습니다.', errors=errors[:50])
        conn.commit()
    finally:
        conn.close()
    return jsonify({'status': 'success', 'message': f'학생 {len(lines)}명을 등록했습니다.'})


@student_attendance_bp.route('/student-attendance/api/classes/<int:class_id>/students/<int:student_id>/remove',
                             methods=['POST'])
@menu_permission_required(MENU_KEY)
def remove_student(class_id, student_id):
    conn = get_db()
    try:
        conn.execute('DELETE FROM parent_class_students WHERE class_id=? AND student_id=?', (class_id, student_id))
        conn.commit()
    finally:
        conn.close()
    return jsonify({'status': 'success', 'message': '과목 명단에서 뺐습니다. 지난 출석 기록은 보존됩니다.'})


@student_attendance_bp.route('/student-attendance/api/students/<int:student_id>/reset-device', methods=['POST'])
@menu_permission_required(MENU_KEY)
def reset_device(student_id):
    conn = get_db()
    try:
        conn.execute('DELETE FROM student_att_devices WHERE student_id=?', (student_id,))
        conn.commit()
    finally:
        conn.close()
    return jsonify({'status': 'success', 'message': RESET_DEVICE_MESSAGE})


@student_attendance_bp.route('/student-attendance/api/classes/<int:class_id>/day')
@menu_permission_required(MENU_KEY)
def admin_day(class_id):
    conn = get_db()
    try:
        info = _admin_class(conn, class_id)
        if not info:
            return _json_error('과목을 찾을 수 없습니다.', 404)
        payload = _day_payload(conn, info, _parse_date(request.args.get('date'), _today()))
    finally:
        conn.close()
    return jsonify(payload)


@student_attendance_bp.route('/student-attendance/api/classes/<int:class_id>/mark', methods=['POST'])
@menu_permission_required(MENU_KEY)
def admin_mark(class_id):
    data = request.get_json(silent=True) or {}
    conn = get_db()
    try:
        info = _admin_class(conn, class_id)
        if not info:
            return _json_error('과목을 찾을 수 없습니다.', 404)
        message, error = _mark(
            conn, info, _int_or_none(data.get('student_id')), _text(data.get('status'), 10),
            _parse_date(data.get('date'), _today()), _actor(),
            time_text=data.get('time'), note=data.get('note'), notify=data.get('notify', True) is not False,
        )
    finally:
        conn.close()
    return _json_error(error) if error else jsonify({'status': 'success', 'message': message})


@student_attendance_bp.route('/student-attendance/api/classes/<int:class_id>/dismiss', methods=['POST'])
@menu_permission_required(MENU_KEY)
def admin_dismiss(class_id):
    conn = get_db()
    try:
        info = _admin_class(conn, class_id)
        if not info:
            return _json_error('과목을 찾을 수 없습니다.', 404)
        count = _dismiss_all(conn, info, _actor())
    finally:
        conn.close()
    return jsonify({'status': 'success', 'count': count, 'message': _bulk_message(info, count, '하원')})


@student_attendance_bp.route('/student-attendance/api/classes/<int:class_id>/absent-rest', methods=['POST'])
@menu_permission_required(MENU_KEY)
def admin_absent_rest(class_id):
    conn = get_db()
    try:
        info = _admin_class(conn, class_id)
        if not info:
            return _json_error('과목을 찾을 수 없습니다.', 404)
        count = _absent_rest(conn, info, _actor())
    finally:
        conn.close()
    return jsonify({'status': 'success', 'count': count, 'message': _bulk_message(info, count, '결석')})


@student_attendance_bp.route('/student-attendance/api/classes/<int:class_id>/monthly')
@menu_permission_required(MENU_KEY)
def admin_monthly(class_id):
    conn = get_db()
    try:
        info = _admin_class(conn, class_id)
        if not info:
            return _json_error('과목을 찾을 수 없습니다.', 404)
        data = _monthly_data(conn, info, request.args.get('month'))
    finally:
        conn.close()
    return jsonify({'status': 'success', **data})


@student_attendance_bp.route('/student-attendance/classes/<int:class_id>/monthly.xlsx')
@menu_permission_required(MENU_KEY)
def monthly_excel(class_id):
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    conn = get_db()
    try:
        info = _admin_class(conn, class_id)
        if not info:
            return '과목을 찾을 수 없습니다.', 404
        data = _monthly_data(conn, info, request.args.get('month'))
    finally:
        conn.close()
    fills = {
        '출석': PatternFill('solid', fgColor='DCFCE7'),
        '지각': PatternFill('solid', fgColor='FEF3C7'),
        '결석': PatternFill('solid', fgColor='FEE2E2'),
    }
    weekend_fill = PatternFill('solid', fgColor='F3F4F6')
    wb = Workbook()
    ws = wb.active
    ws.title = data['month']
    ws.append([f"{data['school_name']} {data['class_name']} 학생 출석부 ({data['month']})"])
    ws['A1'].font = Font(bold=True, size=14)
    ws.append(['학생명'] + [f"{d['day']}\n{d['weekday']}" for d in data['days']] + ['출석', '지각', '결석'])
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for row in data['rows']:
        values = [row['name'] + (' (제외)' if row['removed'] else '')]
        for d in data['days']:
            mark = row['days'].get(str(d['day']))
            values.append((mark['status'][0] + ('\n' + mark['time'] if mark['time'] else '')) if mark else '')
        ws.append(values + [row['present'], row['late'], row['absent']])
    for cells in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in cells:
            cell.alignment = center
            index = cell.column - 2
            if 0 <= index < len(data['days']) and cell.row > 2:
                if cell.value:
                    cell.fill = fills.get({'출': '출석', '지': '지각', '결': '결석'}.get(str(cell.value)[0]), weekend_fill)
                elif data['days'][index]['weekend']:
                    cell.fill = weekend_fill
    for cell in ws[2]:
        cell.font = Font(bold=True)
    ws.row_dimensions[2].height = 30
    ws.column_dimensions['A'].width = 14
    for index in range(len(data['days']) + 3):
        ws.column_dimensions[get_column_letter(index + 2)].width = 7
    ws.freeze_panes = 'B3'
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(
        output, as_attachment=True,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        download_name=f"학생출석_{data['school_name']}_{data['class_name']}_{data['month']}.xlsx",
    )


# ---------------------------------------------------------------- 강사용 실시간 화면(로그인 없음, 비밀 링크)

def _live_actor(info):
    return f"강사 {info['instructor_name']}" if info['instructor_name'] else '강사화면'


def _live_context(live_key):
    conn = get_db()
    try:
        return _class_by_key(conn, 'live_key', live_key)
    finally:
        conn.close()


@student_attendance_bp.route('/student-attendance/live/<live_key>')
def live_page(live_key):
    info = _live_context(live_key)
    if not info:
        return render_template('student_attendance/live.html', invalid=True), 404
    today = _now().date()
    response = Response(render_template(
        'student_attendance/live.html', invalid=False, info=info, live_key=live_key,
        display_url=_display_url(info),
        today=today.isoformat(), weekday=WEEKDAY_LABELS[today.weekday()],
        schedule_label=_schedule_label(info), refresh_seconds=LIVE_QR_WINDOW_SECONDS,
    ))
    response.headers['Cache-Control'] = 'no-store'
    response.headers['Referrer-Policy'] = 'no-referrer'
    return response


def _qr_by_key(column, key):
    conn = get_db()
    try:
        info = _class_by_key(conn, column, key)
        if not info:
            return '사용할 수 없는 링크입니다.', 404
        token = _current_scan_token(conn, info)
    finally:
        conn.close()
    return _svg_response(_scan_url(info['id'], token))


@student_attendance_bp.route('/student-attendance/live/<live_key>/qr.svg')
def live_qr_svg(live_key):
    return _qr_by_key('live_key', live_key)


@student_attendance_bp.route('/student-attendance/live/<live_key>/data')
def live_data(live_key):
    conn = get_db()
    try:
        info = _class_by_key(conn, 'live_key', live_key)
        if not info:
            return _json_error('사용할 수 없는 링크입니다.', 404)
        payload = _day_payload(conn, info, _today())
    finally:
        conn.close()
    response = jsonify(payload)
    response.headers['Cache-Control'] = 'no-store'
    return response


@student_attendance_bp.route('/student-attendance/live/<live_key>/mark', methods=['POST'])
def live_mark(live_key):
    data = request.get_json(silent=True) or {}
    conn = get_db()
    try:
        info = _class_by_key(conn, 'live_key', live_key)
        if not info:
            return _json_error('사용할 수 없는 링크입니다.', 404)
        message, error = _mark(conn, info, _int_or_none(data.get('student_id')),
                               _text(data.get('status'), 10), _today(), _live_actor(info),
                               note=data.get('note'))
    finally:
        conn.close()
    return _json_error(error) if error else jsonify({'status': 'success', 'message': message})


@student_attendance_bp.route('/student-attendance/live/<live_key>/dismiss', methods=['POST'])
def live_dismiss(live_key):
    conn = get_db()
    try:
        info = _class_by_key(conn, 'live_key', live_key)
        if not info:
            return _json_error('사용할 수 없는 링크입니다.', 404)
        count = _dismiss_all(conn, info, _live_actor(info))
    finally:
        conn.close()
    return jsonify({'status': 'success', 'count': count, 'message': _bulk_message(info, count, '하원')})


@student_attendance_bp.route('/student-attendance/live/<live_key>/absent-rest', methods=['POST'])
def live_absent_rest(live_key):
    conn = get_db()
    try:
        info = _class_by_key(conn, 'live_key', live_key)
        if not info:
            return _json_error('사용할 수 없는 링크입니다.', 404)
        count = _absent_rest(conn, info, _live_actor(info))
    finally:
        conn.close()
    return jsonify({'status': 'success', 'count': count, 'message': _bulk_message(info, count, '결석')})


@student_attendance_bp.route('/student-attendance/live/<live_key>/reset-device', methods=['POST'])
def live_reset_device(live_key):
    """휴대폰·앱을 바꿔 막힌 학생을 강사가 교실에서 바로 풀어 준다. (자기 과목 학생만)"""
    student_id = _int_or_none((request.get_json(silent=True) or {}).get('student_id'))
    conn = get_db()
    try:
        info = _class_by_key(conn, 'live_key', live_key)
        if not info:
            return _json_error('사용할 수 없는 링크입니다.', 404)
        student = conn.execute('''
            SELECT s.name FROM parent_class_students cs JOIN parent_students s ON s.id=cs.student_id
            WHERE cs.class_id=? AND cs.student_id=?
        ''', (info['id'], student_id)).fetchone()
        if not student:
            return _json_error('이 과목의 학생이 아닙니다.', 404)
        conn.execute('DELETE FROM student_att_devices WHERE student_id=?', (student_id,))
        conn.commit()
    finally:
        conn.close()
    _emit_change(info, student_id)
    return jsonify({'status': 'success', 'message': f"{student['name']} 학생의 {RESET_DEVICE_MESSAGE}"})


# ---------------------------------------------------------------- 교실 QR 화면(로그인 없음, 조작 불가)

@student_attendance_bp.route('/student-attendance/display/<display_key>')
def display_page(display_key):
    conn = get_db()
    try:
        info = _class_by_key(conn, 'display_key', display_key)
    finally:
        conn.close()
    if not info:
        return render_template('student_attendance/display.html', invalid=True), 404
    today = _now().date()
    response = Response(render_template(
        'student_attendance/display.html', invalid=False, info=info, display_key=display_key,
        today=today.isoformat(), weekday=WEEKDAY_LABELS[today.weekday()],
        refresh_seconds=LIVE_QR_WINDOW_SECONDS,
    ))
    response.headers['Cache-Control'] = 'no-store'
    response.headers['Referrer-Policy'] = 'no-referrer'
    return response


@student_attendance_bp.route('/student-attendance/display/<display_key>/qr.svg')
def display_qr_svg(display_key):
    return _qr_by_key('display_key', display_key)


@student_attendance_bp.route('/student-attendance/display/<display_key>/summary')
def display_summary(display_key):
    """교실 화면에는 인원수만 준다. 학생 명단·연락처는 내보내지 않는다."""
    conn = get_db()
    try:
        info = _class_by_key(conn, 'display_key', display_key)
        if not info:
            return _json_error('사용할 수 없는 링크입니다.', 404)
        summary = _summary(_day_items(conn, info['id'], _today()))
    finally:
        conn.close()
    response = jsonify({'status': 'success', 'qr_mode': info['qr_mode'], 'summary': summary})
    response.headers['Cache-Control'] = 'no-store'
    return response


# ---------------------------------------------------------------- 학생 휴대폰 스캔(로그인 없음)

def _scan_error_page(message, info=None):
    today = _now().date()
    return render_template('student_attendance/scan.html', error=message, info=info,
                           today=today.isoformat(), weekday=WEEKDAY_LABELS[today.weekday()])


@student_attendance_bp.route('/student-attendance/s/<int:class_id>/<token>')
def scan_page(class_id, token):
    conn = get_db()
    try:
        info = _load_class(conn, class_id)
        if not info:
            body = _scan_error_page('등록되지 않은 수업 QR입니다.')
        elif not _scan_token_valid(conn, info, token):
            body = _scan_error_page(
                '지금 사용할 수 있는 QR이 아닙니다. 교실 화면에 떠 있는 QR을 다시 찍어주세요.'
                if info['qr_mode'] == 'live' else
                '오늘 발급된 QR이 아닙니다. 교실에 있는 오늘의 QR을 다시 찍어주세요.', info)
        else:
            today = _now().date()
            body = render_template(
                'student_attendance/scan.html', error=None, info=info, ticket=_make_ticket(class_id),
                today=today.isoformat(), weekday=WEEKDAY_LABELS[today.weekday()],
            )
    finally:
        conn.close()
    response = Response(body)
    response.headers['Cache-Control'] = 'no-store'
    response.headers['Referrer-Policy'] = 'no-referrer'
    return response


def _device_students(conn, device_token, class_id):
    """이 기기에 연결된 학생 중 이 과목 수강생. (수강생 목록, 기기에 연결된 학생 수)"""
    if not device_token:
        return [], 0
    rows = conn.execute('''
        SELECT d.student_id, s.name,
               EXISTS(SELECT 1 FROM parent_class_students cs
                      WHERE cs.class_id=? AND cs.student_id=d.student_id) AS enrolled
        FROM student_att_devices d
        JOIN parent_students s ON s.id=d.student_id AND s.is_active=1
        WHERE d.token_hash=?
        ORDER BY s.name
    ''', (class_id, _hash_token(device_token))).fetchall()
    return [row for row in rows if row['enrolled']], len(rows)


def _register_device(conn, class_info, name, code, device_token):
    """이름 + 보호자 휴대폰 뒷4자리로 수강생을 찾아 이 기기에 연결한다. (학생행, 새 기기키, 오류)"""
    name = re.sub(r'\s+', '', name)
    code = re.sub(r'\D', '', code)
    if not name or len(code) != 4:
        return None, None, ('이름과 보호자 휴대폰 뒷 4자리를 입력해주세요.', 400)
    rows = conn.execute('''
        SELECT DISTINCT s.id, s.name FROM parent_class_students cs
        JOIN parent_students s ON s.id=cs.student_id AND s.is_active=1
        JOIN parent_guardian_students gs ON gs.student_id=s.id
        JOIN parent_guardians g ON g.id=gs.guardian_id AND g.is_active=1
        WHERE cs.class_id=? AND REPLACE(s.name, ' ', '')=? AND substr(g.phone, -4)=?
    ''', (class_info['id'], name, code)).fetchall()
    if len(rows) != 1:
        message = ('같은 정보의 학생이 여러 명입니다. 선생님께 말씀해주세요.' if rows else
                   f"{class_info['class_name']} 수강생 정보와 일치하지 않습니다. "
                   '이름과 보호자 휴대폰 뒷 4자리를 확인하거나 선생님께 말씀해주세요.')
        return None, None, (message, 404)
    student = rows[0]
    token_hash = _hash_token(device_token) if device_token else None
    bound = conn.execute('SELECT token_hash FROM student_att_devices WHERE student_id=?',
                         (student['id'],)).fetchone()
    if bound and bound['token_hash'] != token_hash:
        return None, None, ('이 학생은 이미 다른 휴대폰(또는 다른 앱)에 등록되어 있어요. '
                            '선생님께 [휴대폰 다시 등록]을 눌러 달라고 말씀해 주세요.', 409)
    new_token = None
    if not device_token:
        new_token = device_token = secrets.token_urlsafe(32)
        token_hash = _hash_token(device_token)
    if not bound:
        conn.execute('''
            INSERT INTO student_att_devices (student_id, token_hash, user_agent, bound_at, last_seen_at)
            VALUES (?, ?, ?, ?, ?)
        ''', (student['id'], token_hash, _text(request.headers.get('User-Agent'), 200), _now_text(), _now_text()))
    return student, new_token, None


def _check_in(conn, class_info, student):
    """QR 출석을 기록한다. (이미출석 여부, 상태, 시각)"""
    today, now = _today(), _now()
    now_text = now.strftime('%Y-%m-%d %H:%M:%S')
    conn.execute('UPDATE student_att_devices SET last_seen_at=? WHERE student_id=?', (now_text, student['id']))
    existing = conn.execute(
        'SELECT status, check_in_at FROM student_att_records WHERE class_id=? AND student_id=? AND att_date=?',
        (class_info['id'], student['id'], today),
    ).fetchone()
    if existing and existing['status'] in CHECK_STATUSES:
        conn.commit()
        return True, existing['status'], (existing['check_in_at'] or '')[11:16]
    status = _auto_status(class_info, now)
    # 강사가 미리 결석으로 표시했더라도 학생이 도착해 찍으면 출석으로 바꾼다.
    conn.execute('''
        INSERT INTO student_att_records
            (class_id, student_id, att_date, status, check_in_at, method, recorded_by, updated_at)
        VALUES (?, ?, ?, ?, ?, 'QR', '학생QR', ?)
        ON CONFLICT(class_id, student_id, att_date) DO UPDATE SET
            status=excluded.status, check_in_at=excluded.check_in_at, method='QR',
            recorded_by='학생QR', note=NULL, updated_at=excluded.updated_at
    ''', (class_info['id'], student['id'], today, status, now_text, now_text))
    conn.commit()
    _emit_change(class_info, student['id'], status, student['name'])
    _notify(class_info, student['id'], today, status, now_text[11:16], '학생QR')
    return False, status, now_text[11:16]


@student_attendance_bp.route('/student-attendance/s/<int:class_id>', methods=['POST'])
def scan_check_in(class_id):
    data = request.get_json(silent=True) or {}
    device_token = request.cookies.get(DEVICE_COOKIE, '')
    new_token = None
    conn = get_db()
    try:
        info = _load_class(conn, class_id)
        if not info:
            return _json_error('등록되지 않은 수업입니다.', 404)
        if not _ticket_valid(data.get('ticket'), class_id):
            return _json_error('QR을 찍은 지 오래되었습니다. 교실의 QR을 다시 찍어주세요.', 410)

        students, device_count = _device_students(conn, device_token, class_id)
        name, code = _text(data.get('name'), 40), _text(data.get('code'), 10)
        if name or code:
            limit_key = _client_key(f'scan:{class_id}')
            if _too_many_failures(limit_key):
                return _json_error('이름·번호가 여러 번 맞지 않아 잠시 막혔어요. '
                                   '10분 뒤 다시 하거나 선생님께 말씀해 주세요.', 429)
            student, new_token, error = _register_device(conn, info, name, code, device_token)
            if error:
                conn.rollback()
                if error[1] == 404:
                    _record_failure(limit_key)
                return _json_error(*error)
        elif not students:
            # 처음 쓰는 휴대폰이면 화면의 기본 안내를 그대로 보여준다.
            message = (f"이 휴대폰에 등록된 학생은 {info['class_name']} 수강생이 아니에요. 내 이름으로 확인해 주세요."
                       if device_count else '')
            return jsonify({'status': 'need_register', 'message': message})
        elif len(students) > 1 and not data.get('student_id'):
            return jsonify({'status': 'choose', 'students': [
                {'id': row['student_id'], 'name': row['name']} for row in students
            ]})
        else:
            chosen = _int_or_none(data.get('student_id'))
            match = [row for row in students if row['student_id'] == chosen] if chosen else students[:1]
            if not match:
                return _json_error('이 휴대폰에 등록된 학생이 아닙니다.', 403)
            student = {'id': match[0]['student_id'], 'name': match[0]['name']}

        already, status, time_text = _check_in(conn, info, student)
        # 형제·자매가 같은 휴대폰을 쓰면 완료 화면에서 '다른 학생도 출석하기'를 보여준다.
        device_student_count = len(_device_students(conn, new_token or device_token, class_id)[0])
    finally:
        conn.close()

    response = jsonify({
        'status': 'success',
        'already': already,
        'attendance': status,
        'name': student['name'],
        'class_name': info['class_name'],
        'check_in_time': time_text,
        'device_registered': bool(new_token),
        'device_student_count': device_student_count,
        'parent_notified': bool(info['notify_parents']) and not already,
        'message': '이미 출석 처리되었습니다.' if already else f'{status} 처리되었습니다.',
    })
    if new_token:
        response.set_cookie(
            DEVICE_COOKIE, new_token, max_age=DEVICE_COOKIE_MAX_AGE, httponly=True, samesite='Lax',
            path='/student-attendance/s',
            secure=request.is_secure or request.headers.get('X-Forwarded-Proto', '').startswith('https'),
        )
    return response


# ---------------------------------------------------------------- 학부모 알림 등록 안내(로그인 없음)

@student_attendance_bp.route('/student-attendance/p/<parent_key>', methods=['GET', 'POST'])
def parent_lookup(parent_key):
    """가정통신문 QR로 들어온 보호자를 명단과 대조해 개인 알림등록 페이지로 보낸다.

    출결 알림은 아이의 위치 정보이므로, 관리자가 등록한 보호자 번호와 일치할 때만 연결한다.
    """
    conn = get_db()
    try:
        info = _class_by_key(conn, 'parent_key', parent_key)
        if not info:
            return render_template('student_attendance/parent.html', invalid=True), 404
        error = ''
        if request.method == 'POST':
            child = re.sub(r'\s+', '', _text(request.form.get('child_name'), 40))
            phone = normalize_phone(request.form.get('phone'))
            limit_key = _client_key(f"parent:{info['id']}")
            if _too_many_failures(limit_key):
                error = '입력한 정보가 여러 번 맞지 않아 잠시 막혔습니다. 10분 뒤 다시 시도해주세요.'
            elif not child or not re.fullmatch(r'01\d{8,9}', phone):
                error = '자녀 이름과 보호자 휴대폰번호를 정확히 입력해주세요.'
            else:
                guardian = conn.execute('''
                    SELECT g.id FROM parent_class_students cs
                    JOIN parent_students s ON s.id=cs.student_id AND s.is_active=1
                    JOIN parent_guardian_students gs ON gs.student_id=s.id
                    JOIN parent_guardians g ON g.id=gs.guardian_id AND g.is_active=1
                    WHERE cs.class_id=? AND REPLACE(s.name, ' ', '')=? AND g.phone=?
                    LIMIT 1
                ''', (info['id'], child, phone)).fetchone()
                if guardian:
                    invite = pn._active_invite(conn, guardian['id'])
                    return redirect(url_for('parent_notifications.parent_register', token=invite['token']))
                _record_failure(limit_key)
                error = ('수강생 명단의 보호자 정보와 일치하지 않습니다. '
                         '방과후 신청 때 적은 휴대폰번호로 입력하거나 학교 담당자에게 문의해주세요.')
        response = Response(render_template('student_attendance/parent.html', invalid=False, info=info,
                                            error=error, form=request.form))
        response.headers['Cache-Control'] = 'no-store'
        return response
    finally:
        conn.close()


# ---------------------------------------------------------------- 실시간 소켓

@socketio.on('connect', namespace=SOCKET_NS)
def student_attendance_socket_connect(auth=None):
    auth = auth if isinstance(auth, dict) else {}
    conn = get_db()
    try:
        for column in ('live_key', 'display_key'):
            key = _text(auth.get(column), 100)
            if key:
                info = _class_by_key(conn, column, key)
                if not info:
                    return False
                join_room(f"class:{info['id']}")
                return True
    finally:
        conn.close()
    school_id = _int_or_none(auth.get('school_id'))
    if not session.get('emp_no') or not has_menu_permission(MENU_KEY) or not school_id:
        return False
    join_room(f'school:{school_id}')
    return True
