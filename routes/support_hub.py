"""학교지원 > 학부모·강사 통합지원.

예전의 [강사통합지원]과 [학부모통합지원]을 한 메뉴로 합친 본사 화면이다. 학교를 하나 고르면 그 학교의
강사·과목·학생·학부모 안내를 한곳에서 입력하고 관리한다. 학부모가 들어오는 입구는 이 화면에서 만든
학교 포털(/parent/hub/…) 하나뿐이다.

탭
1. 강사 등록·관리 : 강사를 직접 등록(한 명 / 붙여넣기 / 엑셀)하고 접속 링크·비밀번호·문자를 관리한다. (instructor_hub)
2. 과목·학생     : 과목(강의부서·강사·요일·시간)과 수강생·보호자 명단을 입력하는 **유일한 화면**.
                   엑셀 한 장으로 학교 전체 과목·명단을 올릴 수도 있다. (student_attendance의 함수를 그대로 쓴다)
3. 승인관리       : 강사 한 명이 한 줄 — 프로필과 홍보 리플렛 자료를 함께 보고 승인·반려한다. (instructor_hub)
4. 학부모 안내    : 학교 포털(링크·QR·인쇄용 안내문), 링크 발송(문자·알림톡), 학교 e리플렛 앞·뒤쪽. (parent_hub)

[학생출석관리]는 출석 운영(실시간·월별·교실 QR)만, [학부모알림전송]은 푸시알림 발송·이력만 맡는다.
이 모듈은 화면 구성과 과목·명단 입력 경로만 갖고, 나머지 처리는 위 모듈의 함수를 부른다.
"""

from __future__ import annotations

import io
import re

from flask import (
    Blueprint, abort, flash, redirect, render_template, request, send_file, url_for,
)

from . import instructor_hub as hub
from . import parent_hub as ph
from . import solapi_settings as solapi
from . import student_attendance as sa
from .database import get_db
from .security import menu_permission_required

support_hub_bp = Blueprint('support_hub', __name__)

MENU_KEY = 'school_support_hub'
TABS = (
    ('instructors', '강사 등록·관리', 'fa-user-plus'),
    ('roster', '과목·학생', 'fa-children'),
    ('review', '승인관리', 'fa-circle-check'),
    ('portal', '학부모 안내', 'fa-paper-plane'),
)
TAB_KEYS = {key for key, _label, _icon in TABS}
TAB_ALIASES = {'register': 'instructors', 'schools': 'instructors', 'links': 'instructors',
               'portals': 'portal', 'leaflet': 'portal'}
MAX_ROSTER_ROWS = 1000
MAX_ROSTER_BYTES = 3 * 1024 * 1024
PHONE_PATTERN = re.compile(r'01\d{8,9}')

_text = hub._text
_int = hub._int
_norm = hub.norm


def _url(tab='instructors', school=None, **extra):
    return url_for('support_hub.admin_page', tab=tab, school=school or None, **extra)


def _back(default_tab='roster', school='', **extra):
    """처리를 마친 뒤 돌아갈 곳. 보던 화면(next)이 통합지원 안이면 그대로, 아니면 학교·탭을 맞춰 보낸다."""
    target = request.form.get('next') or ''
    if not extra and target.startswith(('/support-hub', '/instructor-hub', '/parent-hub')) \
            and not target.startswith('//'):
        return redirect(target)
    return redirect(_url(default_tab, school, **extra))


# ---------------------------------------------------------------- 화면 자료

def _school_state(conn):
    """선택 상자와 전체 현황에 쓰는 학교별 요약(강사·과목·학생·보호자·포털·승인 대기)."""
    schools, legacy, pending_total = hub._school_choices(conn)
    stats, _classes = ph._school_stats(conn)
    portals = ph._portals_by_school(conn)
    zero = {'students': 0, 'guardians': 0, 'push': 0, 'classes': 0}
    for item in schools + legacy:
        key = _norm(item['name'])
        item.update({k: v for k, v in stats.get(key, zero).items() if k != 'classes'})
        portal = portals.get(key)
        item['portal'] = None
        if portal:
            pages, _toc = ph.leaflet_pages(conn, portal)
            item['portal'] = {'active': bool(portal['is_active']), 'leaflet_pages': len(pages),
                              'last_sent_at': portal['last_sent_at'] or '', 'last_sent_count': portal['last_sent_count'] or 0}
    return schools, legacy, pending_total


def _subject_options(conn, school_name, edit_class=None):
    """과목명 선택 상자의 항목. 수정 중인 과목의 예전 과목명이 목록에 없어도 사라지지 않게 더한다."""
    options = list(hub.school_departments(conn, school_name))
    old = (edit_class or {}).get('department') or ''
    if old and old not in options:
        options.append(old)
    return options


def _roster_context(conn, selected):
    """② 과목·학생 탭: 이 학교의 과목 목록, 고른 과목의 명단, 과목 입력 양식."""
    school_name = selected['name']
    key = _norm(school_name)
    classes = []
    for row in conn.execute('SELECT id, school_id, school_name FROM parent_classes WHERE is_active=1 '
                            'ORDER BY class_name COLLATE NOCASE, id').fetchall():
        if _norm(row['school_name']) == key or (selected['school_id'] and row['school_id'] == selected['school_id']):
            classes.append(sa._class_payload(conn, row))
    wanted = _int(request.args.get('class'))
    current = next((c for c in classes if c['id'] == wanted), classes[0] if classes else None)
    students = sa._roster(conn, current['id']) if current else []
    edit = _int(request.args.get('edit'))
    edit_class = next((c for c in classes if c['id'] == edit), None) if edit else None
    return {
        'classes': classes, 'current': current, 'students': students,
        'edit_class': edit_class,
        'edit_student': _int(request.args.get('student')),
        'hub_people': hub.school_instructors(conn, school_name),
        'departments': hub._department_suggestions(conn, school_name),
        # 과목명 선택 상자: [강사 등록·관리]에 이 학교 강사로 등록한 강의부서
        'subject_options': _subject_options(conn, school_name, edit_class),
        'weekday_labels': sa.WEEKDAY_LABELS,
        'can_edit': selected['matched'],
    }


def _kakao_info(conn):
    """학부모 안내 알림톡이 설정돼 있는지와 변수 표기(솔라피설정 저장값만 읽고 솔라피에는 묻지 않는다)."""
    try:
        settings = solapi.get_settings(conn)
        return bool(settings.get('parent_kakao_configured')), ph.link_variable_keys(settings.get('parent_variable_names'))
    except Exception:  # 암호화 키 문제 등으로 설정을 못 읽어도 화면은 열려야 한다
        return False, ph.link_variable_keys()


def _portal_context(conn, selected):
    """④ 학부모 안내 탭: 포털·링크 발송 대상·리플렛·준비 점검표."""
    school_name = selected['name']
    portal = ph.portal_for_school(conn, school_name, active_only=False)
    context = {'portal': None, 'checks': []}
    guide = None
    pages, toc = [], []
    if portal:
        kakao_ready, kakao_var_keys = _kakao_info(conn)
        guardians = ph.school_guardians(conn, portal)
        pages, toc = ph.leaflet_pages(conn, portal)
        guide = ph._guide_for_school(conn, portal)
        context.update(
            portal=dict(portal), portal_url=ph._portal_url(portal),
            audience=ph.link_audience(guardians), guardians=guardians,
            designs=[{'key': key, 'label': info['label'], 'tagline': info['tagline'], 'points': info['points']}
                     for key, info in ph.DESIGNS.items()],
            current_design=ph.design_of(portal),
            link_targets=ph.LINK_TARGETS, default_intro=ph.DEFAULT_LINK_INTRO,
            link_methods=ph.LINK_METHODS, kakao_ready=kakao_ready, kakao_var_keys=kakao_var_keys,
            kakao_preview_url=url_for('parent_hub.admin_kakao_preview', portal_id=portal['id']),
            solapi_settings_url='/admin/solapi-settings',
            toc=toc, total_pages=len(pages), sections=ph.LEAFLET_SECTIONS, max_pages=ph.MAX_HQ_PAGES,
            hq={section: [dict(r) for r in ph._hq_pages(conn, portal['id'], section)] for section in ph.LEAFLET_SECTIONS},
            coverage=ph._instructor_coverage(conn, school_name),
            guide_ready=bool(guide),
            preview_url=url_for('parent_hub.hub_leaflet', token=portal['token']),
            school_url=url_for('classroom_guide.index', school=selected['school_id']) if selected['school_id'] else '',
        )
    sent = (portal['last_sent_at'] if portal else '') or ''
    context['checks'] = [
        {'label': '강사 등록', 'ok': selected['instructors'] > 0, 'note': f"{selected['instructors']}명", 'tab': 'instructors'},
        {'label': '과목', 'ok': selected['classes'] > 0, 'note': f"{selected['classes']}개", 'tab': 'roster'},
        {'label': '수강생·보호자', 'ok': selected['students'] > 0,
         'note': f"학생 {selected['students']}명 · 보호자 {selected['guardians']}명", 'tab': 'roster'},
        {'label': '학교 리플렛', 'ok': bool(pages), 'note': f'{len(pages)}쪽', 'tab': 'portal', 'optional': True},
        {'label': '3D 교실안내', 'ok': bool(guide), 'note': '연결됨' if guide else '미설정', 'tab': 'portal', 'optional': True},
        {'label': '포털 공개', 'ok': bool(portal and portal['is_active']),
         'note': ('공개 중' if portal and portal['is_active'] else ('비공개' if portal else '미개설')), 'tab': 'portal'},
        {'label': '링크 발송', 'ok': bool(sent), 'note': (sent[:10] + f" · {portal['last_sent_count']}명") if sent else '아직 안 보냄',
         'tab': 'portal'},
    ]
    return context


@support_hub_bp.route('/support-hub/')
@support_hub_bp.route('/support-hub')
@menu_permission_required(MENU_KEY)
def admin_page():
    tab = request.args.get('tab')
    tab = TAB_ALIASES.get(tab, tab)
    tab = tab if tab in TAB_KEYS else 'instructors'
    conn = get_db()
    try:
        schools, legacy, pending_total = _school_state(conn)
        selected = hub._selected_school(schools, legacy)
        context = {'tab': tab, 'tabs': TABS, 'schools': schools, 'legacy': legacy, 'selected': selected,
                   'pending_total': pending_total, 'status_labels': hub.STATUS_LABELS,
                   'max_rows': hub.MAX_IMPORT_ROWS, 'max_roster_rows': MAX_ROSTER_ROWS,
                   'push_configured': ph.pn._push_ready()}
        extra = {}
        if selected:
            if tab == 'instructors':
                context.update(groups=hub._school_instructors(conn, selected['name']),
                               departments=hub._department_suggestions(conn, selected['name']))
            elif tab == 'roster':
                context.update(_roster_context(conn, selected))
                if context['current']:
                    extra['class'] = context['current']['id']
            elif tab == 'review':
                context.update(review=hub._review_rows(conn, selected['name']))
            else:
                context.update(_portal_context(conn, selected))
        elif tab == 'review':
            context.update(pending_schools=[s for s in schools + legacy if s['pending']])
        else:
            context.update(everyone=hub._everyone(conn), orphans=hub._orphan_accounts(conn))
            context['overview'] = sorted(
                schools + legacy,
                key=lambda s: (not (s['instructors'] or s['students'] or s['classes']), s['name']))
        context['here'] = _url(tab, selected['name'] if selected else '', **extra)
        return render_template('support_hub/admin.html', **context)
    finally:
        conn.close()


# ---------------------------------------------------------------- 과목 입력(② 과목·학생)

def _school_by_name(conn, school_name):
    school_id, name = hub.canonical_school(conn, school_name)
    return (school_id, name) if school_id else (None, name)


@support_hub_bp.route('/support-hub/classes/save', methods=['POST'])
@menu_permission_required(MENU_KEY)
def save_class():
    """과목을 만들거나 고친다(양식 한 장)."""
    conn = get_db()
    try:
        school_id, school_name = _school_by_name(conn, request.form.get('school_name'))
        if not _text(request.form.get('department'), 40):
            flash('과목명을 골라 주세요. 목록이 비어 있으면 [강사 등록·관리]에서 강사를 먼저 등록해 주세요.', 'error')
            return _back('roster', school_name)
        if not _text(request.form.get('class_name'), 60):
            flash('강의반을 입력해 주세요. 예) 로봇과학A', 'error')
            return _back('roster', school_name)
        data = {
            'school_id': school_id, 'id': _int(request.form.get('id')),
            'class_name': request.form.get('class_name'), 'department': request.form.get('department'),
            'instructor_name': request.form.get('instructor_name'),
            'hub_account_id': request.form.get('hub_account_id', ''),
            'weekdays': [_int(v) for v in request.form.getlist('weekdays') if _int(v) is not None],
            'start_time': request.form.get('start_time'), 'end_time': request.form.get('end_time'),
            'late_minutes': request.form.get('late_minutes'),
            'qr_mode': request.form.get('qr_mode'), 'notify_parents': bool(request.form.get('notify_parents')),
        }
        class_id, error, _status = sa.save_class_data(conn, data)
    finally:
        conn.close()
    if error:
        flash(error, 'error')
        return _back('roster', school_name)
    flash(f"{_text(data['class_name'], 60)} 강의반을 저장했습니다. 아래에서 수강생을 등록해 주세요.", 'success')
    return _back('roster', school_name, **{'class': class_id})


@support_hub_bp.route('/support-hub/classes/<int:class_id>/delete', methods=['POST'])
@menu_permission_required(MENU_KEY)
def delete_class(class_id):
    conn = get_db()
    try:
        info = sa._admin_class(conn, class_id)
        if not info:
            abort(404)
        sa.deactivate_class(conn, class_id)
    finally:
        conn.close()
    flash(f"{info['class_name']} 과목을 삭제했습니다. 지난 출석 기록은 보존됩니다.", 'success')
    return _back('roster', info['school_name'])


def _class_or_404(conn, class_id):
    info = sa._admin_class(conn, class_id)
    if not info:
        abort(404)
    return info


@support_hub_bp.route('/support-hub/classes/<int:class_id>/students/save', methods=['POST'])
@menu_permission_required(MENU_KEY)
def save_student(class_id):
    """학생 한 명을 등록하거나(student_id 없음) 고친다."""
    conn = get_db()
    try:
        info = _class_or_404(conn, class_id)
        error = sa._save_student(conn, info, {
            'name': request.form.get('name'), 'grade': request.form.get('grade'),
            'classroom': request.form.get('classroom'), 'guardian_name': request.form.get('guardian_name'),
            'guardian_phone': request.form.get('guardian_phone'), 'guardian_id': request.form.get('guardian_id'),
        }, _int(request.form.get('student_id')))
        if error:
            conn.rollback()
        else:
            conn.commit()
    finally:
        conn.close()
    flash(error or '학생 정보를 저장했습니다.', 'error' if error else 'success')
    return _back('roster', info['school_name'], **{'class': class_id})


@support_hub_bp.route('/support-hub/classes/<int:class_id>/students/bulk', methods=['POST'])
@menu_permission_required(MENU_KEY)
def bulk_students(class_id):
    """'학생명, 학년, 반, 보호자명, 보호자연락처' 여러 줄을 한 과목 명단에 넣는다(하나라도 틀리면 아무도 넣지 않는다)."""
    rows = sa.parse_roster_text(request.form.get('text'))
    conn = get_db()
    try:
        info = _class_or_404(conn, class_id)
        if not rows:
            flash('등록할 학생 명단을 입력해 주세요.', 'error')
        elif len(rows) > 300:
            flash('한 번에 300명까지 등록할 수 있습니다. 나누어 올려 주세요.', 'error')
        else:
            errors = sa.register_students(conn, info, rows)
            if errors:
                conn.rollback()
                flash('명단을 확인해 주세요. 아무도 등록하지 않았습니다.\n' + '\n'.join(errors[:8]), 'error')
            else:
                conn.commit()
                flash(f'학생 {len(rows)}명을 등록했습니다.', 'success')
    finally:
        conn.close()
    return _back('roster', info['school_name'], **{'class': class_id})


@support_hub_bp.route('/support-hub/classes/<int:class_id>/students/<int:student_id>/remove', methods=['POST'])
@menu_permission_required(MENU_KEY)
def remove_student(class_id, student_id):
    conn = get_db()
    try:
        info = _class_or_404(conn, class_id)
        conn.execute('DELETE FROM parent_class_students WHERE class_id=? AND student_id=?', (class_id, student_id))
        conn.commit()
    finally:
        conn.close()
    flash('과목 명단에서 뺐습니다. 지난 출석 기록은 보존됩니다.', 'success')
    return _back('roster', info['school_name'], **{'class': class_id})


@support_hub_bp.route('/support-hub/classes/<int:class_id>/students/<int:student_id>/reset-device', methods=['POST'])
@menu_permission_required(MENU_KEY)
def reset_device(class_id, student_id):
    conn = get_db()
    try:
        info = _class_or_404(conn, class_id)
        conn.execute('DELETE FROM student_att_devices WHERE student_id=?', (student_id,))
        conn.commit()
    finally:
        conn.close()
    flash(sa.RESET_DEVICE_MESSAGE, 'success')
    return _back('roster', info['school_name'], **{'class': class_id})


# -- 학교 전체 엑셀(과목·학생·보호자를 한 장으로)

ROSTER_HEADERS = {
    'class_name': ('과목', '과목명', '강좌', '강좌명', '수업명', '프로그램명', '방과후강좌'),
    'department': ('강의부서', '부서', '분야', '강의분야'),
    'name': ('학생명', '학생이름', '학생', '이름'),
    'grade': ('학년',),
    'classroom': ('반', '학급'),
    'guardian_name': ('보호자명', '보호자이름', '학부모명', '보호자', '학부모'),
    'guardian_phone': ('보호자휴대폰', '보호자연락처', '학부모연락처', '휴대폰', '휴대폰번호', '전화번호', '연락처'),
}
ROSTER_COLUMNS = ('class_name', 'department', 'name', 'grade', 'classroom', 'guardian_name', 'guardian_phone')


def _roster_header(value):
    text = re.sub(r'\(.*?\)|\s+', '', str(value or '')).lower()
    for key, names in ROSTER_HEADERS.items():
        if text in names:
            return key
    return ''


def parse_roster_workbook(upload):
    """엑셀(.xlsx) 첫 시트를 과목·학생 줄로 읽는다. 머리글이 있으면 칸 순서는 상관없다."""
    from openpyxl import load_workbook
    name = str(getattr(upload, 'filename', '') or '')
    if not name.lower().endswith('.xlsx'):
        raise ValueError('엑셀(.xlsx) 파일만 올릴 수 있습니다. .xls 파일은 [다른 이름으로 저장]에서 .xlsx로 바꿔 주세요.')
    data = upload.stream.read(MAX_ROSTER_BYTES + 1)
    if len(data) > MAX_ROSTER_BYTES:
        raise ValueError('엑셀 파일은 3MB까지 올릴 수 있습니다.')
    try:
        sheet_rows = [list(row) for row in load_workbook(io.BytesIO(data), read_only=True, data_only=True)
                      .worksheets[0].iter_rows(values_only=True)]
    except Exception as exc:
        raise ValueError('엑셀 파일을 읽을 수 없습니다. 파일이 깨졌는지 확인해 주세요.') from exc
    rows, mapping, first = [], None, True
    for number, cells in enumerate(sheet_rows, 1):
        if not any(hub._cell_text(cell) for cell in cells):
            continue
        if first:
            first = False
            keys = [_roster_header(cell) for cell in cells]
            if 'name' in keys and 'guardian_phone' in keys:
                mapping = {key: index for index, key in enumerate(keys) if key}
                continue
        if mapping:
            values = [cells[mapping[key]] if key in mapping and mapping[key] < len(cells) else '' for key in ROSTER_COLUMNS]
        else:
            values = (cells + [''] * len(ROSTER_COLUMNS))[:len(ROSTER_COLUMNS)]
        row = {key: hub._cell_text(values[i]) for i, key in enumerate(ROSTER_COLUMNS)}
        row['line'] = number
        rows.append(row)
    return rows


def import_roster(conn, school_id, rows):
    """과목 칸이 있는 줄들을 학교의 과목·학생·보호자로 등록한다. 하나라도 틀리면 아무것도 등록하지 않는다.

    반환: {'ok', 'errors', 'students', 'new_classes', 'classes'}
    """
    errors, prepared = [], []
    for row in rows:
        phone = hub.clean_phone(row['guardian_phone'])
        problems = []
        if not row['class_name']:
            problems.append('과목명이 없습니다')
        if not row['name']:
            problems.append('학생 이름이 없습니다')
        if not PHONE_PATTERN.fullmatch(phone):
            problems.append('보호자 휴대폰번호(010…)를 확인해 주세요')
        if problems:
            errors.append(f"{row['line']}번째 줄 {row['name'] or '(이름 없음)'}: " + ', '.join(problems))
        else:
            prepared.append({**row, 'guardian_phone': phone})
    if errors:
        return {'ok': False, 'errors': errors, 'students': 0, 'new_classes': 0, 'classes': 0}
    class_ids, new_classes = {}, 0
    existing = {_norm(r['class_name']): r['id'] for r in conn.execute(
        'SELECT id, class_name FROM parent_classes WHERE school_id=? AND is_active=1', (school_id,)).fetchall()}
    for row in prepared:
        key = _norm(row['class_name'])
        if key in class_ids:
            continue
        if key in existing:
            class_ids[key] = existing[key]
            continue
        class_id, error, _status = sa.save_class_data(conn, {
            'school_id': school_id, 'class_name': row['class_name'], 'department': row['department']})
        if error:
            errors.append(f"{row['line']}번째 줄 과목 ‘{row['class_name']}’: {error}")
            continue
        class_ids[key] = class_id
        new_classes += 1
    if errors:
        conn.rollback()
        return {'ok': False, 'errors': errors, 'students': 0, 'new_classes': 0, 'classes': 0}
    infos = {}
    for row in prepared:
        class_id = class_ids[_norm(row['class_name'])]
        info = infos.setdefault(class_id, sa._load_class(conn, class_id))
        error = sa._save_student(conn, info, row)
        if error:
            errors.append(f"{row['line']}번째 줄 {row['name']}: {error}")
    if errors:
        conn.rollback()
        return {'ok': False, 'errors': errors, 'students': 0, 'new_classes': new_classes, 'classes': len(class_ids)}
    conn.commit()
    return {'ok': True, 'errors': [], 'students': len(prepared), 'new_classes': new_classes, 'classes': len(class_ids)}


@support_hub_bp.route('/support-hub/roster/import', methods=['POST'])
@menu_permission_required(MENU_KEY)
def import_roster_route():
    upload = request.files.get('file')
    conn = get_db()
    try:
        school_id, school_name = _school_by_name(conn, request.form.get('school_name'))
        if not school_id:
            flash('[학교업무공간]에 등록된 학교만 과목을 만들 수 있습니다.', 'error')
            return _back('roster', school_name)
        try:
            rows = parse_roster_workbook(upload) if upload and upload.filename else []
        except ValueError as exc:
            flash(str(exc), 'error')
            return _back('roster', school_name)
        if not rows:
            flash('등록할 줄이 없습니다. 엑셀 양식에 과목·학생·보호자를 적어 올려 주세요.', 'error')
        elif len(rows) > MAX_ROSTER_ROWS:
            flash(f'한 번에 {MAX_ROSTER_ROWS}줄까지 등록할 수 있습니다. 나누어 올려 주세요.', 'error')
        else:
            result = import_roster(conn, school_id, rows)
            if result['ok']:
                flash(f"학생 {result['students']}명을 등록했습니다. (과목 {result['classes']}개 중 새로 만든 과목 {result['new_classes']}개)", 'success')
            else:
                flash('엑셀을 확인해 주세요. 아무것도 등록하지 않았습니다.\n' + '\n'.join(result['errors'][:8])
                      + (f"\n… 외 {len(result['errors']) - 8}줄" if len(result['errors']) > 8 else ''), 'error')
    finally:
        conn.close()
    return _back('roster', school_name)


@support_hub_bp.route('/support-hub/roster/template.xlsx')
@menu_permission_required(MENU_KEY)
def roster_template():
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.title = '과목·학생'
    sheet.append(['과목', '강의부서', '학생명', '학년', '반', '보호자명', '보호자휴대폰'])
    sheet.append(['로봇과학 A반', '로봇과학', '김민준', '3', '2', '김민준 어머니', '010-1234-5678'])
    sheet.append(['로봇과학 A반', '로봇과학', '이서연', '1', '4', '이서연 아버지', '010-2345-6789'])
    sheet.append(['창의미술', '미술', '김민준', '3', '2', '김민준 어머니', '010-1234-5678'])
    for column, width in zip('ABCDEFG', (18, 14, 12, 8, 8, 18, 18)):
        sheet.column_dimensions[column].width = width
    note = book.create_sheet('작성방법')
    for line in ('1. 한 줄에 한 명(과목 하나)입니다. 한 학생이 과목 두 개를 들으면 두 줄로 적어 주세요.',
                 '2. 과목·학생명·보호자휴대폰은 꼭 적어 주세요. 강의부서·학년·반·보호자명은 비워도 됩니다.',
                 '3. 같은 이름의 과목이 이미 있으면 그 과목에 학생을 더하고, 없으면 새로 만듭니다.',
                 '4. 강의부서는 [강사 등록·관리]에 쓴 이름과 똑같이 적어야 그 부서 강사와 자동으로 연결됩니다.',
                 '5. 예시 세 줄은 지우고 올려 주세요. (.xlsx 파일, 한 번에 1000줄까지 · 하나라도 틀리면 아무것도 등록되지 않습니다)'):
        note.append([line])
    note.column_dimensions['A'].width = 110
    stream = io.BytesIO()
    book.save(stream)
    stream.seek(0)
    return send_file(stream, as_attachment=True, download_name='과목·학생등록양식.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
